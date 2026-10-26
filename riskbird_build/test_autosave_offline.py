"""离线验证自动保存改造：每次保存两份文件（原数据合并 + 清洗后入库），不分平台。

覆盖：
1. 会话目录里只有两个固定文件名，不再出现 企知道.xlsx / 爱企查.xlsx；
2. 写盘失败（文件被 Excel 占用等）时旧文件内容完好、不留 .tmp，采集继续进行；
3. 采集线程一边写 results、GUI 一边取快照写 Excel，不炸也不串数据；
4. 上一轮还没写完时只保留最后一次请求，写完自动补一轮（停止/完成那次必然落盘）；
5. 清洗文件按入库模板出，条数与原数据一致（未找到的行被剔除）。
"""
import os
import threading
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt5.QtWidgets import QApplication  # noqa: E402

app = QApplication([])

import riskbird_gui as G  # noqa: E402

TMP = Path("/tmp/autosave_offline")
if TMP.exists():
    import shutil
    shutil.rmtree(str(TMP), ignore_errors=True)
TMP.mkdir(parents=True, exist_ok=True)


def make_record(i, platform):
    if platform == "企知道":
        return {"公司名": f"企知道测试公司{i}", "状态": "成功", "数据来源": "企知道",
                "法定代表人": f"法人{i}", "注册资本": "100万元", "成立日期": "2020-01-0{i%9+1}",
                "统一社会信用代码": f"91440{i:06d}MA5A{i:02d}", "电话": f"020-800{i:04d}",
                "邮箱": f"c{i}@example.com", "地址": f"广东省测试路{i}号"}
    return {"公司名": f"爱企查测试公司{i}", "状态": "成功", "数据来源": "爱企查",
            "法定代表人": f"法人{i}", "注册资本": "200万元", "成立日期": "2019-05-0{i%9+1}",
            "统一社会信用代码": f"91310{i:06d}MA7B{i:02d}", "电话": f"021-900{i:04d}",
            "经营范围": "研发、销售机电设备"}


w = G.MainWindow()
w._session_folder = TMP                       # 不污染真实桌面
w._get_or_create_session_folder = lambda: TMP


def drain(timeout=90):
    """等后台保存（含补保存）全部收尾"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        app.processEvents()
        worker = w._autosave_worker
        if worker is not None and worker.isRunning():
            worker.wait(5000)
            continue
        if w._autosave_pending is not None:
            time.sleep(0.05)
            continue
        return
    raise AssertionError("自动保存超时未结束")


# ── 1. 两份固定文件，不再分平台 ────────────────────────
w.qizhidao_crawler.results = [make_record(i, "企知道") for i in range(1, 21)]
w.aiqicha_crawler.results = [make_record(i, "爱企查") for i in range(1, 16)]
w.aiqicha_crawler.results.append({"公司名": "查无此公司", "状态": "未找到", "数据来源": "爱企查"})

w._auto_save("测试保存")
drain()

files = sorted(p.name for p in TMP.iterdir())
assert files == [G.RAW_FILE_NAME, G.CLEAN_FILE_NAME], f"目录里应只有两份固定文件：{files}"
assert not list(TMP.glob("*.tmp")), "不应留下临时文件"
from openpyxl import load_workbook
wb = load_workbook(TMP / G.RAW_FILE_NAME)
raw_rows = wb.active.max_row - 1
wb.close()
assert raw_rows == 36, f"原数据应有 36 条（20+16）：{raw_rows}"
wb = load_workbook(TMP / G.CLEAN_FILE_NAME)
clean_rows = wb.active.max_row
wb.close()
assert clean_rows > 0, "清洗文件应有数据"
print(f"① 两份固定文件 ✓ 原数据 {raw_rows} 条 / 清洗文件含 {clean_rows} 行（未找到不进库）")

# ── 2. 同名覆盖：第二次保存替换内容而不是堆文件 ─────────
w.qizhidao_crawler.results.append(make_record(99, "企知道"))
w._auto_save("第二次")
drain()
files = sorted(p.name for p in TMP.iterdir())
assert files == [G.RAW_FILE_NAME, G.CLEAN_FILE_NAME], f"覆盖式保存不该堆文件：{files}"
wb = load_workbook(TMP / G.RAW_FILE_NAME)
assert wb.active.max_row - 1 == 37, f"应覆盖为 37 条：{wb.active.max_row - 1}"
wb.close()
print("② 固定文件名反复覆盖 ✓")

# ── 3. 写盘失败（被占用）时旧文件完好、无 .tmp、不崩 ───────
before = (TMP / G.RAW_FILE_NAME).read_bytes()
real_save = None
import openpyxl
real_save = openpyxl.Workbook.save
logs = []
w._append_log = lambda m: logs.append(m)


class Locked(Exception):
    pass


def raise_permission(self, stream, *a, **kw):
    raise PermissionError("另一个程序正在使用此文件")


openpyxl.Workbook.save = raise_permission
try:
    w._auto_save("模拟被占用")
    drain()
finally:
    openpyxl.Workbook.save = real_save
assert (TMP / G.RAW_FILE_NAME).read_bytes() == before, "写盘失败时旧文件必须保持原样"
assert not list(TMP.glob("*.tmp")), f"失败后不能残留 .tmp：{list(TMP.glob('*.tmp'))}"
assert any("占用" in m for m in logs), f"应提示文件被占用：{logs}"
print("③ 被占用时：旧文件完好 + 无残留 + 有提示 ✓")

# ── 4. 采集线程边写边保存：快照不炸、不串数据 ─────────────
w.qizhidao_crawler.results_lock = threading.RLock()
w.aiqicha_crawler.results_lock = threading.RLock()
stop = threading.Event()
errors = []


def harvester():
    """模拟两个采集线程：不断追加记录，并往刚写入的那条 dict 里继续补字段"""
    i = 0
    while not stop.is_set():
        i += 1
        rec = make_record(1000 + i, "企知道")
        with w.qizhidao_crawler.results_lock:
            w.qizhidao_crawler.results.append(rec)
            target = w.qizhidao_crawler.results[-1]
        for k in ("官网", "简介", "所属行业", "参保人数"):
            target[k] = f"{k}-{i}"        # 同一条记录被逐步填写
        time.sleep(0.005)


def saver():
    """模拟 GUI 线程反复取快照写 Excel"""
    try:
        for _ in range(12):
            combined = w._all_results()
            G.write_merged_excel(combined, str(TMP / "并发写测试.xlsx"))
            time.sleep(0.01)
    except Exception as e:
        errors.append(e)


t = threading.Thread(target=harvester)
s = threading.Thread(target=saver)
t.start()
s.start()
s.join()
stop.set()
t.join()
assert not errors, f"并发下写盘出错：{errors}"
wb = load_workbook(TMP / "并发写测试.xlsx")
assert wb.active.max_row - 1 >= 37
wb.close()
print(f"④ 采集线程与保存并发：无异常 ✓（{len(w._all_results())} 条快照）")

# ── 5. 上一轮没写完时只补最后一次 ───────────────────────
(TMP / "并发写测试.xlsx").unlink()
w._append_log = G.MainWindow._append_log.__get__(w)   # 恢复真实日志方法
pending_cases = []
w._autosave_pending = None
big = [make_record(i, "企知道") for i in range(1, 80)]
w.qizhidao_crawler.results = list(big)
w.aiqicha_crawler.results = []
w._auto_save("第一轮")
first = w._autosave_worker
# 第一轮还在写就再点两次停止/暂停保存：只保留最后一次
w._auto_save("第二轮")
w._auto_save("第三轮·最后一次")
assert w._autosave_pending is not None, "上一轮未完成时应记录待补保存"
assert first is w._autosave_worker, "不应同时起两个写盘线程"
drain()
assert w._autosave_pending is None, "补保存应已消费"
wb = load_workbook(TMP / G.RAW_FILE_NAME)
assert wb.active.max_row - 1 == 79, f"最后一次请求的数据必须落盘：{wb.active.max_row - 1}"
wb.close()
files = sorted(p.name for p in TMP.iterdir())
assert files == [G.RAW_FILE_NAME, G.CLEAN_FILE_NAME], files
print("⑤ 保存请求不排队只保留最后一次，最终数据一致 ✓")

print("AUTOSAVE_OFFLINE_OK")
