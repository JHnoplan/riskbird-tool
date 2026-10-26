"""
qizhidao_crawler.py - 企知道 (qizhidao.com) 企业数据爬虫模块 v2
基于 NUXT 数据提取，搜索→详情页→全量字段采集
"""

import json
import os
import sys
import time
import re
import threading
import urllib.parse
from datetime import datetime
from pathlib import Path

import block_guard

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    NoSuchElementException, TimeoutException,
    StaleElementReferenceException, WebDriverException
)

try:
    from PyQt5.QtCore import QObject, pyqtSignal
    HAS_PYQT = True
except ImportError:
    HAS_PYQT = False


# ── Qt 信号类 ──
if HAS_PYQT:
    class CrawlerSignals(QObject):
        log_signal = pyqtSignal(str)
        progress_signal = pyqtSignal(int, int)
        crawl_finished_signal = pyqtSignal()
        login_status_signal = pyqtSignal(bool, str)
        manual_check_signal = pyqtSignal(str)
        blocked_signal = pyqtSignal(str)   # 被安全验证/风控拦截，需人工在浏览器完成验证
        block_cleared_signal = pyqtSignal(str)  # 人工完成验证，采集已自动继续
        login_required_signal = pyqtSignal(str)  # 未登录，需界面上确认去登录还是继续
        handoff_signal = pyqtSignal(str)  # 本路已放弃等人工，剩余名单转交另一平台（界面要撤掉红条）
else:
    class CrawlerSignals:
        pass


class QiZhiDaoCrawler:
    BASE_URL = "https://www.qizhidao.com"
    DETAIL_BASE = "https://qiye.qizhidao.com/company/"
    PLATFORM_LABEL = "企知道"

    def __init__(self, gui_callback=None):
        self.driver = None
        self.wait = None
        self._logged_in = False
        self._crawling = False
        self._stop_requested = False
        self._pause_requested = False
        self.results = []
        # 采集线程在写、GUI 线程在读（双标签并行时两个线程同时写），靠锁交接快照
        self.results_lock = threading.RLock()
        self.gui_callback = gui_callback
        self._browser_lock = threading.Lock()

        # 双标签并行采集：附加到共享 Chrome 的调试端口，本实例只操作自己独占的那个标签
        self.attach_port = None
        self.shared_session = False
        self._own_handle = None

        if HAS_PYQT:
            self.signals = CrawlerSignals()
        else:
            self.signals = None

        # Cookie 持久化
        self._data_dir = Path(__file__).parent / "riskbird_data"
        self._data_dir.mkdir(exist_ok=True)
        self._cookie_file = self._data_dir / "qizhidao_cookies.json"

        # 速率控制
        self.delay_between = 0.5
        self.max_retries = 2

        # 登录态闸门与空壳计数
        self._login_gate_obj = None
        self._shell_streak = 0
        self._shell_gate_passed = False

        # 夜间无人值守：爱企查撑不住时把剩余名单移交到这里，企知道取来接着采
        # （handoff 由双标签工作线程注入，单平台采集时一直是 None）
        self.handoff = None
        self.handoff_wait = 3600   # 自己名单跑完后最多再等爱企查落定这么久

    # ── 日志 ──
    def _log(self, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        line = f"[{ts}] {msg}"
        if self.signals:
            self.signals.log_signal.emit(line)
        if self.gui_callback:
            self.gui_callback(line)
        print(line)

    # ── 结果交接 / 拦截处理 ──
    def _add_result(self, result: dict):
        """线程安全地追加一条采集结果"""
        with self.results_lock:
            self.results.append(result)

    def snapshot_results(self) -> list:
        """给 GUI 用的快照：连每条记录的 dict 一起复制（单条记录是边采边写的，
        只拷列表会让写 Excel 时遍历 r.keys() 撞上 dict 变大报 RuntimeError）"""
        with self.results_lock:
            return [dict(r) for r in self.results]

    def _emit_blocked(self, reason: str):
        if self.signals:
            try:
                self.signals.blocked_signal.emit(reason or "页面被平台拦截")
            except Exception:
                pass

    def _detect_block(self):
        """只看自己这个窗口当前的页面是不是验证页"""
        if self.driver is None:
            return ""
        return block_guard.detect_selenium(self.driver)

    def _wait_unblock(self, reason: str, max_wait: int = 1800) -> bool:
        """被拦截时原地等人工验证，不推进名单、不写入空记录"""
        ok = block_guard.wait_unblock(
            reason, self._log, self._emit_blocked,
            lambda: self._stop_requested, self._detect_block,
            max_wait=max_wait,
        )
        if ok:
            self._emit_cleared(reason)
        return ok

    def _emit_cleared(self, reason: str):
        if self.signals:
            try:
                self.signals.block_cleared_signal.emit(reason)
            except Exception:
                pass

    # ── 接收爱企查移交的名单（夜间无人值守）──
    def _take_handoff(self, companies: list, wait: float = 0) -> bool:
        """取一次爱企查移交的名单，追加到本侧名单尾部。返回是否真的追加了新名字。

        wait=0 时只看一眼不等（循环内用，绝不能因为对面还在等验证就把本侧也阻住）；
        自己名单跑完时才传一个大于 0 的时长等对面落定，否则会漏掉刚移交过来的几家。
        """
        if self.handoff is None:
            return False
        try:
            if wait > 0:
                self.handoff.wait_for(wait, should_stop=lambda: self._stop_requested)
            new = self.handoff.pending()
        except Exception as e:
            self._log(f"  ⚠ 接收移交名单失败：{e}")
            return False
        if not new:
            return False
        # 去重两层：已经采过的公司名、以及本侧名单里已经排着的名字
        # （移交过来的本来就是爱企查未采的，但原名单重复出现时两边会撞车）
        with self.results_lock:
            known = {block_guard.norm_company(r.get("公司名")) for r in self.results}
        known |= {block_guard.norm_company(n) for n in companies}
        fresh = [n for n in new if block_guard.norm_company(n) not in known]
        if not fresh:
            return False
        companies.extend(fresh)
        self._log(f"🔁 接下爱企查移交的 {len(fresh)} 家（企知道共 {len(companies)} 家）："
                  + "、".join(n[:16] for n in fresh[:3]) + ("…" if len(fresh) > 3 else ""))
        return True

    def _iter_work(self, companies: list, start_index: int):
        """按下标走名单，但每步重读名单长度：爱企查移交的名字是运行时追加进来的。

        不能沿用 range(start, len(companies)) 在开头定死长度，否则移交过来那批永远没人采。
        自己名单跑完时也不能直接收工：爱企查可能正在等人工验证，它撑不住之后才会把
        剩余名单放过来，所以先等对面落定（或超时），落定后又有名字就继续采。
        """
        idx = start_index
        while idx < len(companies):
            self._take_handoff(companies)
            yield idx
            idx += 1
        # 自己这边空了才需要等对面落定；循环途中的停止由采集循环自己记日志
        while not self._stop_requested and self._take_handoff(companies, self.handoff_wait):
            while idx < len(companies):
                if self._stop_requested:
                    return
                yield idx
                idx += 1

    # ── 登录态确认 ──
    def check_login(self):
        """用浏览器 Cookie 判登录态：未登录时详情页只回电话，其余字段全空"""
        if self.driver is None:
            return "unknown", "浏览器还没打开"
        try:
            if "qizhidao.com" not in (self.driver.current_url or ""):
                self.driver.get(self.BASE_URL)
                time.sleep(2)
            names = [c.get("name") for c in (self.driver.get_cookies() or [])]
        except Exception as e:
            return "unknown", f"读取 Cookie 失败：{str(e)[:60]}"
        return block_guard.login_state(names, self.PLATFORM_LABEL)

    def _gate(self):
        """登录闸门（延迟创建，signals 可能在爬虫建好之后才接上）"""
        if self._login_gate_obj is None:
            self._login_gate_obj = block_guard.LoginGate(
                self.PLATFORM_LABEL, self._log,
                lambda r: self.signals.login_required_signal.emit(r) if self.signals else None,
                lambda: self._stop_requested, self.check_login,
                emit_cleared=self._emit_cleared,
            )
        return self._login_gate_obj

    def resolve_login_gate(self, choice: str):
        """界面把用户的选择（continue / login_done / stop）交回采集线程"""
        self._gate().resolve(choice)

    def _ask_login(self, reason: str) -> bool:
        """没有界面时（命令行直接跑）不等人，直接结束本轮，避免默默卡 30 分钟"""
        if self.signals is None:
            self._log(f"🛑 {reason}；没有界面可确认，本轮采集结束")
            return False
        return self._gate().ensure(reason)

    def _login_gate(self) -> bool:
        """采集开始前确认登录态。True=继续，False=本轮结束"""
        state, detail = self.check_login()
        if state == "ok":
            self._log(f"🔑 {detail}")
            return True
        if state != "missing":
            return True          # 浏览器没开、会话未就绪时不拦，交给循环里的会话检查
        return self._ask_login(detail)

    def _shell_gate(self) -> bool:
        """连续采到空壳：多半是登录失效或权限不足。不写空记录，走登录确认流程"""
        if self._shell_gate_passed:
            return True          # 本轮已经问过你了，再遇到空壳照写，不反复弹窗
        self._shell_gate_passed = True
        state, detail = self.check_login()
        head = f"连续 {self._shell_streak} 家详情字段全空"
        reason = (head + f"：{detail}" if state == "missing"
                  else head + "（登录凭证仍在，可能是平台权限或页面结构变了）")
        return self._ask_login(reason)

    # ── Cookie 持久化 ──
    def _save_cookies(self, phone: str = ""):
        try:
            cookies = self.driver.get_cookies()
            data = {
                "cookies": cookies,
                "saved_at": datetime.now().isoformat(),
                "phone": phone,
            }
            with open(self._cookie_file, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            self._log("Cookies 已保存")
        except Exception as e:
            self._log(f"Cookie 保存失败: {e}")

    def _load_cookies(self):
        if self._cookie_file.exists():
            try:
                with open(self._cookie_file, "r", encoding="utf-8") as f:
                    data = json.load(f)
                for c in data.get("cookies", []):
                    try:
                        self.driver.add_cookie(c)
                    except Exception:
                        pass
                self._log(f"已加载 Cookies（保存于 {data.get('saved_at', '未知')}）")
                return True
            except Exception:
                pass
        return False

    def _check_login_status(self) -> bool:
        """检测是否已登录：以登录之后才会下发的 Cookie 为准

        以前是“Cookie 名字里带 token/session/auth 就算已登录”，但企知道未登录
        时也会下发匿名的 token，于是未登录被判成已登录，整份名单会被采成只电话有值的空数据。
        只有在实时 Cookie 压根读不到（unknown）时，才退回看页面上的用户标识。
        """
        state, detail = self.check_login()
        if state == "ok":
            self._log(f"检测到登录凭证：{detail}")
            return True
        if state == "missing":
            self._log(f"未检测到登录凭证：{detail}")
            return False
        try:
            body = self.driver.find_element(By.TAG_NAME, 'body').text[:500]
            if re.search(r'1[3-9]\d{1}\*{4}\d{4}', body) or '普通用户' in body or 'VIP' in body:
                self._log("检测到用户标识（脱敏手机号/普通用户/VIP）")
                return True
        except Exception:
            pass
        return False

    # ── 会话健康检查 ──
    def _is_session_alive(self) -> bool:
        """检查浏览器会话是否存活"""
        try:
            if self.driver is None:
                return False
            # 尝试执行简单 JS
            self.driver.execute_script("return document.title")
            # 尝试获取 URL
            _ = self.driver.current_url
            return True
        except (WebDriverException, Exception):
            return False

    def try_restore_session(self) -> bool:
        """启动时的登录探测：不跟已在跑的浏览器抢配置目录

        调试 Chrome 用的 user-data-dir 同一时间只能给一个实例用。双标签的共享
        浏览器在跑时再 start_browser，新 Chrome 会直接退出，报出来的
        “Cookie 恢复失败: session not created”会让用户误以为账号出了什么问题。
        这时只按已保存的登录凭证给个状态提示；真正开采前还会用浏览器里的
        实时 Cookie 再确认一次（见 _login_gate）。
        """
        if not self._cookie_file.exists():
            return False
        if self.driver is not None:
            state, detail = self.check_login()
            self._logged_in = state != "missing"
            self._log(f"浏览器已就绪：{detail or '未读到登录凭证，开采前会再确认'}")
            return self._logged_in
        if block_guard.debug_chrome_alive(self.attach_port or block_guard.SHARED_DEBUG_PORT):
            state, detail = block_guard.login_state(
                block_guard.saved_cookie_names(self._cookie_file), self.PLATFORM_LABEL)
            self._logged_in = state == "ok"
            if self._logged_in:
                self._log(f"浏览器正在运行，按上次保存的登录凭证判定：{detail}")
            else:
                self._log("浏览器正在运行，未另起 Chrome（避开配置目录占用）："
                          "上次保存的凭证里没有登录标识，开采前会用实时 Cookie 再确认一次")
            return self._logged_in
        try:
            self.start_browser(headed=True)
            self.driver.get(self.BASE_URL)
            time.sleep(2)
            self._load_cookies()
            self.driver.refresh()
            time.sleep(3)
            if self._check_login_status():
                self._logged_in = True
                self._log("从 Cookie 恢复登录成功")
                return True
            else:
                self._log("Cookie 已过期或从未登录，需重新登录")
                return False
        except Exception as e:
            self._log(f"Cookie 恢复失败: {e}")
            return False

    def _restart_browser(self, headed: bool = True):
        """重启浏览器"""
        self._log("正在重启浏览器...")
        try:
            if self.shared_session:
                # 共享模式下绝不能 quit（可能连带关掉另一个会话正在用的浏览器），只关自己的标签
                self._close_own_tab()
            elif self.driver:
                self.driver.quit()
        except Exception:
            pass
        self.driver = None
        self._own_handle = None
        time.sleep(2)
        self.start_browser(headed=headed, attach_port=self.attach_port)
        self._log("浏览器已重启")

    def _close_own_tab(self):
        """只关闭本会话独占的窗口/标签（共享浏览器保持运行）"""
        try:
            if self.driver is not None and self._own_handle:
                handles = self.driver.window_handles
                if self._own_handle in handles:
                    if self.driver.current_window_handle != self._own_handle:
                        self.driver.switch_to.window(self._own_handle)
                    self.driver.close()
        except Exception:
            pass

    def _ensure_own_tab(self) -> bool:
        """共享模式下确保当前操作对象仍是自己的窗口（防止误碰另一个会话的页面）"""
        if not self.shared_session or self.driver is None:
            return True
        try:
            if self.driver.current_window_handle == self._own_handle:
                return True
            handles = self.driver.window_handles
            if self._own_handle not in handles:
                self.driver.switch_to.new_window("window")
                self._own_handle = self.driver.current_window_handle
                self._log("  原窗口已被关闭，已重新开一个独占窗口")
                return True
            self.driver.switch_to.window(self._own_handle)
            return True
        except Exception:
            return False

    # ── 浏览器启动 / 关闭 ──
    def start_browser(self, headed: bool = True, attach_port: int = None):
        with self._browser_lock:
            if self.driver is not None:
                return
            if attach_port:
                # 双标签并行模式：附加到已在运行的共享 Chrome，并新开一个自己独占的标签
                self.attach_port = attach_port
                self.shared_session = True
                opts = Options()
                opts.add_experimental_option("debuggerAddress", f"127.0.0.1:{attach_port}")
                # 注意：附加模式（debuggerAddress）下 chromedriver 不接受 excludeSwitches /
                # useAutomationExtension，带上会报 "unrecognized chrome option"
                if sys.platform == "darwin":
                    chromedriver_path = Path(__file__).parent / "chromedriver-mac-x64" / "chromedriver"
                else:
                    chromedriver_path = Path(__file__).parent / "chromedriver-win64" / "chromedriver.exe"
                if chromedriver_path.exists():
                    self.driver = webdriver.Chrome(
                        service=Service(executable_path=str(chromedriver_path)), options=opts)
                else:
                    self.driver = webdriver.Chrome(options=opts)
                self.driver.switch_to.new_window("window")
                self._own_handle = self.driver.current_window_handle
                self.driver.execute_cdp_cmd(
                    "Page.addScriptToEvaluateOnNewDocument",
                    {"source": "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"}
                )
                self.wait = WebDriverWait(self.driver, 10)
                self._log(f"已附加到共享浏览器（端口 {attach_port}），独占窗口 {self._own_handle[:8]}")
                self._goto_homepage_verified()
                return True
            opts = Options()
            if not headed:
                opts.add_argument("--headless=new")
            opts.add_argument("--window-size=1400,900")
            opts.add_argument("--disable-blink-features=AutomationControlled")
            opts.add_argument("--disable-notifications")
            opts.add_argument("--no-first-run")
            opts.add_argument("--disable-infobars")
            opts.add_argument("--no-sandbox")
            opts.add_argument("--disable-dev-shm-usage")
            # 防止 Chrome 恢复上次会话/弹出"恢复页面"气泡（避免打开 hao123 等被劫持的启动页）
            opts.add_argument("--disable-session-crashed-bubble")
            opts.add_argument("--hide-crash-restore-bubble")
            opts.add_experimental_option("excludeSwitches", ["enable-automation"])
            opts.add_experimental_option("useAutomationExtension", False)

            # 使用 Chrome 配置文件保持登录状态
            if sys.platform == "darwin":
                temp_profile = Path("/tmp/chrome_debug_profile")
                original_profile = os.path.expanduser("~/Library/Application Support/Google/Chrome")
                if temp_profile.exists() and (temp_profile / "Default" / "Cookies").exists():
                    opts.add_argument(f"--user-data-dir={temp_profile}")
                    opts.add_argument("--profile-directory=Default")
                    self._log("使用 Chrome 临时配置（登录态）")
                elif os.path.exists(original_profile):
                    import shutil
                    if temp_profile.exists():
                        shutil.rmtree(temp_profile, ignore_errors=True)
                    self._log("正在复制 Chrome 配置...")
                    shutil.copytree(original_profile, str(temp_profile), dirs_exist_ok=True)
                    opts.add_argument(f"--user-data-dir={temp_profile}")
                    opts.add_argument("--profile-directory=Default")
                    self._log("使用 Chrome 配置（登录态）")
            else:
                # Windows: 复制 Chrome 配置到临时目录，避免单实例锁冲突
                chrome_profile = os.path.expandvars(r"%LOCALAPPDATA%\Google\Chrome\User Data")
                temp_profile = Path(os.environ.get("TEMP", r"C:\temp")) / "chrome_debug_profile"
                if os.path.exists(chrome_profile):
                    import shutil
                    if not temp_profile.exists():
                        self._log("正在复制 Chrome 配置...")
                        shutil.copytree(chrome_profile, str(temp_profile), dirs_exist_ok=True)
                    opts.add_argument(f"--user-data-dir={temp_profile}")
                    opts.add_argument("--profile-directory=Default")
                    self._log("使用 Chrome 配置（登录态）")

            # 使用本地 ChromeDriver（优先项目内，否则用系统 PATH）
            if sys.platform == "darwin":
                chromedriver_path = Path(__file__).parent / "chromedriver-mac-x64" / "chromedriver"
            else:
                chromedriver_path = Path(__file__).parent / "chromedriver-win64" / "chromedriver.exe"
            
            if chromedriver_path.exists():
                service = Service(executable_path=str(chromedriver_path))
                self.driver = webdriver.Chrome(service=service, options=opts)
            else:
                self.driver = webdriver.Chrome(options=opts)

            self.driver.execute_cdp_cmd(
                "Page.addScriptToEvaluateOnNewDocument",
                {"source": "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"}
            )
            self.wait = WebDriverWait(self.driver, 10)
            self._log("浏览器已启动")
            # 显式导航到平台首页并验证（解决配置被 hao123 劫持、多标签页等问题）
            self._goto_homepage_verified()
            return True

    def _goto_homepage_verified(self) -> bool:
        """关闭多余标签页并导航到平台首页，验证成功（解决 hao123 劫持/多标签问题）"""
        try:
            # 1. 只保留第一个标签页，关闭其余（hao123 等启动页可能在多余标签）
            #    共享浏览器下不能这么干：其余标签属于另一个并行会话，必须只操作自己的
            if self.shared_session:
                if not self._ensure_own_tab():
                    return False
            else:
                handles = self.driver.window_handles
                if len(handles) > 1:
                    for h in handles[1:]:
                        try:
                            self.driver.switch_to.window(h)
                            self.driver.close()
                        except Exception:
                            pass
                    try:
                        self.driver.switch_to.window(handles[0])
                    except Exception:
                        pass

            # 2. 导航并验证（最多 3 次）
            domain = self.BASE_URL.split("//")[-1].split("/")[0]
            for attempt in range(3):
                self.driver.get(self.BASE_URL)
                time.sleep(2)
                cur = self.driver.current_url or ""
                if domain in cur:
                    return True
                self._log(f"  首页导航未生效（第{attempt+1}次），重试...")
            self._log(f"  ⚠ 未能自动打开 {domain}，请手动在浏览器访问")
            return False
        except Exception as e:
            self._log(f"  首页导航异常: {e}")
            return False

    def close_browser(self):
        """关闭浏览器"""
        try:
            if self.driver:
                if self.shared_session:
                    # 共享浏览器：只关自己的标签，浏览器留给另一个会话继续用
                    self._close_own_tab()
                self.driver.quit()
        except Exception:
            pass
        self.driver = None
        self._own_handle = None
        self._log("浏览器已关闭")

    def pause_crawl(self):
        """暂停采集"""
        self._pause_requested = True

    def resume_crawl(self):
        """恢复采集"""
        self._pause_requested = False

    def stop_crawl(self):
        """停止采集"""
        self._stop_requested = True

    def login_manual(self) -> bool:
        """打开登录页让用户手动登录"""
        try:
            if self.driver is None:
                self.start_browser(headed=True)
            self.driver.get(self.BASE_URL)
            time.sleep(2)
            self._log("请在浏览器中手动完成登录，完成后点击'我已登录'按钮")
            if self.signals:
                self.signals.manual_check_signal.emit("请在浏览器中手动完成登录")
            return True
        except Exception as e:
            self._log(f"❌ 打开登录页失败：{e}")
            return False

    # ── 搜索公司（调用 suggestPredict API 获取公司hash → 直接跳详情页，零搜索次数） ──
    def search_company(self, name: str) -> dict:
        """搜索公司：调用 suggestPredict API → 获取公司hash → 直接跳详情页（零搜索次数消耗）"""
        result = {
            "公司名": name,
            "状态": "", "法定代表人": "", "注册资本": "",
            "成立日期": "", "统一社会信用代码": "",
            "电话": "", "邮箱": "", "官网": "", "地址": "",
            "数据来源": "企知道", "搜索结果URL": "",
        }

        for attempt in range(2):
            try:
                if self.driver is None:
                    self.start_browser(headed=True, attach_port=self.attach_port)

                # 会话健康检查
                if not self._is_session_alive():
                    self._log("  会话已断开，重启浏览器...")
                    self._restart_browser(headed=True)

                # 双标签并行：确保操作的是自己的标签
                self._ensure_own_tab()

                # 1. 确保在首页（获取cookies/session）
                if 'qizhidao.com' not in (self.driver.current_url or ''):
                    self.driver.get(self.BASE_URL)
                    time.sleep(2)

                # 2. 调用 suggestPredict API 获取公司hash（不消耗搜索次数）
                company_hash = self._get_company_hash(name)
                if not company_hash:
                    self._log(f"  未找到公司: {name}")
                    result["错误"] = "未找到公司（suggestPredict 无结果）"
                    return result

                # 3. 直接构造详情页URL并导航（零搜索次数）
                detail_url = f"https://qiye.qizhidao.com/company/{company_hash}.html"
                self._log(f"  直接跳详情: {detail_url}")
                result["搜索结果URL"] = detail_url
                self.driver.get(detail_url)
                time.sleep(3)

                # 4. 从 NUXT 数据 + DOM 提取全部字段
                result = self._extract_detail_data(result)
                return result

            except (WebDriverException, Exception) as e:
                err_str = str(e)
                if "invalid session" in err_str or "session deleted" in err_str or "disconnected" in err_str:
                    self._log(f"  会话异常，重启浏览器 (attempt {attempt+1}/2)...")
                    try:
                        self.driver.quit()
                    except Exception:
                        pass
                    self.driver = None
                    time.sleep(2)
                    if attempt == 1:
                        result["错误"] = f"会话异常: {err_str[:100]}"
                        return result
                else:
                    self._log(f"  ❌ 搜索失败: {err_str[:100]}")
                    result["错误"] = err_str[:200]
                    return result

        return result

    def _get_company_hash(self, name: str) -> str:
        """通过拦截页面自身的 suggestPredict 请求获取公司hash ID（不消耗搜索次数）"""
        try:
            # 每次都导航到首页（确保拦截器干净、session有效）
            self.driver.get(self.BASE_URL)
            time.sleep(2)

            # 注入 XHR 响应拦截器（必须在输入前注入）
            self.driver.execute_script("""
                window.__suggestPredictResult = null;
                
                var origXHROpen = XMLHttpRequest.prototype.open;
                var origXHRSend = XMLHttpRequest.prototype.send;
                
                XMLHttpRequest.prototype.open = function(method, url) {
                    this._url = url;
                    return origXHROpen.apply(this, arguments);
                };
                
                XMLHttpRequest.prototype.send = function() {
                    var self = this;
                    if (self._url && self._url.includes('suggestPredict')) {
                        self.addEventListener('load', function() {
                            try {
                                window.__suggestPredictResult = JSON.parse(self.responseText);
                            } catch(e) {}
                        });
                    }
                    return origXHRSend.apply(this, arguments);
                };
            """)

            # 在搜索框中输入公司名（触发 suggestPredict API）
            search_input = self.driver.find_element(By.XPATH, "//input[@type='text']")
            search_input.click()
            time.sleep(0.3)
            search_input.clear()
            
            # 逐字输入（模拟真实用户，触发自动补全API）
            for char in name:
                search_input.send_keys(char)
                time.sleep(0.08)
            
            # 等待 API 响应
            time.sleep(3)

            # 获取拦截到的 suggestPredict 结果
            api_result = self.driver.execute_script("return window.__suggestPredictResult;")

            if not api_result:
                self._log(f"  suggestPredict 未拦截到响应，尝试备用方案...")
                return self._get_hash_via_search(name)

            if not api_result.get("success"):
                self._log(f"  suggestPredict 失败: {api_result.get('msg', '')}")
                return self._get_hash_via_search(name)

            # 提取公司hash
            data = api_result.get("data", {})
            company_list = data.get("list", [])
            if company_list:
                # 优先精确匹配
                for item in company_list:
                    if item.get("name") == name:
                        return item.get("id", "")
                # 否则取第一个
                return company_list[0].get("id", "")

            self._log(f"  suggestPredict 无结果，尝试备用方案...")
            return self._get_hash_via_search(name)

        except Exception as e:
            self._log(f"  _get_company_hash 异常: {e}")
            return self._get_hash_via_search(name)

    def _get_hash_via_search(self, name: str) -> str:
        """备用方案：通过搜索页获取公司hash（会消耗搜索次数）"""
        try:
            search_url = f"{self.BASE_URL}/check?searchKey={urllib.parse.quote(name)}&tagNum=1"
            self.driver.get(search_url)
            time.sleep(3)

            result_links = self.driver.find_elements(
                By.XPATH, "//a[contains(@href,'/company/')]"
            )
            if result_links:
                href = result_links[0].get_attribute("href") or ""
                hash_match = re.search(r'/company/([a-f0-9]+)\.html', href)
                if hash_match:
                    self._log(f"  备用方案: 从搜索页获取hash")
                    return hash_match.group(1)

            return ""
        except Exception as e:
            self._log(f"  备用方案失败: {e}")
            return ""

    # ── 从详情页提取数据（NUXT + DOM 混合） ──
    def _extract_detail_data(self, result: dict) -> dict:
        """从详情页提取数据：NUXT参数精确提取 + DOM补充"""
        try:
            page_source = self.driver.page_source

            # ═══ 方式1: 从 NUXT IIFE 参数中精确提取 ═══
            nuxt_data = self._extract_nuxt_data(page_source)
            if nuxt_data:
                self._log("  ✅ NUXT 数据提取成功")
                if nuxt_data.get("operName"):
                    result["法定代表人"] = nuxt_data["operName"]
                if nuxt_data.get("creditCode"):
                    result["统一社会信用代码"] = nuxt_data["creditCode"]
                if nuxt_data.get("registCapi"):
                    result["注册资本"] = nuxt_data["registCapi"]
                if nuxt_data.get("startDate"):
                    result["成立日期"] = nuxt_data["startDate"]
                if nuxt_data.get("address"):
                    result["地址"] = nuxt_data["address"]
                if nuxt_data.get("telephones"):
                    # 分号分隔的电话列表
                    phones = [p.strip() for p in nuxt_data["telephones"].split(";") if p.strip()]
                    result["电话"] = "；".join(phones)
                if nuxt_data.get("emails"):
                    # 分号分隔的邮箱列表
                    emails = [e.strip() for e in nuxt_data["emails"].split(";") if e.strip()]
                    result["邮箱"] = "；".join(emails)
                if nuxt_data.get("scope"):
                    result["经营业务范围"] = nuxt_data["scope"]
                if nuxt_data.get("desc"):
                    result["简介"] = nuxt_data["desc"]
                if nuxt_data.get("industry"):
                    result["所属行业"] = nuxt_data["industry"]

            # ═══ 方式2: 从 DOM 补充/修正 ═══
            # 公司名（NUXT 中没有直接字段，从 DOM 取）
            try:
                name_el = self.driver.find_element(
                    By.XPATH, "//div[contains(@class,'company-name')]//a"
                )
                if name_el.text.strip():
                    result["公司名"] = name_el.text.strip()
            except NoSuchElementException:
                pass

            # 基本信息字段（DOM 方式补充 NUXT 缺失的字段）
            dom_fields = self._extract_dom_fields()
            # 只补充 NUXT 中没有的字段
            for key in ["企业类型", "营业期限", "核准日期", "登记机关",
                        "英文名", "实缴资本", "参保人数", "组织机构代码",
                        "工商注册号", "纳税人资质", "进出口企业代码",
                        "海关注册编码", "纳税人识别号"]:
                if key in dom_fields and key not in result:
                    result[key] = dom_fields[key]

            # 状态（从 DOM 取）
            if dom_fields.get("企业经营状态"):
                result["状态"] = dom_fields["企业经营状态"]
            elif dom_fields.get("经营状态"):
                result["状态"] = dom_fields["经营状态"]

            # 法定代表人（如果 NUXT 没取到，从 DOM 取并清理）
            if not result.get("法定代表人") and dom_fields.get("法定代表人"):
                raw = dom_fields["法定代表人"]
                # 清理：去掉"秦\n秦双玉\n1家关联企业" → "秦双玉"
                lines = [l.strip() for l in raw.split('\n') if l.strip()]
                # 过滤非人名行
                name_candidates = [
                    l for l in lines
                    if not any(kw in l for kw in ['关联', '企业>', '报告', '收藏', '发票', '监控', 'logo'])
                    and 1 < len(l) <= 10
                ]
                result["法定代表人"] = name_candidates[0] if name_candidates else lines[-1]

            # 官网（从 DOM 取）
            try:
                website = self.driver.execute_script("""
                    var items = document.querySelectorAll('.item');
                    for (var item of items) {
                        if (item.textContent.includes('官网')) {
                            var link = item.querySelector('a');
                            if (link && link.href && !link.href.includes('qizhidao')) return link.href;
                        }
                    }
                    return '';
                """)
                if website:
                    result["官网"] = website
            except Exception:
                pass

        except Exception as e:
            self._log(f"  ⚠ 提取详情页数据异常: {e}")

        return result

    def _extract_nuxt_data(self, page_source: str) -> dict:
        """从 NUXT IIFE 参数中提取关键字段"""
        try:
            nuxt_match = re.search(
                r'window\.__NUXT__=\(function\(([^)]+)\)\{return (.+?)\}\((.+?)\)\);',
                page_source, re.DOTALL
            )
            if not nuxt_match:
                return {}

            params = [p.strip() for p in nuxt_match.group(1).split(',')]
            func_body = nuxt_match.group(2)
            args_str = nuxt_match.group(3)

            # 解析参数值
            args = self._parse_nuxt_args(args_str)

            # 建立参数名→值的映射
            param_map = {}
            for i in range(min(len(params), len(args))):
                param_map[params[i]] = args[i]

            # 从 func_body 提取字段→参数名映射，再查值
            result = {}
            field_mapping = {
                'operName': '法定代表人',
                'creditCode': '统一社会信用代码',
                'registCapi': '注册资本',
                'startDate': '成立日期',
                'telephones': '电话',
                'emails': '邮箱',
                'address': '地址',
                'scope': '经营范围',
                'desc': '简介',
                'industry': '行业',
                'registStatus': '经营状态',
                'companyName': '公司名',
            }

            for js_field, cn_key in field_mapping.items():
                m = re.search(rf'{js_field}:(\w+)', func_body)
                if m:
                    var_name = m.group(1)
                    val = param_map.get(var_name, '')
                    if val:
                        result[cn_key] = val

            # 直接用原始字段名也存一份（方便调用方使用）
            for js_field in ['operName', 'creditCode', 'registCapi', 'startDate',
                             'telephones', 'emails', 'address', 'scope', 'desc', 'industry']:
                m = re.search(rf'{js_field}:(\w+)', func_body)
                if m:
                    var_name = m.group(1)
                    val = param_map.get(var_name, '')
                    if val:
                        result[js_field] = val

            return result

        except Exception as e:
            self._log(f"  NUXT 解析异常: {e}")
            return {}

    def _parse_nuxt_args(self, args_str: str) -> list:
        """解析 NUXT IIFE 的参数值列表"""
        args = []
        current = ''
        in_str = False
        str_char = ''
        depth = 0
        prev_char = ''

        for i, c in enumerate(args_str):
            if in_str:
                current += c
                if c == str_char and prev_char != '\\':
                    in_str = False
            elif c in ('"', "'"):
                in_str = True
                str_char = c
                current += c
            elif c in ('{', '['):
                depth += 1
                current += c
            elif c in ('}', ']'):
                depth -= 1
                current += c
            elif c == ',' and depth == 0:
                args.append(self._clean_arg(current.strip()))
                current = ''
            else:
                current += c
            prev_char = c

        if current.strip():
            args.append(self._clean_arg(current.strip()))

        return args

    def _clean_arg(self, val: str) -> str:
        """清理 NUXT 参数值（去引号、解码 unicode）"""
        if not val:
            return val
        if (val.startswith('"') and val.endswith('"')) or \
           (val.startswith("'") and val.endswith("'")):
            val = val[1:-1]
        # 解码 \u002F 等 unicode 转义
        val = val.replace('\\u002F', '/')
        val = val.replace('\\u0026', '&')
        val = val.replace('\\n', '\n')
        return val

    def _extract_dom_fields(self) -> dict:
        """从 DOM 提取基本信息字段（cell flex-center 结构）"""
        fields = {}
        try:
            cells = self.driver.find_elements(
                By.XPATH, "//div[contains(@class,'cell flex-center')]"
            )
            i = 0
            while i < len(cells) - 1:
                label = cells[i].text.strip()
                value = cells[i + 1].text.strip()
                if label and value:
                    # 清理 value（去掉"复制"后缀）
                    value = re.sub(r'\s*复制\s*$', '', value)
                    # 清理 label（去掉 tooltip 文字）
                    label = re.sub(r'\s*来自.*$', '', label)
                    label = re.sub(r'\s*趋势图\s*$', '', label)
                    fields[label] = value
                i += 2
        except Exception as e:
            self._log(f"  DOM 字段提取异常: {e}")
        return fields

    # ── 关闭弹窗 ──
    def _close_modal(self):
        """关闭弹窗"""
        try:
            close_btns = self.driver.find_elements(
                By.XPATH,
                "//div[contains(@class,'el-dialog')]//button[contains(@class,'el-dialog__close')] | "
                "//div[contains(@class,'el-dialog')]//i[contains(@class,'el-icon-close')]"
            )
            for btn in close_btns:
                if btn.is_displayed():
                    btn.click()
                    time.sleep(0.5)
                    return
        except Exception:
            pass
        # 兜底：ESC
        try:
            self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.ESCAPE)
            time.sleep(0.5)
        except Exception:
            pass

    # ── 批量采集 ──
    def crawl(self, companies: list, start_index: int = 0):
        """批量采集公司数据"""
        self._crawling = True
        self._stop_requested = False
        self._pause_requested = False
        total = len(companies)

        self._log(f"🚀 开始采集，共 {total} 家公司，从第 {start_index + 1} 条开始")

        # 未登录时详情页只会回一个电话，状态/法人/注册资本/邮箱/地址全是空的，
        # 继续跑就是把整份名单白烧掉，所以先确认登录态再开跑
        self._shell_streak = 0
        self._shell_gate_passed = False
        if not self._login_gate():
            self._crawling = False
            if self.signals:
                self.signals.crawl_finished_signal.emit()
            return

        consecutive_errors = 0

        for idx in self._iter_work(companies, start_index):
            if self._stop_requested:
                self._log(f"⏹ 用户停止采集，已采集 {len(self.results)} 条")
                break

            # 暂停检查
            while self._pause_requested:
                time.sleep(0.5)
                if self._stop_requested:
                    break

            name = str(companies[idx]).strip()
            if not name:
                continue

            # 被拦截（百度安全验证等）时原地等人工处理：不能把这一家记成失败并往下烧名单
            block = self._detect_block()
            if block and not self._wait_unblock(block):
                break

            self._log(f"[{idx + 1}/{len(companies)}] 正在查询：{name}")

            if self.signals:
                self.signals.progress_signal.emit(idx + 1, len(companies))

            # 会话健康检查
            if not self._is_session_alive():
                self._log("  会话已断开，重启浏览器...")
                self._restart_browser(headed=True)

            try:
                result = self.search_company(name)
                if block_guard.is_shell_record(result):
                    # 不写入空壳，也不推进名单：登录补上后从这一家重新采
                    self._shell_streak += 1
                    self._log(f"  ⚠ 详情字段全空（连续 {self._shell_streak} 家），这一家先不写入结果")
                    if self._shell_streak >= block_guard.SHELL_STREAK_LIMIT \
                            and not self._shell_gate():
                        self._log("⏹ 未登录问题未解决，本轮采集到此结束")
                        break
                    time.sleep(self.delay_between)
                    continue
                self._shell_streak = 0
                self._add_result(result)
                consecutive_errors = 0

                # 打印摘要
                phone_count = len(result.get("电话", "").split("；")) if result.get("电话") else 0
                email_count = len(result.get("邮箱", "").split("；")) if result.get("邮箱") else 0
                self._log(
                    f"  ✅ {result.get('公司名', '')[:20]} | "
                    f"状态:{result.get('状态', '')} | "
                    f"法人:{result.get('法定代表人', '')} | "
                    f"电话:{phone_count}个 | 邮箱:{email_count}个"
                )
            except Exception as e:
                consecutive_errors += 1
                self._log(f"  ❌ 查询失败: {e}")
                self._add_result({
                    "公司名": name, "错误": str(e), "数据来源": "企知道"
                })

                # 连续失败多次，重启浏览器
                if consecutive_errors >= 5:
                    self._log("  连续失败5次，重启浏览器...")
                    self._restart_browser(headed=True)
                    consecutive_errors = 0

            time.sleep(self.delay_between)

        self._crawling = False
        self._log(f"✅ 采集完成！共 {len(self.results)} 条")

        if self.signals:
            self.signals.crawl_finished_signal.emit()

    # ── 导出 Excel ──
    def export_excel(self, output_path: str):
        """导出结果到 Excel"""
        if not self.results:
            self._log("没有数据可导出")
            return False

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

            wb = Workbook()
            ws = wb.active
            ws.title = "企业数据"

            headers = list(self.results[0].keys())
            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"),
            )

            # Column 1: 序号表头
            cell = ws.cell(row=1, column=1, value="序号")
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border

            # Column 2+: 数据表头
            for col, h in enumerate(headers, 2):
                cell = ws.cell(row=1, column=col, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_align
                cell.border = thin_border

            # 数据行
            data_align = Alignment(vertical="center", wrap_text=True)
            for idx, row_data in enumerate(self.results):
                row_num = idx + 2
                # 序号
                cell = ws.cell(row=row_num, column=1, value=idx + 1)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border
                # 数据列
                for col, h in enumerate(headers, 2):
                    val = row_data.get(h, "")
                    cell = ws.cell(row=row_num, column=col, value=val)
                    cell.alignment = data_align
                    cell.border = thin_border

            # 列宽
            ws.column_dimensions['A'].width = 6
            col_widths = {
                "公司名": 30, "状态": 8, "法定代表人": 12, "注册资本": 15,
                "成立日期": 12, "统一社会信用代码": 22, "电话": 40, "邮箱": 35,
                "官网": 25, "地址": 40, "数据来源": 10, "搜索结果URL": 30,
                "经营业务范围": 50, "简介": 50,
            }
            for col, h in enumerate(headers, 2):
                from openpyxl.utils import get_column_letter
                letter = get_column_letter(col)
                ws.column_dimensions[letter].width = col_widths.get(h, 15)

            wb.save(output_path)
            self._log(f"✅ 已导出到：{output_path}")
            return True
        except Exception as e:
            self._log(f"❌ 导出失败：{e}")
            return False
