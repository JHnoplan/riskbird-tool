"""断点续采进度台账（任务清单 + 条目状态机 + 幂等键）。

为什么不能按「已采条数」当续采起点：results 是纯内存，进程一崩就归零，下次
开始永远从第 1 条重采，烧掉的还是已经花出去的平台配额。所以把「这一份名单要采谁」
记成一份带状态的清单，落盘、按公司名幂等、可容损——这是工业界通行做法。

状态机（running 不落盘，所以被强杀也不会留下“永远在跑”的僵尸条目）：
    pending    还没采
    done       采到了有效记录
    no_result  站点确认查不到（永久，续采不再试，省配额）
    failed     验证码超时/网络错/额度用尽/登录失效等临时失败（续采会重试）

文件：<会话目录>/采集进度.json。写盘走「临时文件 + os.replace」原子替换：
一半被强杀最多丢掉最后一次更新，不会把整份进度写成半个坏文件。

名单指纹（文件名+大小+mtime+首行列数）用来判断“这次选的名单是不是上次那份”，
对不上就不弹续采，避免拿着 A 名单的进度去采 B 名单。
"""

import json
import os
import re
import threading
from datetime import datetime

try:
    import block_guard   # 公司名归一化必须与续采跳过判据同一套口径，不能各写一份
except Exception:        # 打包/测试环境缺依赖时退回本地等价实现
    def block_guard_norm(name) -> str:
        s = re.sub(r"\s+", "", str(name or "").strip())
        return s.replace("（", "(").replace("）", ")").upper()

    class _LocalBlockGuard:
        norm_company = staticmethod(block_guard_norm)

    block_guard = _LocalBlockGuard()

CHECKPOINT_FILE_NAME = "采集进度.json"
CHECKPOINT_VERSION = 1
FLUSH_EVERY = 5           # 每采完 5 条落一次盘（暂停/停止/完成/关窗另外强制落）

STATUS_PENDING = "pending"
STATUS_DONE = "done"
STATUS_NO_RESULT = "no_result"
STATUS_FAILED = "failed"
ALL_STATUS = (STATUS_PENDING, STATUS_DONE, STATUS_NO_RESULT, STATUS_FAILED)

# 站点确认查不到：再试一遍还是空，续采直接跳过（与 is_valid_record 的关键字表同口径）
_NO_RESULT_KEYWORDS = ("未找到", "不存在", "无结果", "未匹配", "没有找到", "查无", "无此企业")
# 临时失败：验证码超时、网络断、额度用尽、登录失效……续采值得再试一次
_FAILED_KEYWORDS = ("失败", "异常", "额度", "待重试", "超时", "验证码", "登录",
                    "网络", "中断", "连接", "会话", "受限", "频繁")


def norm_key(name) -> str:
    """幂等键：名单写法与页面返回名的空格/全半角差异不能算成两家公司"""
    return block_guard.norm_company(name)


def classify_record(record) -> str:
    """一条采集结果 → 条目状态（判据复用清洗侧的有效性判定）"""
    if not record:
        return STATUS_FAILED
    name = str(record.get("公司名") or record.get("企业名称") or "").strip()
    if not name:
        return STATUS_FAILED
    status = str(record.get("状态") or "")
    try:
        import customer_library_cleaner as _cleaner
        if _cleaner.is_valid_record(record):
            return STATUS_DONE
    except Exception:
        if not any(kw in status for kw in _NO_RESULT_KEYWORDS + _FAILED_KEYWORDS):
            return STATUS_DONE
    for kw in _NO_RESULT_KEYWORDS:
        if kw in status:
            return STATUS_NO_RESULT
    for kw in _FAILED_KEYWORDS:
        if kw in status:
            return STATUS_FAILED
    return STATUS_FAILED


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _first_row_cols(path: str) -> int:
    """首行列数（指纹的一部分）：读不了返回 -1，不因为一个附属字段把续采卡死"""
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
                return len(row or [])
        finally:
            wb.close()
    except Exception:
        return -1


def files_fingerprint(paths) -> list:
    """名单文件指纹：同名同大小同 mtime 同列数才认是同一份名单"""
    out = []
    for p in paths or []:
        try:
            st = os.stat(p)
            out.append({"name": os.path.basename(p), "size": int(st.st_size),
                        "mtime": int(st.st_mtime), "cols": _first_row_cols(p)})
        except Exception:
            out.append({"name": os.path.basename(str(p)), "size": -1,
                        "mtime": -1, "cols": -1})
    return sorted(out, key=lambda d: str(d.get("name") or ""))


def fingerprint_matches(saved, current) -> bool:
    """指纹比对：文件集合必须完全一致；附属字段读不出来（-1）时不作为不匹配的理由"""
    def norm(items):
        got = {}
        for it in items or []:
            if not isinstance(it, dict):
                continue
            got[str(it.get("name") or "")] = (int(it.get("size") or -1),
                                              int(it.get("mtime") or -1),
                                              int(it.get("cols") or -1))
        return got

    a, b = norm(saved), norm(current)
    if not a or not b or set(a) != set(b):
        return False
    for name, (size, mtime, cols) in b.items():
        old_size, old_mtime, old_cols = a[name]
        for old, new in ((old_size, size), (old_mtime, mtime), (old_cols, cols)):
            if old >= 0 and new >= 0 and old != new:
                return False
    return True


class CrawlCheckpoint:
    """一次采集会话的进度台账（内存状态机 + 可落盘快照）"""

    def __init__(self, session_id: str = "", folder=None, log=None, key_func=None):
        self.session_id = str(session_id or "")
        self.folder = str(folder) if folder else ""
        self.log = log or (lambda _m: None)
        self._key = key_func or norm_key
        self.files = []
        self.snapshot = []      # 原始名单顺序快照：双标签取奇偶、进度百分比都靠它
        self.items = {}         # 归一化公司名 → {"name","status","platform","updated_at"}
        self.updated_at = ""
        self._dirty = 0
        # 双标签并行时两个线程都会写状态、自动保存线程也会顺带落盘：
        # 不加锁会出现两个线程抢同一个 .tmp，一个 replace 走另一个就报 ENOENT。
        self._lock = threading.RLock()

    # ── 构造 / 读取 ────────────────────────────────
    @property
    def path(self) -> str:
        return os.path.join(self.folder, CHECKPOINT_FILE_NAME) if self.folder else ""

    @classmethod
    def create(cls, session_id: str, folder, companies: list,
               file_paths=None, log=None, key_func=None) -> "CrawlCheckpoint":
        cp = cls(session_id, folder, log=log, key_func=key_func)
        cp.files = files_fingerprint(file_paths or [])
        cp.set_snapshot(companies)
        cp.updated_at = _now()
        cp.flush(force=True)
        return cp

    @classmethod
    def open_folder(cls, folder, log=None, key_func=None):
        """读会话目录里的进度文件 → (进度对象或 None, 说明)

        读不出来（半截写/手工改坏）不抛异常：返回 None + 提示语，由界面决定按新任务开始。
        """
        path = os.path.join(str(folder), CHECKPOINT_FILE_NAME)
        if not os.path.exists(path):
            return None, "没有进度文件"
        try:
            with open(path, "r", encoding="utf-8") as fh:
                data = json.load(fh)
            if not isinstance(data, dict) or not data.get("snapshot"):
                raise ValueError("进度文件内容不完整")
            cp = cls(str(data.get("session_id") or os.path.basename(str(folder))),
                     folder, log=log, key_func=key_func)
            cp.load_from(data)
            return cp, ""
        except Exception as e:
            return None, (f"进度文件已损坏（{e}），本次按新任务开始；"
                          f"已采数据仍可从原数据 Excel 恢复")

    def load_from(self, data: dict) -> None:
        self.session_id = str(data.get("session_id") or self.session_id)
        self.files = data.get("files") or []
        self.snapshot = [str(n) for n in (data.get("snapshot") or [])]
        items = data.get("items") or {}
        self.items = {}
        for key, val in items.items():
            if isinstance(val, dict):
                status = str(val.get("status") or STATUS_PENDING)
                if status not in ALL_STATUS:
                    status = STATUS_PENDING       # 未知状态一律当没采过，宁可重采不丢数
                self.items[str(key)] = {
                    "name": str(val.get("name") or ""),
                    "status": status,
                    "platform": str(val.get("platform") or ""),
                    "updated_at": str(val.get("updated_at") or ""),
                }
            else:
                self.items[str(key)] = {"name": "", "status": STATUS_PENDING,
                                        "platform": "", "updated_at": ""}
        self.updated_at = str(data.get("updated_at") or "")
        self._dirty = 0

    def to_dict(self) -> dict:
        return {
            "version": CHECKPOINT_VERSION,
            "session_id": self.session_id,
            "files": self.files,
            "snapshot": self.snapshot,
            "items": self.items,
            "updated_at": self.updated_at,
        }

    # ── 名单 ──────────────────────────────────────
    def set_snapshot(self, companies: list) -> None:
        """原始名单顺序快照（去重保序），每一项初始为 pending"""
        seen, snap = set(), []
        for name in companies or []:
            text = str(name or "").strip()
            key = self._key(text)
            if not text or not key or key in seen:
                continue
            seen.add(key)
            snap.append(text)
            self.items.setdefault(key, {"name": text, "status": STATUS_PENDING,
                                        "platform": "", "updated_at": ""})
        self.snapshot = snap

    def matches_files(self, file_paths) -> bool:
        return fingerprint_matches(self.files, files_fingerprint(file_paths))

    # ── 状态机 ────────────────────────────────────
    def status_of(self, name) -> str:
        return str((self.items.get(self._key(name)) or {}).get("status") or STATUS_PENDING)

    def is_done(self, name) -> bool:
        return self.status_of(name) in (STATUS_DONE, STATUS_NO_RESULT)

    def mark(self, name, status: str, platform: str = "") -> bool:
        """更新一条的状态；每采满 5 条自动落盘。返回是否发生变更"""
        key = self._key(name)
        if not key or status not in ALL_STATUS:
            return False
        with self._lock:
            item = self.items.get(key) or {"name": str(name or "").strip(),
                                           "status": STATUS_PENDING,
                                           "platform": "", "updated_at": ""}
            old = item.get("status")
            # 已经采到过有效数据，后面一次临时失败不能把“到手的数据”说成没采
            if old == STATUS_DONE and status == STATUS_FAILED:
                return False
            changed = False
            if old != status:
                item["status"] = status
                item["updated_at"] = _now()
                changed = True
            if platform and item.get("platform") != platform:
                item["platform"] = platform
                changed = True
            if not changed:
                return False
            item["name"] = item.get("name") or str(name or "").strip()
            self.items[key] = item
            self._dirty += 1
            need_flush = self._dirty >= FLUSH_EVERY
        if need_flush:
            self.flush()
        return True

    def mark_record(self, record, platform: str = "") -> bool:
        """按采集结果落状态（done / no_result / failed 由 classify_record 判）"""
        if not record:
            return False
        name = str(record.get("公司名") or record.get("企业名称") or "").strip()
        return self.mark(name, classify_record(record), platform)

    def mark_records(self, records, platform: str = "") -> int:
        return sum(1 for r in records or [] if self.mark_record(r, platform))

    def rebuild_done(self, records, platform: str = "") -> int:
        """进度文件损坏时的兜底：按原数据 Excel 里已有的记录重建状态"""
        n = 0
        for record in records or []:
            name = str(record.get("公司名") or record.get("企业名称") or "").strip()
            if not name:
                continue
            key = self._key(name)
            item = self.items.setdefault(key, {"name": name, "status": STATUS_PENDING,
                                               "platform": "", "updated_at": ""})
            status = classify_record(record)
            if item.get("status") == STATUS_PENDING or status == STATUS_DONE:
                item["status"] = status
                item["updated_at"] = _now()
                item["platform"] = platform or str(record.get("数据来源") or "") or item.get("platform", "")
                n += 1
                self._dirty += 1
        if self._dirty:
            self.flush()
        return n

    # ── 待采集合 ──────────────────────────────────
    def pending(self, companies=None) -> list:
        """待采 = 名单 − done − no_result（failed 计入待采重试）"""
        out, seen = [], set()
        for name in (companies if companies is not None else self.snapshot):
            key = self._key(name)
            if not key or key in seen:
                continue
            seen.add(key)
            if self.status_of(name) in (STATUS_DONE, STATUS_NO_RESULT):
                continue
            out.append(name)
        return out

    def summary(self, companies=None) -> dict:
        """给界面/日志的口径：名单 M 家里各状态多少"""
        names = list(companies if companies is not None else self.snapshot)
        counts = {s: 0 for s in ALL_STATUS}
        for name in names:
            counts[self.status_of(name)] += 1
        counts["total"] = len(names)
        counts["todo"] = counts[STATUS_PENDING] + counts[STATUS_FAILED]
        return counts

    def text_summary(self, companies=None) -> str:
        c = self.summary(companies)
        return (f"本次待采 {c['todo']} 家（名单 {c['total']} 家 / 已完成 {c[STATUS_DONE]} 家 / "
                f"失败待重试 {c[STATUS_FAILED]} 家 / 无结果跳过 {c[STATUS_NO_RESULT]} 家）")

    # ── 落盘 ──────────────────────────────────────
    def flush(self, force: bool = False) -> bool:
        """原子写盘；没有进度文件路径或没有新变更（且 not force）时跳过"""
        if not self.path:
            return False
        with self._lock:
            if self._dirty == 0 and not force:
                return True
            self.updated_at = _now()
            snapshot = json.dumps(self.to_dict(), ensure_ascii=False, indent=2)
            tmp = f"{self.path}.tmp"
            try:
                os.makedirs(os.path.dirname(self.path), exist_ok=True)
                with open(tmp, "w", encoding="utf-8") as fh:
                    fh.write(snapshot)
                os.replace(tmp, self.path)
            except Exception as e:
                self.log(f"⚠ 采集进度写盘失败：{e}")
                try:
                    if os.path.exists(tmp):
                        os.remove(tmp)
                except OSError:
                    pass
                return False
            self._dirty = 0
            return True


# ── 断点发现 ──────────────────────────────────────
def inspect_folder(folder, companies, file_paths=None, log=None):
    """看某个会话目录能不能续采 → (可续采的进度对象或 None, 说明)

    给界面弹窗前用：进度文件读不出来或名单指纹对不上都不该劝用户“继续”。
    """
    cp, msg = CrawlCheckpoint.open_folder(folder, log=log)
    if cp is None:
        return None, msg
    if file_paths and not cp.matches_files(file_paths):
        return None, "本次选择的名单文件与该任务不是同一份，已按新任务处理"
    counts = cp.summary(companies)
    if not counts["todo"]:
        return None, "该任务的名单已全部采完，无需续采"
    return cp, ""


def load_records_from_raw_excel(folder):
    """会话目录里的原数据 Excel → 记录列表（损坏进度文件的兜底数据源）"""
    try:
        import task_registry as _tasks
    except Exception:
        return []
    path = os.path.join(str(folder), _tasks.RAW_FILE_NAME)
    try:
        return _tasks.read_records_from_xlsx(path)
    except Exception:
        return []
