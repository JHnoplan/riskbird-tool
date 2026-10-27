"""customer_library_cleaner 离线自测：模板对齐、电话/邮箱清洗、跨平台合并。"""
import customer_library_cleaner as C

TEMPLATE = "/Users/jhnoplan/Downloads/客户库数据导入模板.xlsx"

results = [
    # 风鸟：正常记录，电话串混合格式
    {"公司名": "深圳市腾讯计算机系统有限公司", "状态": "存续", "法定代表人": "马化腾",
     "注册资本": "6500万美元", "成立日期": "1998-11-11", "统一社会信用代码": "9144030010458574X",
     "电话": "13800138000；0755-88881111；13912345678", "邮箱": "service@Tencent.com; support@tencent.com",
     "官网": "www.tencent.com", "地址": "深圳市南山区高新科技园", "数据来源": "风鸟"},
    # 企知道：与风鸟同公司（跨平台重复，电话邮箱应并集）
    {"公司名": "深圳市腾讯计算机系统有限公司", "状态": "存续", "法定代表人": "",
     "电话": "010-65985888转802；13700137000", "邮箱": "hr@tencent.com",
     "官网": "", "地址": "", "数据来源": "企知道"},
    # 爱企查：无表头脏值 + 未找到应剔除
    {"公司名": " 新誉轨道交通科技有限公司 ", "状态": "未找到", "法定代表人": "",
     "电话": "", "邮箱": "", "数据来源": "爱企查"},
    # 失败剔除
    {"公司名": "某某有限公司", "状态": "失败: 会话异常", "数据来源": "风鸟"},
    # 空壳记录剔除（无公司名）
    {"公司名": "", "状态": "", "数据来源": "风鸟"},
    # 爱企查：全角括号 + 分号分隔多电话 + 脏邮箱
    {"公司名": "华为技术有限公司（研发）", "状态": "在业", "法定代表人": "任正非/孙亚芳",
     "电话": "18800001111;075528780808;12345;abc", "邮箱": "sales@huawei.com(官网) Service@HUAWEI.com _bad@x",
     "官网": "www.huawei.com", "地址": "深圳市龙岗区坂田华为基地", "数据来源": "爱企查"},
    # 无联系方式但信息有效的记录（保留，备注提示）
    {"公司名": "小而无名工作室", "状态": "开业", "法定代表人": "",
     "电话": "", "邮箱": "", "数据来源": "风鸟"},
]

print("── 电话拆分 ──")
print(C.split_phones("13800138000；0755-88881111；010-65985888转802；12345;abc"))
print(C.split_emails("sales@huawei.com(官网) _bad@x;Service@HUAWEI.com"))
mob, land = C.split_phones("400-800-1234；800 820 5566；075528780808")
assert land == ["400-800-1234", "800-820-5566", "0755-28780808"], land
assert mob == [], mob

rows = C.build_clean_rows(results, log=print)
print("\n── 清洗行 ──")
for r in rows:
    print(r)

assert len(rows) == 2, f"期望2条有联系方式的有效记录，实际{len(rows)}"
# 无任何联系方式的行不入库（上传就是脏数据），但原数据文件里仍然全量保留
assert not any("小而无名工作室" in r["企业名称"] for r in rows), \
    "手机号/固话/邮箱三空的行必须被入库出口剔除"
tencent = next(r for r in rows if "腾讯" in r["企业名称"])
assert tencent["客户名称"] == "马化腾" and tencent["岗位"] == "法定代表人"
assert set(tencent["数据来源"].split(",")) == {"风鸟", "企知道"}
assert "13800138000" in tencent["手机号"] and "13700137000" in tencent["手机号"]
assert "0755-88881111" in tencent["固话"] and "010-65985888转802" in tencent["固话"]
# 多值分隔符：手机号/固话/邮箱全部用英文逗号（以前手机号、固话用“/”，导入时会被当成一个号码）
assert tencent["手机号"].count(",") == 2 and "/" not in tencent["手机号"], tencent["手机号"]
assert tencent["固话"].count(",") == 1 and "/" not in tencent["固话"], tencent["固话"]
assert "," in tencent["邮箱"] and "/" not in tencent["邮箱"], tencent["邮箱"]
huawei = next(r for r in rows if "华为" in r["企业名称"])
assert huawei["客户名称"] == "任正非"
assert "18800001111" in huawei["手机号"] and "12345" not in huawei["固话"]
assert huawei["邮箱"] == "sales@huawei.com,service@huawei.com", \
    "手机号/固话/邮箱多值统一用英文逗号分隔"

out = "/tmp/清洗后客户库入库数据_test.xlsx"
tpl = C.locate_template() or TEMPLATE
print("\n模板定位:", tpl)
C.write_clean_excel(rows, out, template_path=tpl, log=print)

# 校验输出与模板表头逐列一致
import openpyxl
ws = openpyxl.load_workbook(out).active
head = [c.value for c in ws[1]]
tw = openpyxl.load_workbook(TEMPLATE).active
t_header = [c.value if c.value is not None else "" for c in next(tw.iter_rows(min_row=1, max_row=1))]
assert head == t_header, f"表头不一致\n{head}\n{t_header}"
r2 = [c.value for c in ws[2]]
print("\n输出行2:", r2)
# A-I 与 J-R 两块数据一致
assert r2[0] == r2[9] and r2[1] == r2[10] and r2[3] == r2[12] and r2[8] == r2[17]
print("\nALL ASSERTS PASSED")

# ── 风鸟新增状态与多联系方式（详情页"更多"能给出 12-16 个电话） ──
import re as _re
from customer_library_cleaner import is_valid_record, split_phones, split_emails

for st in ("访问频繁待重试", "额度用尽待重试", "未找到",
           "模糊查询额度用尽待重试", "模糊查询额度用尽待重试（登记状态:在营）", "在营"):
    valid = is_valid_record({"公司名": "X公司", "状态": st, "电话": "13800001111"})
    assert valid == (st == "在营"), (st, valid)

mb, ld = split_phones("13815037273；17368303772；0519-88776088；021-69926000；16970087955")
assert all(_re.fullmatch(r"1[3-9]\d{9}", m) for m in mb), mb
assert len(mb) == 3 and len(ld) == 2, (mb, ld)
assert len(split_emails("a@b.com；c@d.org；bad@@x")) == 2
print("\n✅ 风鸟新增状态剔除与多联系方式拆分断言通过")

# ── 电话只分类不删：多个号码全量保留，分组/国家码写法不再被洗没 ──
many = "；".join(f"1380000{i:04d}" for i in range(12))
mb12, _ = split_phones(many)
assert len(mb12) == 12, f"12 个手机号必须全留，实际 {len(mb12)} 个"
row12 = C.make_row({"公司名": "号码很多的公司", "电话": many, "状态": "存续"}, ["企知道"])
assert row12["手机号"].count(",") == 11, row12["手机号"]

# 空格/横线分组的手机号（139 0806 0502）以前会被分隔符切成三段全丢
assert split_phones("139 0806 0502")[0] == ["13908060502"], split_phones("139 0806 0502")
assert split_phones("139-0806-0503")[0] == ["13908060503"], split_phones("139-0806-0503")
# 带国家码的手机号脱码后归到手机号，单独的 +86 不会变成垃圾号码
assert split_phones("+8613800138000；0086-139-0013-9000；+86")[0] == ["13800138000", "13900139000"]
# 标准格式认不出的一律落到固话，不再静默丢弃（E.164 国际号码最长 15 位）
assert split_phones("+49-30-12345678")[1] == ["+49-30-12345678"], split_phones("+49-30-12345678")
# 多段横线的固话会被归一并成「区号-号码」，不是整段丢掉
assert split_phones("0755-8888-1111")[1] == ["0755-88881111"], split_phones("0755-8888-1111")
# 但位数不够的碎片（分机、序号）不能被当成号码洗进固话
assert split_phones("12345；abc；分机802；010") == ([], []), split_phones("12345；abc；分机802；010")
print("✅ 电话只分类不删（全量保留 + 分组/国家码 + 兜底进固话）断言通过")

# ── 分类边界：手机号优先识别，带区号/前导 0/分机都不再被固话正则切坏 ──
for _raw in ("13900139000转801", "13900139000-801", "013900139000",
             "075513900139000", "0 139 0013 9000", "13900139000转"):
    _mb, _ld = split_phones(_raw)
    assert _mb == ["13900139000"] and _ld == [], ("手机号被切成固话碎片", _raw, _mb, _ld)
# 固话+分机不能被上面的规则误伤（数字串里正好藏着 1[3-9] 开头的 11 位）
assert split_phones("0519-8389830转6049") == ([], ["0519-8389830转6049"]), \
    split_phones("0519-8389830转6049")
assert split_phones("010-65985888转802") == ([], ["010-65985888转802"])
# 「区号 空格 号码」以前会被分隔符拆开、区号整段丢掉
assert split_phones("0755 88881111") == ([], ["0755-88881111"]), split_phones("0755 88881111")
assert split_phones("020 12345678；0755 88881111")[1] == ["020-12345678", "0755-88881111"]
print("✅ 分类边界（手机号优先 + 区号空格连写）断言通过")

# ── 裸号补区号：只有同一行区号唯一才补 ──
_r1 = C.make_row({"公司名": "区号唯一的公司有限公司",
                  "电话": "0755-88881111；88990011；88990012转201", "状态": "存续"},
                 ["风鸟"], _st := {})
assert _r1["固话"] == "0755-88881111,0755-88990011,0755-88990012转201", _r1["固话"]
assert _st["area_filled"] == 2 and _st["area_kept"] == 0, _st
_r2 = C.make_row({"公司名": "跨城有限公司",
                  "电话": "0755-88881111；020-33334444；88990011", "状态": "存续"},
                 ["风鸟"], _st2 := {})
assert _r2["固话"] == "0755-88881111,020-33334444,88990011", _r2["固话"]
assert _st2["area_filled"] == 0 and _st2["area_kept"] == 1, _st2
_r3 = C.make_row({"公司名": "只有裸号的公司有限公司", "电话": "88990011", "状态": "存续"},
                 ["风鸟"], _st3 := {})
assert _r3["固话"] == "88990011" and _st3["area_filled"] == 0, (_r3["固话"], _st3)
print("✅ 裸号补区号（同行区号唯一才补，多区号不猜）断言通过")

# ── 入库出口剔除三空行 ──
assert C._has_contact({"手机号": "", "固话": "", "邮箱": "a@b.com"}) is True
assert C._has_contact({"手机号": "", "固话": "", "邮箱": ""}) is False
assert not any("工作室" in r["企业名称"] for r in rows), "三空行不能留在入库文件里"
print("✅ 无任何联系方式的行不入库断言通过")
