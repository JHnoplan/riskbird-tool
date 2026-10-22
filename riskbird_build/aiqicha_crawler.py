"""
aiqicha_crawler.py - 爱企查 (aiqicha.baidu.com) 企业数据爬虫模块
"""

import json
import os
import sys
import time
import re
import threading
import urllib.parse
from datetime import datetime
from pathlib import Path

from selenium import webdriver
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import NoSuchElementException, TimeoutException

try:
    from PyQt5.QtCore import QObject, pyqtSignal
    HAS_PYQT = True
except ImportError:
    HAS_PYQT = False


if HAS_PYQT:
    class CrawlerSignals(QObject):
        log_signal = pyqtSignal(str)
        progress_signal = pyqtSignal(int, int)
        crawl_finished_signal = pyqtSignal()
else:
    class CrawlerSignals:
        pass


class AiQiChaCrawler:
    BASE_URL = "https://aiqicha.baidu.com"
    SEARCH_URL = f"{BASE_URL}/s"

    def __init__(self, gui_callback=None):
        self.driver = None
        self.wait = None
        self._crawling = False
        self._stop_requested = False
        self._pause_requested = False
        self.results = []
        self.gui_callback = gui_callback
        self._browser_lock = threading.Lock()

        if HAS_PYQT:
            self.signals = CrawlerSignals()
        else:
            self.signals = None

        self.delay_between = 1.0

    def _log(self, msg: str):
        ts = datetime.now().strftime("%H:%M:%S")
        full_msg = f"[{ts}] {msg}"
        if self.gui_callback:
            self.gui_callback(full_msg)
        if self.signals:
            self.signals.log_signal.emit(full_msg)
        print(full_msg)

    def start_browser(self, headed: bool = True):
        with self._browser_lock:
            if self.driver is not None:
                return
            opts = Options()
            if not headed:
                opts.add_argument("--headless=new")
            opts.add_argument("--window-size=1400,900")
            opts.add_argument("--disable-blink-features=AutomationControlled")
            opts.add_argument("--disable-notifications")
            opts.add_experimental_option("excludeSwitches", ["enable-automation"])
            opts.add_experimental_option("useAutomationExtension", False)

            self.driver = webdriver.Chrome(options=opts)
            self.driver.execute_cdp_cmd(
                "Page.addScriptToEvaluateOnNewDocument",
                {"source": "Object.defineProperty(navigator,'webdriver',{get:()=>undefined})"},
            )
            self.wait = WebDriverWait(self.driver, 15)
            self._log("🌐 Chrome 浏览器已启动")

    def close_browser(self):
        if self.driver:
            try:
                self.driver.quit()
            except Exception:
                pass
            self.driver = None

    def _is_session_alive(self) -> bool:
        """检查浏览器会话是否存活"""
        try:
            if self.driver is None:
                return False
            self.driver.execute_script("return document.title")
            _ = self.driver.current_url
            return True
        except Exception:
            return False

    def search_company(self, name: str) -> dict:
        """搜索公司并返回第一条结果"""
        result = {
            "公司名": name,
            "法定代表人": "",
            "注册资本": "",
            "成立日期": "",
            "统一社会信用代码": "",
            "电话": "",
            "邮箱": "",
            "地址": "",
            "经营状态": "",
            "搜索结果 URL": "",
        }
        try:
            if self.driver is None:
                self.start_browser(headed=True)

            search_url = f"{self.SEARCH_URL}?q={urllib.parse.quote(name)}"
            self.driver.get(search_url)
            time.sleep(3)

            # 尝试多种选择器定位结果列表
            cards = []
            selectors = [
                "//div[contains(@class,'search-result')]//div[contains(@class,'item')]",
                "//div[contains(@class,'list-item')]",
                "//div[contains(@class,'company-item')]",
                "//div[contains(@class,'result-card')]",
                "//a[contains(@class,'title')]/ancestor::div[contains(@class,'item')]",
            ]
            for sel in selectors:
                try:
                    cards = self.driver.find_elements(By.XPATH, sel)
                    if cards:
                        self._log(f"  找到 {len(cards)} 条结果 (选择器：{sel[:30]}...)")
                        break
                except Exception:
                    continue

            if not cards:
                page_text = self.driver.find_element(By.TAG_NAME, "body").text[:500]
                self._log(f"  未找到结果，页面内容前 500 字：{page_text}")
                result["经营状态"] = "未找到"
                return result

            card = cards[0]
            card_text = card.text

            # 提取公司名
            try:
                name_el = card.find_element(By.XPATH, ".//a[contains(@class,'title')] | .//h3 | .//span[contains(@class,'name')]")
                result["公司名"] = name_el.text.strip()
                href = name_el.get_attribute("href") or ""
                if href:
                    result["搜索结果 URL"] = href
            except NoSuchElementException:
                pass

            # 提取各项信息
            patterns = {
                "法定代表人": r"法定代表人 [：:\s]*(\S+)",
                "注册资本": r"注册资本 [：:\s]*([\d.]+\s*万\S*)",
                "成立日期": r"成立日期 [：:\s]*(\d{4}[-/]\d{1,2}[-/]\d{1,2})",
                "统一社会信用代码": r"(?:统一社会信用代码 | 信用代码)[：:\s]*([A-Z0-9]{18})",
                "电话": r"电话 [：:\s]*([\d\-]+)",
                "邮箱": r"邮箱 [：:\s]*([\w.+-]+@[\w.-]+)",
            }
            for key, pattern in patterns.items():
                m = re.search(pattern, card_text)
                if m:
                    value = m.group(1).strip()
                    if key in ["电话", "邮箱"]:
                        value = re.sub(r'[更多\s]+$', '', value)
                    result[key] = value

            # 地址
            m = re.search(r"(?:地址 | 通信地址)[：:\s]*(.+?)(?:电话 | 邮箱 | 官网|$)", card_text, re.DOTALL)
            if m:
                result["地址"] = m.group(1).strip()[:100]

            # 经营状态
            for status in ["存续", "在业", "在营", "注销", "吊销"]:
                if status in card_text:
                    result["经营状态"] = status
                    break

            self._log(f"  ✅ 搜索到：{result['公司名']}")

        except TimeoutException:
            self._log(f"  搜索 '{name}' 超时")
            result["经营状态"] = "搜索超时"
        except Exception as e:
            self._log(f"  搜索 '{name}' 异常：{e}")
            result["经营状态"] = f"异常：{str(e)[:50]}"

        return result

    def test_search(self, company_name: str) -> dict:
        """测试搜索单个公司"""
        self._log(f"🔍 测试搜索：{company_name}")
        result = self.search_company(company_name)
        if result.get("经营状态") and result["经营状态"] not in ["未找到", "搜索超时"]:
            self._log(f"✅ 测试成功！采集到数据")
        else:
            self._log(f"  测试失败：{result.get('经营状态', '未知错误')}")
        return result

    def crawl(self, companies: list, start_index: int = 0):
        """批量采集"""
        self._crawling = True
        self._stop_requested = False
        self._pause_requested = False
        total = len(companies)

        self._log(f"🚀 开始采集，共 {total} 家公司，从第 {start_index + 1} 条开始")

        for idx in range(start_index, total):
            if self._stop_requested:
                self._log(f"  用户停止采集，已采集 {len(self.results)} 条")
                break

            while self._pause_requested:
                time.sleep(0.5)
                if self._stop_requested:
                    break

            name = str(companies[idx]).strip()
            if not name:
                continue

            self._log(f"[{idx + 1}/{total}] 正在查询：{name}")

            if self.signals:
                self.signals.progress_signal.emit(idx + 1, total)

            result = self.search_company(name)
            self.results.append(result)
            time.sleep(self.delay_between)

        self._crawling = False
        self._log(f"✅ 采集完成！共 {len(self.results)} 条数据")
        if self.signals:
            self.signals.crawl_finished_signal.emit()

    def stop_crawl(self):
        self._stop_requested = True

    def pause_crawl(self):
        self._pause_requested = True

    def resume_crawl(self):
        self._pause_requested = False

    def export_excel(self, filepath: str):
        """导出 Excel"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "企业数据"

            if not self.results:
                return

            headers = ["公司名", "法定代表人", "注册资本", "成立日期", "统一社会信用代码",
                       "电话", "邮箱", "地址", "经营状态", "搜索结果 URL"]
            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="1a3c6e", end_color="1a3c6e", fill_type="solid")

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for row_idx, result in enumerate(self.results, 2):
                for col_idx, key in enumerate(headers, 1):
                    value = result.get(key, "")
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            wb.save(filepath)
            self._log(f"✅ 已导出 {len(self.results)} 条数据到：{filepath}")

        except Exception as e:
            self._log(f"❌ 导出失败：{e}")
