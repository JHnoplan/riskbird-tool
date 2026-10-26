"""离线验证云展动力对接：接口判据、上传文件构造、任务索引、界面弹窗与导入闭环。

全程不联网、不启动浏览器（mock 掉 requests.request）：
① 接口3 check_phone：格式不对本地就拦；账号不存在抛后端文案；code=1 时透传 msg。
② 接口4 list_groups：自动翻页取全量、模糊搜索参数、国外 tab 带 foreignFlag。
③ 接口5 create_group：空名/超长本地拦住，成功必须拿到 id。
④ 上传文件：表头严格 9 列、整行空白不写、超长字段截断（后端是整条不入库）。
⑤ 任务索引：会话=任务、started_at 不被续采刷掉、按采集时间倒序、账号绑定持久化。
⑥ 界面：未绑定提示文案、分组必选、任务必选、新建分组校验、零分组引导。
⑦ 导入闭环：逐任务上传 + 超 5000 行自动分片 + 导入期间锁界面、走完解锁并弹汇总。
⑧ 历史会话补登记与桌面目录归一。
⑨ 采集完自动导入：未绑定/取消勾选要退回，采完只传本次会话并在日志提示。
⑩⑪ 「清洗数据」按钮：外部表智能匹配列名写出入库格式，坏文件说清原因且不留空文件。
"""
import json
import os
import shutil
import tempfile
import time
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PyQt5.QtWidgets import QApplication  # noqa: E402

app = QApplication([])

import customer_library_cleaner as C  # noqa: E402
import openpyxl  # noqa: E402
import riskbird_gui as G  # noqa: E402
import yzd_client as Y  # noqa: E402
from riskbird_gui import MainWindow, YzdGroupDialog  # noqa: E402

TMP = tempfile.mkdtemp(prefix="yzd_test_")
T = G._tasks
ORIG_TOOL_DIR = T.tool_dir               # ⑧ 要拿回真实现验 set_desktop_dir，不能拿到下面被换掉的
T.tool_dir = lambda desktop_dir=None: Path(TMP)   # 测试数据只写临时目录，不碰真桌面


# ══════════════ ① ~ ④ 客户端层（mock requests） ══════════════
class Resp:
    def __init__(self, payload=None, status=200, broken=False):
        self.status_code = status
        self.payload = payload
        self.broken = broken

    def json(self):
        if self.broken:
            raise ValueError("not json")
        return self.payload


class FakeHttp:
    """按调用顺序吐出预先准备的响应，并记录每次请求"""

    def __init__(self, queue):
        self.queue = list(queue)
        self.calls = []

    def __call__(self, method, url, **kwargs):
        # params/data/json 会被客户端逐页就地复用同一个 dict，
        # 不拷快照就只能看到最后一帧（current 永远是末页），测不到翻页参数
        snap = {}
        for k, v in kwargs.items():
            snap[k] = dict(v) if isinstance(v, dict) and k in ("params", "data", "json") else v
        self.calls.append({"method": method, "url": url, "kwargs": snap})
        if not self.queue:
            raise AssertionError("测试没准备足够的响应")
        return self.queue.pop(0)


def ok(data):
    return Resp({"code": 0, "msg": "success", "data": data})


def biz_fail(msg):
    return Resp({"code": 1, "msg": msg, "data": None})


def use_http(queue):
    fake = FakeHttp(queue)
    Y.requests.request = fake
    return fake


CANNOT = "请输入正确的手机号"
NOEXIST = "云展账号不存在"

# ① 接口3
fake = use_http([ok(True)])
assert Y.YzdClient().check_phone("15088131371") is True
assert "exist/15088131371" in fake.calls[0]["url"], "应打接口3 的 exist 路径"

for bad in ("", "1508813", "25088131371", "150881313711", "abc"):
    try:
        Y.YzdClient().check_phone(bad)
        raise AssertionError(f"格式不对却没拦住：{bad}")
    except Y.YzdError as e:
        assert str(e) == CANNOT, f"格式错的文案要对得上后端口径：{e}"

fake = use_http([ok(False)])
try:
    Y.YzdClient().check_phone("13800000001")
    raise AssertionError("账号不存在却没报错")
except Y.YzdError as e:
    assert str(e) == NOEXIST, e
assert len(fake.calls) == 1, "格式校验不该占网络往返（上面那批已验证）"

fake = use_http([biz_fail("请输入正确的手机号")])      # 后端 msg 原样透出
try:
    Y.YzdClient().check_phone("13900000000")
except Y.YzdError as e:
    assert str(e) == CANNOT, e
else:
    raise AssertionError("code=1 必须当失败")

fake = use_http([Resp(None, status=502)])              # HTTP 非 200
try:
    Y.YzdClient().check_phone("13900000000")
except Y.YzdError as e:
    assert "502" in str(e), e
else:
    raise AssertionError("HTTP 非 200 必须当失败")

fake = use_http([Resp(broken=True)])                   # 返回 HTML/网关页
try:
    Y.YzdClient().check_phone("13900000000")
except Y.YzdError as e:
    assert "JSON" in str(e), e
else:
    raise AssertionError("非 JSON 响应必须当失败")
print("① 接口3 校验手机号：格式/存在/code/msg 透传 ✓")

# ② 接口4：两页自动翻页
page1 = ok({"records": [{"id": "1", "groupName": "A组", "clientNum": 3}], "total": 4, "pages": 2, "current": 1})
page2 = ok({"records": [{"id": "2", "groupName": "B组", "clientNum": 0},
                        {"id": "3", "groupName": "C组"}, {"id": "4", "groupName": "D组"}],
            "total": 4, "pages": 2, "current": 2})
fake = use_http([page1, page2])
groups = Y.YzdClient("15088131371").list_groups(foreign_flag=1, group_name="组")
assert [g["groupName"] for g in groups] == ["A组", "B组", "C组", "D组"], groups
assert len(fake.calls) == 2, "两页应该翻两次"
q1 = fake.calls[0]["kwargs"]["params"]
assert q1["phone"] == "15088131371" and q1["foreignFlag"] == 1, q1
assert q1["groupName"] == "组" and q1["current"] == 1 and q1["size"] == 100, q1
assert fake.calls[1]["kwargs"]["params"]["current"] == 2
# 空分组账号（测试号现状）：records 为空要立刻停住，不能空转 20 页
fake = use_http([ok({"records": [], "total": 0, "pages": 0})])
assert Y.YzdClient("15088131371").list_groups() == []
assert len(fake.calls) == 1, "空结果不该继续翻页"
print("② 接口4 分组分页：自动取全量 + 空结果早停 ✓")

# ③ 接口5：新建分组
for bad, expect in (("", "不能为空"), ("   ", "不能为空"),
                    ("这个是超过十个字符的分组名", "10个字符"), ("十个字不行但十六个字符绝对超限", "10个字符")):
    assert not bad.strip() or len(bad.strip()) > Y.GROUP_NAME_MAX, f"测试用例本身要能触发校验：{bad!r}"
    try:
        Y.YzdClient("15088131371").create_group(bad)
        raise AssertionError(f"分组名校验没拦住：{bad!r}")
    except Y.YzdError as e:
        assert expect in str(e), f"{expect} → {e}"

fake = use_http([ok({"id": "99", "groupName": "新组1"})])
g = Y.YzdClient("15088131371").create_group("新组", foreign_flag=1)
assert g["id"] == "99"
call = fake.calls[0]
assert call["method"] == "POST" and call["url"].endswith("/api/out/member/sms/group/save")
assert call["kwargs"]["params"] == {"phone": "15088131371"}, call["kwargs"]["params"]
assert call["kwargs"]["json"] == {"groupName": "新组", "foreignFlag": 1}
assert call["kwargs"]["headers"]["Content-Type"] == "application/json"
fake = use_http([ok(None)])                    # 后端没给 id 不能当成功
try:
    Y.YzdClient("15088131371").create_group("另一组")
    raise AssertionError("缺 id 却当成功")
except Y.YzdError as e:
    assert "id" in str(e), e
print("③ 接口5 新建分组：本地校验 + 必须拿到 id ✓")

# ④ 上传文件构造
long_name = "甲" * 300
rows = [
    {"客户名称": "张三", "企业名称": "某某有限公司", "岗位": "法定代表人",
     "手机号": "13800138000,13900139000", "固话": "0755-12345678", "邮箱": "a@x.com,b@x.com",
     "微信号": "", "数据来源": "企知道", "备注": "备注内容", "外呼备注": "不该上传的列"},
    {"客户名称": long_name, "手机号": "1" * 130, "企业名称": "长名公司", "备注": "y" * 3000},
    {k: "" for k in Y.IMPORT_HEADERS},                     # 整行空白
    {"客户名称": None, "企业名称": "只有企业名"},           # None 要当空串
]
p = os.path.join(TMP, "upload.xlsx")
assert Y.write_import_workbook(rows, p) == 3, "空白行不该计入"
wb = openpyxl.load_workbook(p)
ws = wb.active
head = [c.value for c in ws[1]]
assert tuple(head) == Y.IMPORT_HEADERS, f"表头必须与接口文档模板完全一致：{head}"
assert "外呼备注" not in head, "非模板列不该写进上传文件"
body = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
assert body[0][0] == "张三" and body[0][7] == "企知道"
assert len(body[1][0]) == Y.FIELD_LIMITS["客户名称"], "超长按上限截断而不是整条丢"
assert len(body[1][3]) == Y.FIELD_LIMITS["手机号"]
assert len(body[1][8]) == Y.FIELD_LIMITS["备注"]
assert not any(isinstance(v, str) and v.strip() == "None" for v in body[2]), \
    f"None 要当空处理，不能写成字面量“None”：{body[2]}"
assert body[2][1] == "只有企业名", body[2]
wb.close()
assert Y.clamp_fields({"客户名称": None})["客户名称"] == ""
print("④ 上传文件：9 列表头 / 跳空行 / 超长截断 ✓")

# 接口6：multipart 字段与返回计数
fake = use_http([ok({"importTotal": 10, "saveTotal": 9, "failTotal": 1})])
res = Y.YzdClient("15088131371").import_excel(p, group_id="99", foreign_flag=0)
assert res == {"importTotal": 10, "saveTotal": 9, "failTotal": 1}, res
form = fake.calls[0]["kwargs"]["data"]
assert form["phone"] == "15088131371" and form["groupId"] == "99", form
assert form["foreignFlag"] == "0" and form["clientType"] == "1"
assert "file" in fake.calls[0]["kwargs"]["files"]
try:
    Y.YzdClient("15088131371").import_excel(os.path.join(TMP, "不存在.xlsx"))
    raise AssertionError("文件不存在却没报错")
except Y.YzdError as e:
    assert "不存在" in str(e), e
print("④ 接口6 multipart：phone/groupId/foreignFlag/file ✓")

# 批内去重（接口6 不做重复校验，完全一样的行不能传两遍）
dup_rows = [
    {"客户名称": "张三", "手机号": "13800138000,13900139000", "企业名称": "甲公司", "数据来源": "企知道"},
    {"客户名称": "张三", "手机号": "13900139000,13800138000", "企业名称": "甲公司",
     "数据来源": "爱企查"},                     # 只是顺序不同 → 同一条
    {"客户名称": "张三", "手机号": "13800138000", "企业名称": "甲公司", "数据来源": "风鸟"},
    {"客户名称": "李四", "邮箱": "a@x.com", "企业名称": "乙公司", "数据来源": "企知道"},
]
kept = Y.dedupe_rows(dup_rows)
assert len(kept) == 3, kept
assert kept[0]["数据来源"] == "企知道,爱企查", kept[0]
assert Y.row_identity(dup_rows[0]) == Y.row_identity(dup_rows[1]), "乱序联系方式要归一成同一个 key"
assert Y.row_identity(dup_rows[0]) != Y.row_identity(dup_rows[2]), "号码不一样就是两条客户"
assert Y.dedupe_rows([]) == [] and Y.dedupe_rows(None) == []
print("④ 批内去重：同名同联系方式只留一条 + 数据来源合并 ✓")


# ══════════════ ⑤ 任务索引与绑定账号 ══════════════
assert T.make_task_name(["/x/名单A.xlsx"]) == "名单A.xlsx"
assert T.make_task_name(["/x/名单A.xlsx", "/x/名单B.xlsx", "/x/名单C.xlsx"]) == "名单A.xlsx 等 3 个文件"
assert T.make_task_name([]) == "未命名采集任务"

t1 = T.new_task("2026-10-01_09-00-00", ["/x/名单A.xlsx"], os.path.join(TMP, "2026-10-01_09-00-00"), count=12)
t1["started_at"] = "2026-10-01 09:00:00"      # 真实使用时二者一致，测试里写死才能验证排序
assert t1["status"] == T.TASK_STATUS_RUNNING and t1["raw_file"].endswith(T.RAW_FILE_NAME)
T.upsert_task(t1)
t2 = T.new_task("2026-10-05_10-00-00", ["/x/名单B.xlsx", "/x/名单C.xlsx"],
                os.path.join(TMP, "2026-10-05_10-00-00"))
t2["started_at"] = "2026-10-05 10:00:00"
T.upsert_task(t2)
assert len(T.load_tasks()) == 2
old_started = t1["started_at"]
again = T.new_task(t1["id"], ["/x/名单A.xlsx"], t1["folder"], count=30)   # 续采重复登记
again["started_at"] = old_started                                          # 界面上会保留最早时间
T.upsert_task(again)
assert len(T.load_tasks()) == 2, "同一会话目录不能插出两条任务"
assert T.update_task(t1["id"], count=40) is True
assert T.update_task("不存在的时间戳", count=1) is False, "update_task 不能插空壳"
assert T.load_tasks()[0]["count"] == 40
T.finish_task(t1["id"], T.TASK_STATUS_DONE, 40)
done = next(t for t in T.load_tasks() if t["id"] == t1["id"])
assert done["status"] == T.TASK_STATUS_DONE and done["finished_at"]
order = [t["id"] for t in T.sorted_for_import(T.load_tasks())]
assert order[0] == "2026-10-05_10-00-00", f"导入列表要按采集时间倒序：{order}"
assert order[1] == "2026-10-01_09-00-00", f"导入列表要按采集时间倒序：{order}"
assert T.latest_task()["id"] == "2026-10-05_10-00-00"
assert json.loads(open(T.index_path(), encoding="utf-8").read())[0]["id"] == t1["id"]

assert T.load_account() == ""
T.save_account("15088131371")
assert T.load_account() == "15088131371"
T.clear_account()
assert T.load_account() == ""

# 原数据回读：写出 → 读回，字段与条数都不能少（openpyxl 不用 read_only，避免行数被截断）
records = [
    {"公司名": "深圳某某科技有限公司", "状态": "在营", "法定代表人": "张三",
     "注册资本": "500万", "地址": "深圳市南山区", "电话": "13800138000", "邮箱": "a@x.com",
     "数据来源": "企知道"},
    {"公司名": "广州某某贸易有限公司", "状态": "在营", "法定代表人": "",
     "注册资本": "", "地址": "", "电话": "020-88888888", "邮箱": "", "数据来源": "爱企查"},
    {"公司名": "空行公司", "状态": "未找到", "数据来源": "企知道"},
]
raw_path = os.path.join(TMP, T.RAW_FILE_NAME)
G.write_merged_excel(records, raw_path, log=lambda _m: None)
back = T.read_records_from_xlsx(raw_path)
assert len(back) == 3, f"回读条数不能少：{len(back)}"
assert back[0]["公司名"] == "深圳某某科技有限公司" and back[0]["法定代表人"] == "张三"
assert back[1]["电话"] == "020-88888888"
try:
    T.read_records_from_xlsx(os.path.join(TMP, "没有这个文件.xlsx"))
    raise AssertionError("文件不存在却没报错")
except FileNotFoundError:
    pass
try:
    T.read_records_from_xlsx("")              # 任务没填 raw_file 也要能拦住
    raise AssertionError("空路径却没报错")
except FileNotFoundError:
    pass
print("⑤ 任务索引 + 绑定账号 + 原数据回读 ✓")


# ══════════════ ⑥ 界面：绑定校验 + 弹窗必选 ══════════════
class FakeBox:
    """记录弹出的提示，并回答输入类弹窗"""
    Warning = Information = Question = 0
    Yes = No = 0
    AcceptRole = 0
    answers = []
    shown = []

    @classmethod
    def information(cls, parent, title, text, *a):
        cls.shown.append(("info", title, text))

    @classmethod
    def question(cls, parent, title, text, *a):
        cls.shown.append(("ask", title, text))
        return cls.Yes           # 覆盖确认统一选“是”，让清洗能写同一个文件

    @classmethod
    def warning(cls, parent, title, text, *a):
        cls.shown.append(("warn", title, text))

    @classmethod
    def critical(cls, parent, title, text, *a):
        cls.shown.append(("crit", title, text))


class FakeGroupsWorker:
    """替掉真线程：start() 当场把分组吐回来（分组按 flag 预置）"""

    data = {0: [], 1: []}

    def __init__(self, phone, flag, keyword=""):
        self.phone, self.flag, self.keyword = phone, flag, keyword
        self.finished_groups = self

    def connect(self, fn):
        self._fn = fn

    def isRunning(self):
        return False

    def wait(self, *_a):
        return True

    def start(self):
        groups = [g for g in self.data.get(self.flag, []) if self.keyword in g["groupName"]]
        self._fn(self.flag, groups, "")


class FakeCreateWorker:
    """新建分组同时把分组追到假分组数据里（后端就是这样一个同账号共享列表）"""

    def __init__(self, phone, name, flag):
        self.name, self.flag = name, flag
        self.finished_create = self

    def connect(self, fn):
        self._fn = fn

    def isRunning(self):
        return False

    def wait(self, *_a):
        return True

    def start(self):
        group = {"id": f"new-{self.name}", "groupName": self.name, "foreignFlag": self.flag}
        FakeGroupsWorker.data.setdefault(self.flag, []).append(group)
        self._fn(group, "")


class FakeInput:
    answers = []

    @staticmethod
    def getText(parent, title, label, *a):
        return FakeInput.answers.pop(0) if FakeInput.answers else ("", False)


G.QMessageBox = FakeBox
G.QInputDialog = FakeInput
G.YzdGroupsWorker = FakeGroupsWorker
G.YzdCreateGroupWorker = FakeCreateWorker

w = MainWindow()
assert w.yzd_status_lbl.text() == "未绑定云展动力账号", w.yzd_status_lbl.text()
assert w.import_yzd_btn.isEnabled()

FakeBox.shown.clear()
w._yzd_phone = ""
w._open_yzd_import()      # 未绑定直接点导入
assert FakeBox.shown and FakeBox.shown[-1][2] == "请先绑定云展动力账号", FakeBox.shown

# 绑定：格式错当场拦（不发请求），账号不存在用后端文案
FakeBox.shown.clear()
FakeInput.answers = [("12345678901", True)]
w._bind_yzd_account()
assert FakeBox.shown[-1][2] == CANNOT, FakeBox.shown
assert w._yzd_phone == ""

use_http([ok(False)])
FakeBox.shown.clear()
FakeInput.answers = [("15000000000", True)]
w._bind_yzd_account()
w._yzd_worker.wait(3000)
app.processEvents()
assert FakeBox.shown[-1][2] == NOEXIST, FakeBox.shown
assert w._yzd_phone == "" and w.yzd_status_lbl.text() == "未绑定云展动力账号"

use_http([ok(True)])
FakeBox.shown.clear()
FakeInput.answers = [("15088131371", True)]
w._bind_yzd_account()
w._yzd_worker.wait(3000)
app.processEvents()
assert w._yzd_phone == "15088131371", w._yzd_phone
assert T.load_account() == "15088131371", "绑定态要落盘"
assert w.yzd_status_lbl.text() == "已绑定云展账号：15088131371", w.yzd_status_lbl.text()
assert FakeBox.shown[-1][1] == "绑定成功"
print("⑥ 界面绑定：未绑定/格式错/账号不存在/成功持久化 ✓")

# 分组弹窗：零分组引导、单选必选、新建分组、tab 切换
FakeGroupsWorker.data = {0: [{"id": "11", "groupName": "光伏组", "clientNum": 5, "foreignFlag": 0}],
                         1: [{"id": "22", "groupName": "海外组", "clientNum": 0, "foreignFlag": 1}]}
dlg = YzdGroupDialog("15088131371")
app.processEvents()
assert dlg.group_list.count() == 1, dlg.group_list.count()
assert "光伏组" in dlg.group_list.item(0).text() and "客户数 5" in dlg.group_list.item(0).text()

FakeBox.shown.clear()
dlg._accept_latest()                       # 没选分组：必选拦住，不能带着空分组 accept
assert dlg.group_list.selectedItems() == []
assert FakeBox.shown[-1][2] == "请选择一个分组", FakeBox.shown
assert dlg.mode == "" and dlg.chosen_group is None
assert "必选" in dlg.hint_lbl.text() and "c62828" in dlg.hint_lbl.styleSheet()

dlg.group_list.setCurrentRow(0)
dlg._show_tasks()                          # 选了分组才能进「导入其他」页
assert dlg.stack.currentIndex() == 1, "该切到「导入其他」页"
assert dlg.task_list.count() == 2, dlg.task_list.count()
first = dlg.task_list.item(0).text()
assert "名单B.xlsx 等 2 个文件" in first, f"任务名要取采集文档名：{first}"
assert "2026-10-05" in first, f"列表要带采集时间：{first}"

FakeBox.shown.clear()
dlg._accept_others()                       # 一个都没勾
assert FakeBox.shown[-1][2] == "请至少选择一个任务", FakeBox.shown
dlg.all_chk.setChecked(True)
assert all(dlg.task_list.item(i).checkState() == 2 for i in range(dlg.task_list.count()))
assert dlg.task_ok_btn.text().startswith("确认导入（2）"), dlg.task_ok_btn.text()
dlg._accept_others()
assert dlg.mode == dlg.MODE_OTHERS and len(dlg.chosen_tasks) == 2
assert dlg.chosen_group["id"] == "11"

# 「刷新」：客户数取后端返回值，导完想知道最新数量必须能绕过本地缓存重查
FakeGroupsWorker.data[0] = [{"id": "11", "groupName": "光伏组", "clientNum": 175, "foreignFlag": 0}]
dlg._groups[0] = [{"id": "11", "groupName": "光伏组", "clientNum": 5, "foreignFlag": 0}]
dlg._refresh_groups()
app.processEvents()
assert "客户数 175" in dlg.group_list.item(0).text(), dlg.group_list.item(0).text()
assert "60 秒" in dlg.group_status_lbl.toolTip(), dlg.group_status_lbl.toolTip()
assert dlg.refresh_btn.isEnabled(), "刷新完要能再点（不锁住按钮会让人以为没反应）"

# 国外 tab + 新建分组
dlg2 = YzdGroupDialog("15088131371")
dlg2.tab_btns[1].click()
app.processEvents()
assert "海外组" in dlg2.group_list.item(0).text(), dlg2.group_list.item(0).text()
FakeInput.answers = [("名字超过十个字符上限的分组", True)]
FakeBox.shown.clear()
dlg2._create_group()
assert FakeBox.shown[-1][2] == "群组名称最多输入10个字符！", FakeBox.shown
FakeInput.answers = [("新组", True)]
dlg2._create_group()
app.processEvents()
assert any("新组" in dlg2.group_list.item(i).text() for i in range(dlg2.group_list.count()))
assert dlg2._current_group()["id"] == "new-新组", "新建完应默认选中新分组"

# 零分组账号：给出可操作的引导，而不是一片空白
FakeGroupsWorker.data = {0: [], 1: []}
dlg3 = YzdGroupDialog("15088131371")
app.processEvents()
assert dlg3.group_list.count() == 0
assert "新建分组" in dlg3.group_status_lbl.text(), dlg3.group_status_lbl.text()
print("⑥ 分组/任务弹窗：必选校验 + 新建分组 + 零分组引导 ✓")


# ══════════════ ⑦ 导入闭环（逐任务 + 分片 + 锁界面） ══════════════
G.YZD_IMPORT_CHUNK = 2          # 把分片阈值调小，才能在不造 5000 行的情况下验证多次调用
uploads = []


class FakeImportHttp:
    """接口6 的假后端：把上传文件里的行读出来说明“入库”，并校验表头"""

    def __call__(self, method, url, **kwargs):
        assert url.endswith("/api/out/member/sms/client/info/import"), url
        form = kwargs["data"]
        assert form["groupId"] == "11" and form["phone"] == "15088131371", form
        file_field = kwargs["files"]["file"]
        payload = file_field[1] if isinstance(file_field, tuple) else file_field
        data = payload.read() if hasattr(payload, "read") else open(payload, "rb").read()
        path = os.path.join(TMP, "uploaded.xlsx")
        with open(path, "wb") as dst:
            dst.write(data)
        wb = openpyxl.load_workbook(path)
        head = [c.value for c in wb.active[1]]
        assert tuple(head) == Y.IMPORT_HEADERS, head
        n = max(0, wb.active.max_row - 1)
        wb.close()
        uploads.append(n)
        return ok({"importTotal": n, "saveTotal": n, "failTotal": 0})


Y.requests.request = FakeImportHttp()

w._yzd_import_worker = None
session_dir = os.path.join(TMP, "2026-10-26_08-00-00")
os.makedirs(session_dir, exist_ok=True)
w._session_folder = Path(session_dir)
w._excel_paths = ["/x/最新名单.xlsx"]
w._register_task(w._session_folder)
assert w._task_id == "2026-10-26_08-00-00"
fake_records = [{"公司名": f"公司{i}", "状态": "在营", "法定代表人": f"法人{i}",
                 "电话": f"1380000{i:04d}", "邮箱": f"c{i}@x.com", "数据来源": "企知道"}
                for i in range(1, 6)]
# 两家不同公司但法人同名同联系方式 → 清洗后是两行，批内去重该只传一条
fake_records += [
    {"公司名": "甲公司", "状态": "在营", "法定代表人": "张三", "电话": "13811112222",
     "邮箱": "zhang@x.com", "数据来源": "企知道"},
    {"公司名": "乙公司", "状态": "在营", "法定代表人": "张三", "电话": "13811112222",
     "邮箱": "zhang@x.com", "数据来源": "爱企查"},
]
w.qizhidao_crawler.results = list(fake_records)

FakeBox.shown.clear()
w._start_yzd_import({"id": "11", "groupName": "光伏组", "foreignFlag": 0}, None)
assert w._yzd_importing is True, "导入进行中要置位（closeEvent 靠它拦人）"
assert w.isEnabled() is False, "导入期间主界面必须锁住"
assert w._yzd_progress is not None and w._yzd_progress.isVisible(), "要有加载动画窗口"
assert w._yzd_progress.bar.minimum() == 0 and w._yzd_progress.bar.maximum() == 0, "不确定进度条才会跑动"
prog = w._yzd_progress
prog.close()                      # 没导完不给关（否则用户以为导入停了）
assert prog.isVisible() is True, "导入中的加载窗口必须还在"

deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
w._yzd_import_worker.wait(3000) if w._yzd_import_worker else None
app.processEvents()

assert w._yzd_importing is False, "导入结束必须解锁"
assert w.isEnabled() is True
assert w._yzd_progress is None, "加载窗口要收掉"
assert prog.isVisible() is False, "close() 被忽略过但导完必须真关掉"
assert uploads and sum(uploads) == 6, f"7 条清洗后去重成 6 条，应分片传完：{uploads}"
assert uploads.count(0) == 0
title_text = FakeBox.shown[-1][2]
assert "共入库 6 条客户" in title_text, FakeBox.shown
assert "批内去重" in w.log_text.toPlainText(), "去重动作要在日志里看得见"
# 逐任务口径：一个任务一次汇总，两个任务时列表里两行
task = next(t for t in T.load_tasks() if t["id"] == w._task_id)
assert task["yzd_group"] == "光伏组" and task["yzd_saved"] == 6, task
assert task["yzd_import_at"], task

# 历史任务走“回读 xlsx”这条路：当前会话没有内存数据时读文件
# （会话目录里只有一份包含“未找到”公司的 3 行原数据，清洗后应只剩 2 行）
w.qizhidao_crawler.results = []
old_task = next(t for t in T.load_tasks() if t["id"] == "2026-10-05_10-00-00")
old_task["raw_file"] = raw_path
T.upsert_task(old_task)
uploads.clear()
FakeBox.shown.clear()
w._start_yzd_import({"id": "11", "groupName": "光伏组", "foreignFlag": 0}, [old_task])
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
assert uploads and sum(uploads) == 2, f"历史任务应回读原数据文件后上传：{uploads}"
assert "名单B" in FakeBox.shown[-1][2], FakeBox.shown[-1][2]

# 后端业务失败（code=1）不能当成功，且要走完解锁
class FailHttp:
    def __call__(self, method, url, **kwargs):
        return biz_fail("导入的文件格式不正确")


Y.requests.request = FailHttp()
uploads.clear()
FakeBox.shown.clear()
w._start_yzd_import({"id": "11", "groupName": "光伏组", "foreignFlag": 0}, [old_task])
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
assert w._yzd_importing is False and w.isEnabled() is True, "后端失败同样要解锁"
assert FakeBox.shown[-1][0] in ("warn", "crit"), FakeBox.shown
assert "导入的文件格式不正确" in FakeBox.shown[-1][2], FakeBox.shown[-1][2]

# 空任务：读不到数据不该发起任何上传
Y.requests.request = FakeImportHttp()
empty_task = {"id": "空任务", "name": "没采到东西.xlsx", "raw_file": os.path.join(TMP, "不存在.xlsx")}
uploads.clear()
FakeBox.shown.clear()
w._start_yzd_import({"id": "11", "groupName": "光伏组", "foreignFlag": 0}, [empty_task])
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
assert uploads == [], f"读不到数据不该上传：{uploads}"
assert "读取数据失败" in FakeBox.shown[-1][2], FakeBox.shown[-1][2]
assert w._yzd_importing is False
print("⑦ 导入闭环：逐任务/分片/锁定解锁/后端失败透传 ✓")

# ══════════════ ⑧ 历史会话补登记 + 桌面目录归一 ══════════════
REAL_TOOL_DIR = ORIG_TOOL_DIR
hist_dir = os.path.join(TMP, "2026-09-11_08-30-00")
os.makedirs(hist_dir, exist_ok=True)
G.write_merged_excel(fake_records[:3], os.path.join(hist_dir, T.RAW_FILE_NAME), log=lambda _m: None)
os.makedirs(os.path.join(TMP, "空目录没数据"), exist_ok=True)          # 不该被补登
os.makedirs(os.path.join(TMP, "随手建的文件夹"), exist_ok=True)          # 没有原数据文件
os.makedirs(os.path.join(TMP, ".隐藏目录"), exist_ok=True)
G.write_merged_excel(fake_records[:1], os.path.join(TMP, ".隐藏目录", T.RAW_FILE_NAME),
                     log=lambda _m: None)

T.tool_dir = lambda desktop_dir=None: Path(TMP)
assert T.backfill_tasks() == 1, "只补有原数据的历史会话"
assert T.backfill_tasks() == 0, "重复调不能把同一个目录再补一条"
tasks = T.load_tasks()
assert len(tasks) == 4, [t["id"] for t in tasks]   # 原有 3 个 + 补登 1 个
hist = next(t for t in tasks if t["id"] == "2026-09-11_08-30-00")
assert hist["name"] == "历史采集 2026-09-11 08:30:00" and hist["count"] == 3, hist
assert hist["status"] == T.TASK_STATUS_DONE and hist["raw_file"].endswith(T.RAW_FILE_NAME)
assert T.sorted_for_import(tasks)[0]["id"] == w._task_id, "新会话仍然排最前"

# GUI 传入的桌面目录优先于自己算的那份（OneDrive 重定向场景）
T.tool_dir = REAL_TOOL_DIR
T.set_desktop_dir(os.path.join(TMP, "desk"))
assert T.tool_dir().name == T.TOOL_DIR_NAME and str(T.tool_dir()).startswith(os.path.join(TMP, "desk")), T.tool_dir()
assert str(T.tool_dir(os.path.join(TMP, "other"))).startswith(os.path.join(TMP, "other")), \
    "显式传的桌面目录要盖过默认值"
T.set_desktop_dir("")
T.tool_dir = lambda desktop_dir=None: Path(TMP)
print("⑧ 历史会话补登记 + 桌面目录归一（OneDrive 重定向不跑偏）✓")


# ══════════════ ⑨ 采集完自动导入 + 「清洗数据」按钮 ══════════════
def _pos_in_layout(root_layout, widget):
    """在嵌套布局里找控件 → (根布局第几行, 那一行里的第几格)，用来断言“在哪、什么顺序”"""
    def _row_index(lay):
        if lay is None:
            return None
        for i in range(lay.count()):
            item = lay.itemAt(i)
            if item.widget() is widget:
                return i
            got = _row_index(item.layout())
            if got is not None:
                return got
        return None

    for top in range(root_layout.count()):
        cell = _row_index(root_layout.itemAt(top).layout())
        if cell is not None:
            return top, cell
    return None


group_lay = w.export_btn.parentWidget().layout()
epos = _pos_in_layout(group_lay, w.export_btn)
cpos = _pos_in_layout(group_lay, w.clean_btn)
apos = _pos_in_layout(group_lay, w.auto_import_check)
assert cpos and cpos[0] == epos[0] and cpos[1] == epos[1] + 1, \
    f"「清洗数据」必须紧跟在「导出 Excel」后面：export={epos} clean={cpos}"
assert apos and apos[0] > epos[0], f"自动导入勾选项要在按钮行下方（开始采集下面）：{apos} vs {epos}"
assert w.auto_import_check.isChecked() is False, "默认不能勾选"
assert w.clean_btn.isEnabled() and not w.export_btn.isEnabled(), \
    "没采到数据时导出置灰，但清洗其他渠道数据的能力不受影响"

# 未绑定账号：勾选当场拦住并把勾退回，不能让用户以为已经开了
use_http([ok(False)])
FakeBox.shown.clear()
saved_phone = w._yzd_phone
w._yzd_phone = ""
w.auto_import_check.setChecked(True)
assert w.auto_import_check.isChecked() is False, "未绑定时不能保持勾选"
assert w._auto_import_group is None
assert "绑定云展动力账号" in FakeBox.shown[-1][2], FakeBox.shown
w._yzd_phone = saved_phone

# 取消选分组：勾选同样要退回，不能留一个“看着开了其实没目标”的状态
def _reject_pick(self, *a):
    assert self._pick_only is True, "自动导入要复用只选分组的弹窗"
    return 0
orig_exec = G.YzdGroupDialog.exec_
G.YzdGroupDialog.exec_ = _reject_pick
FakeBox.shown.clear()
w.auto_import_check.setChecked(True)
assert w.auto_import_check.isChecked() is False and w._auto_import_group is None
assert FakeBox.shown == [], "用户自己点取消不需要弹窗教训，记日志就够"
assert "没选定分组" in w.log_text.toPlainText(), w.log_text.toPlainText()[-400:]


def _accept_pick(self, *a):
    assert self._pick_only is True, "自动导入要复用只选分组的弹窗"
    assert self.others_btn.isHidden(), "只选分组时不能出现「导入其他」（否则像马上就开始导了）"
    assert self.latest_btn.text() == "选用该分组", self.latest_btn.text()
    self.chosen_group = {"id": "11", "groupName": "光伏组", "clientNum": 5, "foreignFlag": 0}
    return G.QDialog.Accepted


G.YzdGroupDialog.exec_ = _accept_pick
FakeBox.shown.clear()
w.auto_import_check.setChecked(True)
assert w._auto_import_group and w._auto_import_group["id"] == "11", w._auto_import_group
assert "已设置采集完自动导入" in w.log_text.toPlainText() and "光伏组" in w.log_text.toPlainText()

# 采完自动上传：固定传本次会话，不拿历史任务凑数
Y.requests.request = FakeImportHttp()
uploads.clear()
FakeBox.shown.clear()
w.qizhidao_crawler.results = list(fake_records)
w._on_crawl_finished()
assert w._yzd_importing is True, "采集完成应直接进入自动导入"
assert w.isEnabled() is False and w._yzd_progress is not None, "自动导入同样要锁界面"
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
if w._yzd_import_worker:
    w._yzd_import_worker.wait(3000)
app.processEvents()
assert uploads and sum(uploads) == 6, f"自动导入应把本次会话去重后传完：{uploads}"
log_all = w.log_text.toPlainText()
assert "按勾选自动导入云展客户库" in log_all and "采集完成" in log_all, "日志里要能看到自动导入的提示"
assert "客户数" in log_all and "60 秒" in log_all, "要提示客户数是后端缓存值，不然后缀旧数字看着像没导进去"
assert w._yzd_importing is False and w.isEnabled() is True

# 没采到数据 / 没分组 → 自动导入跳过，不能报错也不能静默上传旧任务
w.qizhidao_crawler.results = []
uploads.clear()
FakeBox.shown.clear()
w._maybe_auto_import("采集完成")
assert uploads == [] and "没采到数据" in w.log_text.toPlainText()
w.auto_import_check.setChecked(False)
assert w._auto_import_group is None and w.auto_import_check.isChecked() is False
G.YzdGroupDialog.exec_ = orig_exec
print("⑨ 采集完自动导入：勾选拦未绑定/取消、采完自动传本次会话 ✓")

# 「清洗数据」：外部渠道的表 → 智能匹配列 → 写出入库格式
class FakeFileDialogHolder:
    open_files = []
    save_path = ""


G.QFileDialog = type("FakeFileDialog", (), {
    "getOpenFileNames": staticmethod(lambda *a, **k: (list(FakeFileDialogHolder.open_files), "")),
    "getSaveFileName": staticmethod(lambda *a, **k: (FakeFileDialogHolder.save_path, "")),
})
ext_src = os.path.join(TMP, "其他渠道名单.xlsx")
from openpyxl import Workbook as _WB

_wb = _WB()
_ws = _wb.active
_ws.append(["序号", "客户名称", "单位名称", "手机", "电子邮箱", "备注"])
_ws.append([1, "陈师傅", "东莞宏发精密五金有限公司", 13800138000, "chen@hf.com", "展会拿的名片"])
_ws.append([2, "刘工", "佛山新力机械实业有限公司", "0757-88886666", "liu@xl.com", ""])
_ws.append([3, "", "东莞宏发精密五金有限公司", "13900001111", "", "第二个联系人"])
_wb.save(ext_src)

FakeFileDialogHolder.open_files = [ext_src]
FakeFileDialogHolder.save_path = os.path.join(TMP, G.CLEAN_FILE_NAME)
FakeBox.shown.clear()
w._clean_external_data()
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
assert w._cleaning is False and w.isEnabled() is True, "清洗结束必须解锁"
assert FakeBox.shown[-1][:2] == ("info", "清洗完成"), FakeBox.shown
assert "共清洗 3 条" in FakeBox.shown[-1][2], FakeBox.shown[-1][2]
assert os.path.exists(FakeFileDialogHolder.save_path)
wb = openpyxl.load_workbook(FakeFileDialogHolder.save_path)
ws = wb.active
head = [c.value for c in ws[1]]
body = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
wb.close()
assert head[0] == "客户名称" and "企业名称" in head and "手机号" in head, head
assert "刘工" in [r[0] for r in body], body
assert "13800138000" in [str(c) for row in body for c in row], "数字手机号要被洗成完整号码"
assert any("展会拿的名片" in str(c) for row in body for c in row), "备注列不能丢"
assert "列匹配结果" in w.log_text.toPlainText(), "识别了哪些列要写进日志"
print("⑩ 清洗数据：智能匹配外部列名 → 写出《清洗后客户库入库数据》 ✓")

# 洗不了的文件：说清原因，不能默默产出一个空文件
bad_src = os.path.join(TMP, "旧格式.xls")
with open(bad_src, "w", encoding="utf-8") as fh:
    fh.write("fake")
FakeFileDialogHolder.open_files = [bad_src]
FakeFileDialogHolder.save_path = os.path.join(TMP, "不该生成.xlsx")
FakeBox.shown.clear()
w._clean_external_data()
deadline = time.time() + 20
while time.time() < deadline and not FakeBox.shown:
    app.processEvents()
    time.sleep(0.02)
app.processEvents()
assert FakeBox.shown[-1][0] == "warn" and "没读到任何数据" in FakeBox.shown[-1][2], FakeBox.shown
assert not os.path.exists(FakeFileDialogHolder.save_path), "失败不能留一个空文件在磁盘上"
assert w._cleaning is False and w.isEnabled() is True
print("⑪ 清洗数据：读不了的文件说清原因且不留空文件 ✓")

shutil.rmtree(TMP, ignore_errors=True)
print("\n云展动力对接离线验证全部通过 ✅")
