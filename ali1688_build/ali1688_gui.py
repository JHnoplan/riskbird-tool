"""
ali1688_gui.py - 1688 企业数据自动采集工具（PyQt5 GUI）

面向办公用户：导入关键词（txt/csv/xlsx/docx 或直接粘贴，自动识别多个关键词）
→ 后台采集（列表分页 → 详情 → 联系方式，按公司去重）→ 边采边自动保存 → 导出 Excel。
支持：采集间隔可调、被风控/需登录时弹窗提示人工处理后自动继续、暂停/继续/停止续采。

运行：python ali1688_gui.py
"""

import os
import re
import sys
import time
import zipfile
import threading
from pathlib import Path
from datetime import datetime

from PyQt5.QtCore import Qt, QThread, pyqtSignal, QTimer
from PyQt5.QtGui import QFont
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QTextEdit, QFileDialog, QSpinBox, QDoubleSpinBox,
    QProgressBar, QMessageBox, QGroupBox, QPlainTextEdit, QDialog,
    QLineEdit, QCheckBox, QListWidget, QListWidgetItem, QAbstractItemView,
    QStackedWidget, QButtonGroup, QInputDialog,
)

from ali1688_crawler import Ali1688Crawler
import customer_library_cleaner as _cleaner

try:
    # 云展动力客户库对接（接口 3/4/5/6）+ 采集任务索引/绑定账号持久化
    import yzd_client as _yzd
    import task_registry as _tasks
    HAS_YZD = True
except ImportError:
    HAS_YZD = False


# ── 关键词文件读取（无第三方依赖，docx 走 zip+xml）────────────────
def read_keywords_from_file(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".txt", ".csv", ".md", ".json", ""):
        return p.read_text(encoding="utf-8", errors="ignore")
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=False, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None and str(cell).strip():
                        lines.append(str(cell).strip())
        wb.close()
        return "\n".join(lines)
    if ext == ".docx":
        # 只提取文本节点，不引入完整 XML 解析器（规避实体/外部实体解析风险）
        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        lines = []
        for para in re.split(r"</w:p>", xml):
            texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", para, flags=re.S)
            line = "".join(texts)
            line = (line.replace("&amp;", "&").replace("&lt;", "<")
                        .replace("&gt;", ">").replace("&quot;", '"'))
            if line.strip():
                lines.append(line.strip())
        return "\n".join(lines)
    # 其它按纯文本兜底
    return p.read_text(encoding="utf-8", errors="ignore")


# ── 数据清洗：直接复用成熟的 customer_library_cleaner（对齐《客户库数据导入模板》）──


# ── 采集线程 ───────────────────────────────────────────────────
class CrawlWorker(QThread):
    def __init__(self, crawler, keywords):
        super().__init__()
        self.crawler = crawler
        self.keywords = keywords

    def run(self):
        try:
            self.crawler.crawl(self.keywords, start_index=0)
        except Exception as e:
            self.crawler._log(f"❌ 采集线程异常：{e}")


# ── 自动保存线程 ───────────────────────────────────────────────
class AutoSaveWorker(QThread):
    saved = pyqtSignal(str, int)   # 路径, 条数
    failed = pyqtSignal(str)

    def __init__(self, crawler, rows, session_dir):
        super().__init__()
        self.crawler = crawler
        self.rows = rows            # 已取好的快照（记录级拷贝）
        self.session_dir = session_dir

    def run(self):
        try:
            ts_dir = Path(self.session_dir)
            ts_dir.mkdir(parents=True, exist_ok=True)
            # 原始明细
            raw = ts_dir / "1688企业数据_原始.xlsx"
            self._write_raw(raw, self.rows)
            # 清洗后入库（完全对齐《客户库数据导入模板.xlsx》）
            cl = ts_dir / "1688企业数据_清洗后入库数据.xlsx"
            self._write_cleaned(cl, self.rows)
            self.saved.emit(str(raw), len(self.rows))
        except PermissionError:
            self.failed.emit("文件正被 Excel/WPS 打开占用，本次自动保存跳过，关闭后下轮补上")
        except Exception as e:
            self.failed.emit(f"自动保存失败：{e}")

    @staticmethod
    def _write_raw(path, rows):
        from openpyxl import Workbook
        headers = Ali1688Crawler.EXPORT_HEADERS
        tmp = str(path) + ".tmp"
        wb = Workbook(); ws = wb.active; ws.title = "1688企业数据"
        ws.append(headers)
        for i, r in enumerate(rows, 1):
            ws.append([i] + [r.get(h, "") for h in headers[1:]])
        wb.save(tmp)
        os.replace(tmp, str(path))

    def _write_cleaned(self, path, rows):
        # 统一走 crawler.clean_records：清洗 → 剔除「无手机号/固话/邮箱」的行（原数据不动）
        clean_rows = self.crawler.clean_records(rows, log=lambda m: None)
        template = _cleaner.locate_template([self.session_dir, str(Path.home() / "Downloads")])
        tmp = str(path) + ".tmp"
        _cleaner.write_clean_excel(clean_rows, tmp, template_path=template, log=lambda m: None)
        os.replace(tmp, str(path))


# ══════════════════════════════════════════════
#  云展动力：账号绑定 / 客户库分组 / 导入（与企业数据自动采集工具功能完全一致）
# ═══════════════════════════════════════════════
YZD_IMPORT_CHUNK = 5000   # 接口6 建议单次 ≤5000 行，超出自动切分成多次上传


def _yzd_font(size=10, bold=False):
    return QFont("Microsoft YaHei" if sys.platform == "win32" else "Arial",
                 size, QFont.Bold if bold else QFont.Normal)


class YzdBindWorker(QThread):
    """接口3：校验手机号格式与账号是否存在（真实网络请求，必须放线程里跑）。"""
    finished_bind = pyqtSignal(bool, str)   # ok, 成功=手机号 / 失败=提示

    def __init__(self, phone):
        super().__init__()
        self.phone = phone

    def run(self):
        try:
            _yzd.YzdClient().check_phone(self.phone)
            self.finished_bind.emit(True, self.phone)
        except _yzd.YzdError as e:
            self.finished_bind.emit(False, str(e))
        except Exception as e:
            self.finished_bind.emit(False, f"校验手机号异常：{e}")


class YzdGroupsWorker(QThread):
    """接口4：拉取分组列表（foreign_flag 0=国内 1=国外）"""
    finished_groups = pyqtSignal(int, list, str)   # flag, groups, error

    def __init__(self, phone, foreign_flag, keyword=""):
        super().__init__()
        self.phone = phone
        self.foreign_flag = foreign_flag
        self.keyword = keyword

    def run(self):
        try:
            groups = _yzd.YzdClient(self.phone).list_groups(
                foreign_flag=self.foreign_flag, group_name=self.keyword)
            self.finished_groups.emit(self.foreign_flag, groups, "")
        except Exception as e:
            self.finished_groups.emit(self.foreign_flag, [], str(e))


class YzdCreateGroupWorker(QThread):
    """接口5：新建分组（同名后端会自动追加序号后缀，这里不自造）"""
    finished_create = pyqtSignal(dict, str)   # group, error

    def __init__(self, phone, group_name, foreign_flag):
        super().__init__()
        self.phone = phone
        self.group_name = group_name
        self.foreign_flag = foreign_flag

    def run(self):
        try:
            group = _yzd.YzdClient(self.phone).create_group(
                self.group_name, foreign_flag=self.foreign_flag)
            self.finished_create.emit(group, "")
        except Exception as e:
            self.finished_create.emit({}, str(e))


class YzdImportWorker(QThread):
    """接口6：把若干采集任务逐个清洗成 9 列 Excel 后导入指定分组。

    与风鸟版唯一差异：1688 的记录要先经 crawler.clean_records（映射列 + 剔除无
    任何联系方式的行），再交给 _yzd.dedupe_rows 做批内去重，保证与本地清洗口径一致。
    """
    logged = pyqtSignal(str)
    progress = pyqtSignal(int, int, str)                    # done, total, 阶段文案
    task_done = pyqtSignal(str, int, int, int, str)         # 任务名, importTotal, saveTotal, failTotal, error
    finished_all = pyqtSignal(list, str)                    # 逐任务结果, 致命错误

    def __init__(self, crawler, phone, jobs, group_id, foreign_flag, tmp_dir):
        super().__init__()
        self.crawler = crawler
        self.phone = phone
        self.jobs = jobs            # [{"name":…, "records": list|None, "raw_file":…}]
        self.group_id = group_id
        self.foreign_flag = foreign_flag
        self.tmp_dir = tmp_dir

    def _records(self, job):
        """当前会话优先用内存快照，历史任务回读会话目录里的原数据文件"""
        if job.get("records") is not None:
            return job["records"]
        return _tasks.read_records_from_xlsx(job.get("raw_file", ""))

    def run(self):
        """整个流程包一层兜底：中途抛出的异常会把线程直接结束，而界面还在等
        finished_all 解锁，用户看到的就是一个关不掉的加载框。"""
        try:
            self._run_all()
        except Exception as e:
            self.finished_all.emit([], f"导入线程异常：{e}")

    def _run_all(self):
        total = len(self.jobs)
        results = []
        client = _yzd.YzdClient(self.phone, log=self.logged.emit)
        for idx, job in enumerate(self.jobs, 1):
            name = job.get("name") or f"任务{idx}"
            self.progress.emit(idx - 1, total, f"正在导入 {idx}/{total}：{name}")
            try:
                records = self._records(job)
            except Exception as e:
                self.logged.emit(f"⚠ 读取「{name}」数据失败：{e}")
                results.append((name, 0, 0, 0, f"读取数据失败：{e}"))
                self.task_done.emit(name, 0, 0, 0, f"读取数据失败：{e}")
                continue
            if not records:
                msg = "该任务没有原数据"
                self.logged.emit(f"⚠ 「{name}」跳过：{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue
            try:
                # 1688 专属：原始记录 → 映射清洗 → 剔除无手机号/固话/邮箱的行
                rows = self.crawler.clean_records(records, log=self.logged.emit)
                # 批内去重：接口6 自身不做重复校验，同一批里完全一样的行只传一次
                rows = _yzd.dedupe_rows(rows, log=self.logged.emit)
            except Exception as e:
                msg = f"清洗失败：{e}"
                self.logged.emit(f"⚠ 「{name}」{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue
            if not rows:
                msg = "清洗后没有可导入的记录"
                self.logged.emit(f"⚠ 「{name}」跳过：{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue

            chunks = [rows[i:i + YZD_IMPORT_CHUNK] for i in range(0, len(rows), YZD_IMPORT_CHUNK)] or [[]]
            imp = sav = fail = 0
            err = ""
            for c_idx, chunk in enumerate(chunks, 1):
                tag = f"{name}（{c_idx}/{len(chunks)}）" if len(chunks) > 1 else name
                self.progress.emit(idx - 1, total,
                                   f"正在导入 {idx}/{total}：{tag} · {len(chunk)} 条")
                file_path = os.path.join(self.tmp_dir, f"yzd_import_{idx}_{c_idx}.xlsx")
                try:
                    written = _yzd.write_import_workbook(chunk, file_path, log=self.logged.emit)
                    if written == 0:
                        raise RuntimeError("清洗结果全为空行，没有可上传的数据")
                    self.logged.emit(f"⬆ {tag}：已生成上传文件 {written} 行，正在提交云展…")
                    res = client.import_excel(file_path, group_id=self.group_id,
                                              foreign_flag=self.foreign_flag)
                    imp += res["importTotal"]
                    sav += res["saveTotal"]
                    fail += res["failTotal"]
                    self.logged.emit(f"  ↳ 云展返回：导入 {res['importTotal']} / 成功 {res['saveTotal']} / 失败 {res['failTotal']}")
                except Exception as e:
                    err = str(e)
                    self.logged.emit(f"❌ {tag} 导入失败：{err}")
                    break
                finally:
                    try:
                        if os.path.exists(file_path):
                            os.remove(file_path)
                    except Exception:
                        pass
            results.append((name, imp, sav, fail, err))
            self.task_done.emit(name, imp, sav, fail, err)
        self.progress.emit(total, total, "导入结束")
        self.finished_all.emit(results, "")


class YzdProgressDialog(QDialog):
    """导入期间的模态遮罩：锁住主界面 + busy 进度条动画。没做完不给关。"""

    def __init__(self, parent, title="正在导入云展客户库"):
        super().__init__(parent)
        self._allow_close = False
        self.setWindowTitle(title)
        self.setModal(True)
        self.setWindowFlags(Qt.Dialog | Qt.CustomizeWindowHint | Qt.WindowTitleHint)
        self.setMinimumWidth(460)
        layout = QVBoxLayout(self)
        layout.setSpacing(14)
        layout.setContentsMargins(24, 22, 24, 20)
        self.info_lbl = QLabel("准备中…")
        self.info_lbl.setFont(_yzd_font(11))
        self.info_lbl.setWordWrap(True)
        layout.addWidget(self.info_lbl)
        self.bar = QProgressBar()
        self.bar.setRange(0, 0)          # busy 动画
        self.bar.setTextVisible(False)
        self.bar.setMinimumHeight(20)
        layout.addWidget(self.bar)
        self.detail_lbl = QLabel("导入期间请勿关闭程序，也不要执行其他操作。")
        self.detail_lbl.setFont(_yzd_font(9))
        self.detail_lbl.setStyleSheet("color: #999;")
        self.detail_lbl.setWordWrap(True)
        layout.addWidget(self.detail_lbl)

    def set_text(self, main: str, detail: str = ""):
        self.info_lbl.setText(main)
        if detail:
            self.detail_lbl.setText(detail)

    def allow_close(self):
        self._allow_close = True

    def closeEvent(self, event):
        event.accept() if self._allow_close else event.ignore()


class YzdGroupDialog(QDialog):
    """「导入云展客户库」弹窗：第一页国内/国外 tab + 搜索 + 新建分组 + 分组单选，
    第二页历史采集任务多选。pick_group_only=True 用于自动导入预选分组。"""
    MODE_LATEST = "latest"
    MODE_OTHERS = "others"
    HINT_IMPORT = "选完分组后点右下「导入最新」或「导入其他」"
    HINT_PICK = "选一个分组作为采集完自动导入的目标，然后点右下「选用该分组」"

    def __init__(self, phone, parent=None, pick_group_only=False):
        super().__init__(parent)
        self.phone = phone
        self._pick_only = bool(pick_group_only)
        self._hint_text = self.HINT_PICK if self._pick_only else self.HINT_IMPORT
        self.mode = ""
        self.chosen_group = None
        self.chosen_tasks = []
        self._groups = {0: None, 1: None}
        self._flag = 0
        self._worker = None
        self._create_worker = None
        self._pending_name = ""
        self._created_id = ""
        self._task_map = {}
        self.setWindowTitle("选择自动导入的客户库分组" if self._pick_only else "导入云展客户库")
        self.setMinimumSize(560, 460)
        self.resize(600, 500)
        self._init_ui()
        self._load_groups()

    def _init_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 16)
        root.setSpacing(10)
        self.stack = QStackedWidget()
        root.addWidget(self.stack, 1)
        self.stack.addWidget(self._build_group_page())
        self.stack.addWidget(self._build_task_page())

    def _build_group_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setSpacing(10)

        tab_row = QHBoxLayout()
        self.tab_btns = {}
        self.tab_group = QButtonGroup(self)
        self.tab_group.setExclusive(True)
        for flag, text in ((0, "国内客户库"), (1, "国外客户库")):
            btn = QPushButton(text)
            btn.setCheckable(True)
            btn.setChecked(flag == 0)
            btn.setMinimumHeight(32)
            btn.setStyleSheet(
                "QPushButton { background: #eef1f6; color: #1a3c6e; border: 1px solid #d0d8e8; }"
                "QPushButton:checked { background: #4472C4; color: white; font-weight: bold; }")
            btn.clicked.connect(lambda _c, f=flag: self._switch_flag(f))
            self.tab_group.addButton(btn, flag)
            self.tab_btns[flag] = btn
            tab_row.addWidget(btn)
        tab_row.addStretch()
        self.refresh_btn = QPushButton("刷新")
        self.refresh_btn.setMinimumHeight(32)
        self.refresh_btn.setMinimumWidth(64)
        self.refresh_btn.setToolTip(
            "重新向云展后端查一次分组和客户数。\n"
            "「客户数」是后端统计值（最长约 60 秒才更新），刚导完立刻看可能还是旧数字。")
        self.refresh_btn.clicked.connect(self._refresh_groups)
        tab_row.addWidget(self.refresh_btn)
        self.new_group_btn = QPushButton("+ 新建分组")
        self.new_group_btn.setMinimumHeight(32)
        self.new_group_btn.clicked.connect(self._create_group)
        tab_row.addWidget(self.new_group_btn)
        lay.addLayout(tab_row)

        search_row = QHBoxLayout()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("搜索分组名称（模糊匹配）")
        self.search_edit.setMinimumHeight(30)
        self.search_edit.returnPressed.connect(self._search)
        search_row.addWidget(self.search_edit, 1)
        self.search_btn = QPushButton("搜索")
        self.search_btn.setMinimumHeight(30)
        self.search_btn.setMinimumWidth(70)
        self.search_btn.clicked.connect(self._search)
        search_row.addWidget(self.search_btn)
        lay.addLayout(search_row)

        self.group_status_lbl = QLabel("正在加载分组…")
        self.group_status_lbl.setFont(_yzd_font(9))
        self.group_status_lbl.setStyleSheet("color: #666;")
        self.group_status_lbl.setWordWrap(True)
        lay.addWidget(self.group_status_lbl)

        self.group_list = QListWidget()
        self.group_list.setMinimumHeight(180)
        self.group_list.setSelectionMode(QAbstractItemView.SingleSelection)
        self.group_list.itemSelectionChanged.connect(self._update_hint)
        lay.addWidget(self.group_list, 1)

        self.hint_lbl = QLabel(self._hint_text)
        self.hint_lbl.setFont(_yzd_font(9))
        self.hint_lbl.setStyleSheet("color: #666;")
        self.hint_lbl.setWordWrap(True)
        lay.addWidget(self.hint_lbl)

        btn_row = QHBoxLayout()
        self.cancel_btn = QPushButton("取消")
        self.cancel_btn.setMinimumHeight(36)
        self.cancel_btn.setMinimumWidth(90)
        self.cancel_btn.clicked.connect(self.reject)
        btn_row.addWidget(self.cancel_btn)
        btn_row.addStretch()
        self.others_btn = QPushButton("导入其他")
        self.others_btn.setMinimumHeight(36)
        self.others_btn.setMinimumWidth(110)
        self.others_btn.clicked.connect(self._show_tasks)
        btn_row.addWidget(self.others_btn)
        self.latest_btn = QPushButton("导入最新")
        self.latest_btn.setMinimumHeight(36)
        self.latest_btn.setMinimumWidth(110)
        self.latest_btn.setDefault(True)
        self.latest_btn.clicked.connect(self._accept_latest)
        btn_row.addWidget(self.latest_btn)
        if self._pick_only:
            self.others_btn.setVisible(False)
            self.latest_btn.setText("选用该分组")
        lay.addLayout(btn_row)
        return page

    def _build_task_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setSpacing(10)

        top = QHBoxLayout()
        self.back_btn = QPushButton("← 返回")
        self.back_btn.setMinimumHeight(32)
        self.back_btn.clicked.connect(lambda: self.stack.setCurrentIndex(0))
        top.addWidget(self.back_btn)
        title = QLabel("选择要导入的采集任务（可多选，必选）")
        title.setFont(_yzd_font(11, True))
        top.addWidget(title, 1)
        self.all_chk = QCheckBox("全选")
        self.all_chk.toggled.connect(self._toggle_all)
        top.addWidget(self.all_chk)
        lay.addLayout(top)

        self.task_count_lbl = QLabel("共 0 个任务")
        self.task_count_lbl.setFont(_yzd_font(9))
        self.task_count_lbl.setStyleSheet("color: #666;")
        lay.addWidget(self.task_count_lbl)

        self.task_list = QListWidget()
        self.task_list.setMinimumHeight(200)
        self.task_list.setSelectionMode(QAbstractItemView.NoSelection)
        self.task_list.itemChanged.connect(lambda _i: self._update_task_hint())
        lay.addWidget(self.task_list, 1)

        self.task_hint_lbl = QLabel("至少选择一个任务")
        self.task_hint_lbl.setFont(_yzd_font(9))
        self.task_hint_lbl.setStyleSheet("color: #666;")
        self.task_hint_lbl.setWordWrap(True)
        lay.addWidget(self.task_hint_lbl)

        btn_row = QHBoxLayout()
        btn_row.addStretch()
        self.task_cancel_btn = QPushButton("取消")
        self.task_cancel_btn.setMinimumHeight(36)
        self.task_cancel_btn.clicked.connect(self.reject)
        btn_row.addWidget(self.task_cancel_btn)
        self.task_ok_btn = QPushButton("确认导入")
        self.task_ok_btn.setMinimumHeight(36)
        self.task_ok_btn.setMinimumWidth(110)
        self.task_ok_btn.clicked.connect(self._accept_others)
        btn_row.addWidget(self.task_ok_btn)
        lay.addLayout(btn_row)
        return page

    def _set_busy(self, busy: bool):
        for w in (self.new_group_btn, self.search_btn, self.search_edit,
                  self.refresh_btn, self.latest_btn, self.others_btn):
            w.setEnabled(not busy)
        if busy:
            QApplication.setOverrideCursor(Qt.WaitCursor)
        else:
            QApplication.restoreOverrideCursor()

    def _switch_flag(self, flag: int):
        if flag == self._flag and self._groups[flag] is not None:
            return
        self._flag = flag
        cached = self._groups.get(flag)
        if cached is None:
            self._load_groups()
        else:
            self._render_groups(cached, "")

    def _search(self):
        self._groups[self._flag] = None
        self._load_groups()

    def _refresh_groups(self):
        self._groups[self._flag] = None
        self._load_groups()

    def _load_groups(self):
        if self._worker is not None and self._worker.isRunning():
            return
        keyword = self.search_edit.text().strip()
        self.group_status_lbl.setText("正在加载分组…")
        self._set_busy(True)
        self._worker = YzdGroupsWorker(self.phone, self._flag, keyword)
        self._worker.finished_groups.connect(self._on_groups)
        self._worker.start()

    def _on_groups(self, flag: int, groups: list, error: str):
        self._groups[flag] = groups if not error else None
        if flag != self._flag:
            self._set_busy(False)
            return
        self._render_groups(groups, error)
        self._set_busy(False)

    def _render_groups(self, groups: list, error: str):
        self.group_list.clear()
        if error:
            self.group_status_lbl.setText(f"分组加载失败：{error}")
            self.group_status_lbl.setStyleSheet("color: #c62828;")
            self._update_hint()
            return
        self.group_status_lbl.setStyleSheet("color: #666;")
        if not groups:
            keyword = self.search_edit.text().strip()
            self.group_status_lbl.setText(
                f"没搜到分组，换个关键词试试" if keyword else
                "该账号下还没有分组，请先点「+ 新建分组」创建一个")
            self._update_hint()
            return
        self.group_status_lbl.setText(
            f"共 {len(groups)} 个分组，选中一个后点右下角导入" if not self._pick_only else
            f"共 {len(groups)} 个分组，选一个当自动导入目标")
        self.group_status_lbl.setToolTip(
            "括号里的「客户数」直接取自云展后端接口返回（不是本地统计）。\n"
            "后端是缓存值，最长约 60 秒才更新一次；刚导完看到数字没变是正常的，"
            "点右上「刷新」重新查一次即可。")
        for g in groups:
            num = g.get("clientNum")
            text = f"{g.get('groupName', '')}"
            if num is not None:
                text += f"　（客户数 {num}）"
            item = QListWidgetItem(text)
            item.setData(Qt.UserRole, g)
            self.group_list.addItem(item)
        if self._created_id and self._select_group_by_id(self._created_id):
            self._created_id = ""
        self._update_hint()

    def _current_group(self):
        items = self.group_list.selectedItems()
        return items[0].data(Qt.UserRole) if items else None

    def _update_hint(self, error: bool = False):
        if self._current_group():
            self.hint_lbl.setText("")
            self.hint_lbl.setStyleSheet("color: #666;")
        elif error:
            self.hint_lbl.setText("请先选择一个分组（必选）")
            self.hint_lbl.setStyleSheet("color: #c62828;")
        else:
            self.hint_lbl.setText(self._hint_text)
            self.hint_lbl.setStyleSheet("color: #666;")

    def _need_group(self):
        group = self._current_group()
        if group is None:
            self._update_hint(error=True)
            QMessageBox.warning(self, "提示", "请选择一个分组")
        return group

    def _create_group(self):
        if not self.phone:
            QMessageBox.warning(self, "提示", "请先绑定云展动力账号")
            return
        name, ok = QInputDialog.getText(self, "新建分组", "群组名称（最多 10 个字符）：")
        if not ok:
            return
        name = (name or "").strip()
        if not name:
            QMessageBox.warning(self, "提示", "群组名称不能为空！")
            return
        if len(name) > _yzd.GROUP_NAME_MAX:
            QMessageBox.warning(self, "提示", f"群组名称最多输入{_yzd.GROUP_NAME_MAX}个字符！")
            return
        if self._create_worker is not None and self._create_worker.isRunning():
            return
        self._set_busy(True)
        self.group_status_lbl.setText(f"正在创建分组「{name}」…")
        self._pending_name = name
        self._create_worker = YzdCreateGroupWorker(self.phone, name, self._flag)
        self._create_worker.finished_create.connect(self._on_group_created)
        self._create_worker.start()

    def _on_group_created(self, group: dict, error: str):
        if error:
            self._set_busy(False)
            self.group_status_lbl.setText(f"新建分组失败：{error}")
            QMessageBox.warning(self, "新建分组失败", error)
            return
        self._groups = {0: None, 1: None}
        self._created_id = str(group.get("id") or "")
        self.search_edit.clear()
        self.group_status_lbl.setText(f"已创建分组「{group.get('groupName') or self._pending_name}」")
        self._load_groups()

    def _select_group_by_id(self, gid: str):
        for i in range(self.group_list.count()):
            item = self.group_list.item(i)
            data = item.data(Qt.UserRole) or {}
            if str(data.get("id")) == gid:
                self.group_list.setCurrentItem(item)
                return True
        return False

    def _accept_latest(self):
        group = self._need_group()
        if group is None:
            return
        self.mode = self.MODE_LATEST
        self.chosen_group = group
        self.accept()

    def _show_tasks(self):
        group = self._need_group()
        if group is None:
            return
        self.chosen_group = group
        self._fill_tasks()
        self.stack.setCurrentIndex(1)

    def _fill_tasks(self):
        self.task_list.clear()
        self.all_chk.blockSignals(True)
        self.all_chk.setChecked(False)
        self.all_chk.blockSignals(False)
        tasks = _tasks.sorted_for_import(_tasks.load_tasks())
        self._task_map = {}
        for t in tasks:
            name = t.get("name") or t.get("id") or "未命名任务"
            text = (f"{name}\n"
                    f"采集时间 {t.get('started_at') or '—'}"
                    f"　{t.get('count') or 0} 条"
                    f"　[{t.get('status') or '—'}]")
            item = QListWidgetItem(text)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Unchecked)
            item.setData(Qt.UserRole, t)
            self.task_list.addItem(item)
            self._task_map[id(item)] = t
        if not tasks:
            self.task_count_lbl.setText(
                "暂无采集任务记录：先在主界面完成一次「开始采集 → 采集完成」")
        else:
            self.task_count_lbl.setText(
                f"共 {len(tasks)} 个任务（按采集时间倒序）")
        self._update_task_hint()

    def _toggle_all(self, checked: bool):
        for i in range(self.task_list.count()):
            self.task_list.item(i).setCheckState(Qt.Checked if checked else Qt.Unchecked)

    def _checked_tasks(self):
        out = []
        for i in range(self.task_list.count()):
            item = self.task_list.item(i)
            if item.checkState() == Qt.Checked:
                task = item.data(Qt.UserRole)
                if task:
                    out.append(task)
        return out

    def _update_task_hint(self, error: bool = False):
        n = len(self._checked_tasks())
        if n:
            self.task_hint_lbl.setText("")
            self.task_hint_lbl.setStyleSheet("color: #666;")
        elif error:
            self.task_hint_lbl.setText("请至少选择一个任务（必选）")
            self.task_hint_lbl.setStyleSheet("color: #c62828;")
        else:
            self.task_hint_lbl.setText("勾选要导入的采集任务（可多选）")
            self.task_hint_lbl.setStyleSheet("color: #666;")
        self.task_ok_btn.setText(f"确认导入（{n}）" if n else "确认导入")

    def _accept_others(self):
        tasks = self._checked_tasks()
        if not tasks:
            self._update_task_hint(error=True)
            QMessageBox.warning(self, "提示", "请至少选择一个任务")
            return
        self.mode = self.MODE_OTHERS
        self.chosen_tasks = tasks
        self.accept()

    def closeEvent(self, event):
        for w in (self._worker, self._create_worker):
            if w is not None and w.isRunning():
                w.wait(3000)
        event.accept()


# ── 主窗口 ─────────────────────────────────────────────────────
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("1688 企业数据自动采集工具")
        self.resize(920, 720)
        self.crawler = Ali1688Crawler()
        self.worker = None
        self.autosave = None
        self._pending_save = False
        self._opening_browser = False
        self.keywords = []
        self._kw_source_file = ""
        self.session_dir = None
        self._first_start = None
        # 云展动力客户库（绑定账号 + 导入）
        if HAS_YZD:
            _tasks.set_desktop_dir(self._get_desktop_dir())
        self._yzd_phone = _tasks.load_account() if HAS_YZD else ""
        self._yzd_worker = None            # 绑定账号线程（存引用防被 GC）
        self._yzd_import_worker = None     # 导入线程
        self._yzd_dialog = None            # 选分组/选任务弹窗
        self._yzd_progress = None          # 导入锁定遮罩
        self._yzd_importing = False        # 导入中标志（closeEvent 也看它）
        self._yzd_jobs = []                # 本轮导入的任务快照
        self._yzd_group_name = ""          # 本轮导入目标分组名
        self._auto_import_group = None     # 采集完自动导入的目标分组
        self._task_id = ""                 # 本次采集任务 id（=会话目录名）
        self._build_ui()
        self._connect_signals()
        self._refresh_yzd_status()
        # 每 10 分钟自动保存一次
        self._timer = QTimer(self)
        self._timer.timeout.connect(self.request_autosave)
        self._timer.start(10 * 60 * 1000)

    def _get_desktop_dir(self) -> Path:
        """获取桌面路径（兼容 Windows OneDrive 重定向）"""
        desktop = Path.home() / "Desktop"
        if desktop.exists():
            return desktop
        for env in ("OneDrive", "OneDriveConsumer", "OneDriveCommercial"):
            od = os.environ.get(env)
            if od:
                d = Path(od) / "Desktop"
                if d.exists():
                    return d
        try:
            desktop.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        return desktop

    def _build_ui(self):
        root = QWidget()
        lay = QVBoxLayout(root)
        self.setCentralWidget(root)

        # 关键词输入区
        gb_kw = QGroupBox("① 关键词（上传文档或直接粘贴，自动识别多个关键词，换行/逗号/顿号分隔）")
        v = QVBoxLayout(gb_kw)
        hr = QHBoxLayout()
        self.btn_file = QPushButton("导入关键词文件 (txt/csv/xlsx/docx)")
        self.btn_paste = QPushButton("从剪贴板粘贴")
        self.btn_parse = QPushButton("解析关键词")
        hr.addWidget(self.btn_file); hr.addWidget(self.btn_paste); hr.addWidget(self.btn_parse)
        v.addLayout(hr)
        self.kw_text = QPlainTextEdit()
        self.kw_text.setPlaceholderText("每行一个关键词，例如：\nAHU净化空调箱\n光伏储能电池\nLED广告屏")
        self.kw_text.setMaximumHeight(140)
        v.addWidget(self.kw_text)
        self.lbl_kw = QLabel("识别到 0 个关键词")
        v.addWidget(self.lbl_kw)
        lay.addWidget(gb_kw)

        # 参数区
        gb_p = QGroupBox("② 采集参数")
        hp = QHBoxLayout(gb_p)
        hp.addWidget(QLabel("采集间隔(秒)："))
        self.spin_delay = QDoubleSpinBox()
        self.spin_delay.setRange(1.0, 60.0); self.spin_delay.setValue(3.0); self.spin_delay.setSingleStep(0.5)
        hp.addWidget(self.spin_delay)
        hp.addWidget(QLabel("每关键词最多页数："))
        self.spin_pages = QSpinBox(); self.spin_pages.setRange(1, 50); self.spin_pages.setValue(50)
        hp.addWidget(self.spin_pages)
        hp.addWidget(QLabel("   提示：挂着过夜请把间隔调到 6~10 秒更安全"))
        lay.addWidget(gb_p)

        # 控制按钮
        hb = QHBoxLayout()
        self.btn_open = QPushButton("打开浏览器 / 登录1688")
        self.btn_start = QPushButton("开始采集")
        self.btn_pause = QPushButton("暂停")
        self.btn_stop = QPushButton("停止")
        self.btn_export = QPushButton("导出 Excel")
        for b in (self.btn_open, self.btn_start, self.btn_pause, self.btn_stop, self.btn_export):
            b.setMinimumHeight(36); hb.addWidget(b)
        self.btn_pause.setEnabled(False); self.btn_stop.setEnabled(False)
        lay.addLayout(hb)

        # 采集完自动导入（默认不勾选）
        self.auto_import_check = QCheckBox("采集完自动导入云展客户库")
        self.auto_import_check.setToolTip(
            "勾选时会先让你选一个客户库分组；采集完成后自动把本次采集的数据\n"
            "上传到该分组，并在日志里提示导入结果（不需再手动点「导入云展客户库」）。\n"
            "需先绑定云展动力账号。手动点「停止」不会触发自动导入。\n"
            "不勾选（默认）：采集完只能自己点按钮导入。")
        self.auto_import_check.toggled.connect(self._on_auto_import_toggled)
        lay.addWidget(self.auto_import_check)

        # 云展动力客户库（绑定账号 + 导入）
        yzd_group = QGroupBox("云展动力客户库")
        yl = QHBoxLayout(yzd_group)
        self.yzd_status_lbl = QLabel("未绑定云展动力账号")
        self.yzd_status_lbl.setFont(QFont("Microsoft YaHei" if sys.platform == "win32" else "Arial", 10))
        self.yzd_status_lbl.setMinimumWidth(200)
        yl.addWidget(self.yzd_status_lbl)
        yl.addStretch()
        self.bind_yzd_btn = QPushButton("绑定云展动力账号")
        self.bind_yzd_btn.setMinimumHeight(34)
        self.bind_yzd_btn.setMinimumWidth(150)
        self.bind_yzd_btn.setToolTip(
            "输入云展动力手机号，调云展接口校验格式与账号是否存在，校验通过后才保存绑定。\n"
            "绑定信息存在：桌面/1688企业数据自动采集工具/_云展账号.json（跨重启仍有效）。")
        self.bind_yzd_btn.clicked.connect(self._bind_yzd_account)
        yl.addWidget(self.bind_yzd_btn)
        self.import_yzd_btn = QPushButton("导入云展客户库")
        self.import_yzd_btn.setMinimumHeight(34)
        self.import_yzd_btn.setMinimumWidth(140)
        self.import_yzd_btn.setToolTip(
            "选择客户库分组后，把采集结果上传到云展动力客户库。\n"
            "· 导入最新：只传最近一次采集任务；\n"
            "· 导入其他：从历史采集任务里挑（可多选）。\n"
            "上传中间会锁定界面并显示加载动画，走完弹结果汇总。")
        self.import_yzd_btn.clicked.connect(self._open_yzd_import)
        yl.addWidget(self.import_yzd_btn)
        lay.addWidget(yzd_group)

        # 提示条
        self.banner = QLabel("")
        self.banner.setVisible(False)
        self.banner.setWordWrap(True)
        self.banner.setStyleSheet("background:#fff3cd;color:#856404;border:1px solid #ffeeba;"
                                  "padding:8px;border-radius:4px;font-weight:bold;")
        lay.addWidget(self.banner)

        # 进度
        self.progress = QProgressBar(); self.progress.setValue(0)
        lay.addWidget(self.progress)
        self.lbl_stats = QLabel("就绪")
        lay.addWidget(self.lbl_stats)

        # 日志
        self.log = QTextEdit(); self.log.setReadOnly(True)
        self.log.setFont(QFont("Menlo", 10) if sys.platform != "win32" else QFont("Consolas", 9))
        lay.addWidget(self.log, 1)

    def _connect_signals(self):
        s = self.crawler.signals
        s.log_signal.connect(self._on_log)
        s.progress_signal.connect(self._on_progress)
        s.stats_signal.connect(lambda t: self.lbl_stats.setText(t))
        s.crawl_finished_signal.connect(self._on_finished)
        s.blocked_signal.connect(lambda r: self._show_banner(
            f"⚠ 触发 1688 安全验证：请在弹出的 Chrome 窗口里手动完成滑块/验证，通过后会自动继续采集。", "#f8d7da", "#721c24"))
        s.block_cleared_signal.connect(lambda r: self._hide_banner())
        s.login_required_signal.connect(lambda r: self._show_banner(
            "🔑 需要登录 1688：请在 Chrome 窗口里扫码/账号登录，登录成功后自动继续。", "#cce5ff", "#004085"))
        self.btn_file.clicked.connect(self._on_import_file)
        self.btn_paste.clicked.connect(self._on_paste)
        self.btn_parse.clicked.connect(self._on_parse)
        self.btn_open.clicked.connect(self._on_open_browser)
        self.btn_start.clicked.connect(self._on_start)
        self.btn_pause.clicked.connect(self._on_pause)
        self.btn_stop.clicked.connect(self._on_stop)
        self.btn_export.clicked.connect(self._on_export)

    # ── 日志/进度 ──────────────────────────────────────────────
    def _on_log(self, msg):
        self.log.append(msg)
        self.log.verticalScrollBar().setValue(self.log.verticalScrollBar().maximum())

    def _on_progress(self, cur, total):
        # 发现式采集总数未知：cur 为已采公司数，用忙碌进度条 + 计数标签体现
        self.lbl_stats.setText(f"已采集 {cur} 家公司")

    def _show_banner(self, text, bg, fg):
        self.banner.setText(text)
        self.banner.setStyleSheet(f"background:{bg};color:{fg};padding:8px;border-radius:4px;font-weight:bold;")
        self.banner.setVisible(True)

    def _hide_banner(self):
        self.banner.setVisible(False)

    # ── 关键词 ─────────────────────────────────────────────────
    def _on_import_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择关键词文件", "", "支持的文件 (*.txt *.csv *.xlsx *.xlsm *.docx *.md);;所有文件 (*)")
        if not path:
            return
        try:
            text = read_keywords_from_file(path)
            self.kw_text.setPlainText(text)
            self._kw_source_file = path
            self._on_parse()
            self._on_log(f"已导入文件：{path}")
        except Exception as e:
            QMessageBox.warning(self, "导入失败", f"读取文件出错：{e}")

    def _on_paste(self):
        t = QApplication.clipboard().text()
        if t:
            self.kw_text.appendPlainText(t)
            self._on_parse()

    def _on_parse(self):
        self.keywords = Ali1688Crawler.parse_keywords(self.kw_text.toPlainText())
        self.lbl_kw.setText(f"识别到 {len(self.keywords)} 个关键词：" +
                            ("、".join(self.keywords[:12]) + ("…" if len(self.keywords) > 12 else "") if self.keywords else "无"))

    # ── 浏览器/登录 ────────────────────────────────────────────
    def _on_open_browser(self):
        # 连点会起多个后台线程各自启浏览器：加一把重入锁，启动期间再点直接忽略
        if getattr(self, "_opening_browser", False):
            return
        self._opening_browser = True
        self.btn_open.setEnabled(False)
        self._on_log("正在启动浏览器…")

        def task():
            try:
                ok = self.crawler.start_browser(headed=True)
                if ok:
                    if self.crawler._goto(self.crawler.BASE_URL):
                        self.crawler._log("✅ 已打开 1688 首页，请确认已登录后再点「开始采集」")
                    else:
                        self.crawler._log("⚠ 浏览器已连接，但打开 1688 首页失败，请再点一次「打开浏览器」")
            finally:
                self._opening_browser = False

        threading.Thread(target=task, daemon=True).start()

        # 轮询回主线程解锁按钮：成功/失败都能及时恢复，不用盲等固定秒数
        def poll():
            if getattr(self, "_opening_browser", False):
                QTimer.singleShot(500, poll)
            else:
                self.btn_open.setEnabled(True)
        QTimer.singleShot(500, poll)

    # ── 采集控制 ───────────────────────────────────────────────
    def _on_start(self):
        if not self.keywords:
            self._on_parse()
        if not self.keywords:
            QMessageBox.information(self, "提示", "请先输入或导入关键词并解析。")
            return
        if self.worker and self.worker.isRunning():
            return
        self.crawler.delay_between = self.spin_delay.value()
        self.crawler.max_pages = self.spin_pages.value()
        if not self._first_start:
            self._first_start = datetime.now().strftime("%Y%m%d_%H%M%S")
            self.session_dir = str(self._get_desktop_dir() / "1688企业数据自动采集工具" / self._first_start)
        self._register_task()
        self.btn_start.setEnabled(False)
        self.btn_pause.setEnabled(True)
        self.btn_stop.setEnabled(True)
        self.progress.setRange(0, 0)   # 忙碌指示
        self.worker = CrawlWorker(self.crawler, self.keywords)
        self.worker.start()
        self.request_autosave()   # 先建目录

    def _on_pause(self):
        if self.crawler._pause_requested:
            self.crawler.resume_crawl()
            self.btn_pause.setText("暂停")
        else:
            self.crawler.pause_crawl()
            self.btn_pause.setText("继续")

    def _on_stop(self):
        self._manual_stop = True
        self.crawler.stop_crawl()

    def _on_finished(self):
        self.btn_start.setEnabled(True)
        self.btn_pause.setEnabled(False)
        self.btn_stop.setEnabled(False)
        self.btn_pause.setText("暂停")
        self.progress.setRange(0, 100)
        self.progress.setValue(100)
        self._hide_banner()
        self.request_autosave()   # 收工必落盘
        # 任务收尾：自然完成=已完成，人工停止=已停止（停止不触发自动导入）
        stopped = bool(getattr(self, "_manual_stop", False))
        self._manual_stop = False
        self._finish_task(_tasks.TASK_STATUS_STOPPED if stopped else _tasks.TASK_STATUS_DONE,
                          len(self.crawler.snapshot_results()))
        if not stopped:
            self._maybe_auto_import("采集完成")

    # ── 自动保存 / 导出 ────────────────────────────────────────
    def request_autosave(self):
        if not self.session_dir:
            return
        if self.autosave and self.autosave.isRunning():
            self._pending_save = True
            return
        rows = self.crawler.snapshot_results()
        if not rows:
            return
        self.autosave = AutoSaveWorker(self.crawler, rows, self.session_dir)
        self.autosave.saved.connect(self._on_saved)
        self.autosave.failed.connect(lambda m: self._on_log("⚠ " + m))
        self.autosave.finished.connect(self._after_autosave)
        self.autosave.start()

    def _after_autosave(self):
        if self._pending_save:
            self._pending_save = False
            self.request_autosave()

    def _on_saved(self, path, n):
        self._on_log(f"💾 已自动保存 {n} 条 → {path}")
        # 条数跟着自动保存同步进任务索引，「导入其他」里的条数才不会是旧值
        if HAS_YZD and self._task_id:
            try:
                _tasks.update_task(self._task_id, count=n)
            except Exception:
                pass

    def _on_export(self):
        rows = self.crawler.snapshot_results()
        if not rows:
            QMessageBox.information(self, "提示", "还没有采集到数据。")
            return
        default = f"1688企业数据_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "导出 Excel", str(Path.home() / "Desktop" / default),
                                              "Excel 文件 (*.xlsx)")
        if not path:
            return
        ok = self.crawler.export_excel(path)
        if ok:
            cpath = str(Path(path).with_name(Path(path).stem + "_清洗后入库数据.xlsx"))
            try:
                res = self.crawler.export_cleaned_excel(cpath)
            except Exception as e:
                res = ""
                self._on_log(f"清洗导出失败：{e}")
            if res:
                QMessageBox.information(
                    self, "导出完成",
                    f"原始明细：\n{path}\n\n清洗后入库数据（对齐客户库模板）：\n{cpath}")
            else:
                QMessageBox.information(self, "导出完成",
                                        f"原始明细已导出（无满足入库条件的数据）：\n{path}")

    # ── 云展动力：任务登记 / 绑定 / 导入 ───────────────────────
    def _current_task_id(self):
        return Path(self.session_dir).name if self.session_dir else ""

    def _register_task(self):
        """一次「开始采集」= 一个任务：id=会话目录名，任务名取关键词文档名/关键词。
        重复点「开始采集」（暂停后继续、采完又跑同一批）复用同一任务，保留最早 started_at。"""
        if not HAS_YZD or not self.session_dir:
            return
        try:
            tid = self._current_task_id()
            name_src = [self._kw_source_file] if self._kw_source_file else list(self.keywords)
            task = _tasks.new_task(tid, name_src, self.session_dir,
                                   count=len(self.crawler.snapshot_results()))
            existing = next((t for t in _tasks.load_tasks() if t.get("id") == tid), None)
            if existing:
                task["started_at"] = existing.get("started_at") or task["started_at"]
            _tasks.upsert_task(task)
            self._task_id = tid
            if existing is None:
                self._on_log(f"📁 已登记采集任务：{task['name']}")
        except Exception as e:
            self._on_log(f"⚠ 采集任务登记失败：{e}")

    def _finish_task(self, status: str, count: int):
        if not HAS_YZD or not self._task_id:
            return
        try:
            if _tasks.finish_task(self._task_id, status, count):
                self._on_log(f"📁 任务「{status}」已记录：合计 {count} 条")
        except Exception as e:
            self._on_log(f"⚠ 任务结束状态写入失败：{e}")

    def _refresh_yzd_status(self):
        if not hasattr(self, "yzd_status_lbl"):
            return
        if not HAS_YZD:
            self.yzd_status_lbl.setText("云展模块未加载（缺 yzd_client / task_registry）")
            self.yzd_status_lbl.setStyleSheet(
                "QLabel { padding: 4px 10px; border-radius: 4px; background: #fdecea; color: #c62828; }")
            self.bind_yzd_btn.setEnabled(False)
            self.import_yzd_btn.setEnabled(False)
            return
        if self._yzd_phone:
            self.yzd_status_lbl.setText(f"已绑定云展账号：{self._yzd_phone}")
            self.yzd_status_lbl.setStyleSheet(
                "QLabel { padding: 4px 10px; border-radius: 4px; background: #e8f5e9; color: #2e7d32; }")
            self.bind_yzd_btn.setText("更换绑定账号")
        else:
            self.yzd_status_lbl.setText("未绑定云展动力账号")
            self.yzd_status_lbl.setStyleSheet(
                "QLabel { padding: 4px 10px; border-radius: 4px; background: #f0f0f0; color: #999; }")
            self.bind_yzd_btn.setText("绑定云展动力账号")

    def _bind_yzd_account(self):
        if not HAS_YZD:
            QMessageBox.warning(self, "提示", "云展模块未加载，无法绑定账号")
            return
        if self._yzd_worker is not None and self._yzd_worker.isRunning():
            return
        tip = "请输入云展动力手机号"
        if self._yzd_phone:
            tip += f"（当前绑定：{self._yzd_phone}）"
        phone, ok = QInputDialog.getText(self, "绑定云展动力账号", tip,
                                         QLineEdit.Normal, self._yzd_phone)
        if not ok:
            return
        phone = (phone or "").strip().replace(" ", "").replace("-", "")
        if not phone or not _yzd.PHONE_RE.match(phone):
            QMessageBox.warning(self, "提示", "请输入正确的手机号")
            return
        self.bind_yzd_btn.setEnabled(False)
        self._on_log(f"🔗 正在校验云展动力账号 {phone} …")
        self._yzd_worker = YzdBindWorker(phone)
        self._yzd_worker.finished_bind.connect(self._on_bind_finished)
        self._yzd_worker.start()

    def _on_bind_finished(self, ok: bool, msg: str):
        self.bind_yzd_btn.setEnabled(HAS_YZD)
        if ok:
            self._yzd_phone = msg
            saved = _tasks.save_account(msg)
            self._on_log(f"✅ 已绑定云展动力账号：{msg}"
                         + ("" if saved else "（本地保存失败，下次启动需重新绑定）"))
            QMessageBox.information(self, "绑定成功", f"已绑定云展动力账号：{msg}")
        else:
            self._on_log(f"⚠ 绑定云展账号失败：{msg}")
            QMessageBox.warning(self, "绑定失败", msg)
        self._refresh_yzd_status()

    def _open_yzd_import(self):
        if not HAS_YZD:
            QMessageBox.warning(self, "提示", "云展模块未加载，无法导入客户库")
            return
        if self._yzd_importing:
            QMessageBox.information(self, "提示", "正在导入云展客户库，请等当前导入完成")
            return
        if not self._yzd_phone:
            QMessageBox.warning(self, "提示", "请先绑定云展动力账号")
            return
        try:
            added = _tasks.backfill_tasks()
            if added:
                self._on_log(f"📁 已补登记历史采集会话 {added} 个，可在「导入其他」里选中")
        except Exception as e:
            self._on_log(f"⚠ 历史会话补登记失败：{e}")
        dlg = YzdGroupDialog(self._yzd_phone, self)
        self._yzd_dialog = dlg
        if dlg.exec_() != QDialog.Accepted:
            return
        group = dlg.chosen_group or {}
        tasks = dlg.chosen_tasks if dlg.mode == YzdGroupDialog.MODE_OTHERS else None
        self._start_yzd_import(group, tasks)

    def _on_auto_import_toggled(self, checked: bool):
        if not checked:
            self._auto_import_group = None
            self._on_log("已关闭「采集完自动导入云展客户库」，采集完成后需手动点导入")
            return
        if not HAS_YZD:
            self._revert_auto_import("云展模块未加载，无法自动导入客户库")
            return
        if self._yzd_importing:
            self._revert_auto_import("当前正在导入，请等它结束后再开启自动导入")
            return
        if not self._yzd_phone:
            self._revert_auto_import("请先点「绑定云展动力账号」，绑定后才能选自动导入的目标分组")
            return
        dlg = YzdGroupDialog(self._yzd_phone, self, pick_group_only=True)
        self._yzd_dialog = dlg
        if dlg.exec_() != QDialog.Accepted or not dlg.chosen_group:
            self._auto_import_group = None
            chk = getattr(self, "auto_import_check", None)
            if chk is not None:
                chk.blockSignals(True); chk.setChecked(False); chk.blockSignals(False)
            self._on_log("ℹ 没选定分组，「采集完自动导入」已取消")
            return
        self._auto_import_group = dlg.chosen_group
        flag = self._flag_from_group(self._auto_import_group)
        self._on_log(
            f"🤖 已设置采集完自动导入：分组「{self._auto_import_group.get('groupName') or ''}」"
            f"（{'国外' if flag else '国内'}客户库），采集完成后自动上传并在日志提示结果")

    def _revert_auto_import(self, msg: str):
        self._auto_import_group = None
        chk = getattr(self, "auto_import_check", None)
        if chk is not None:
            chk.blockSignals(True); chk.setChecked(False); chk.blockSignals(False)
        self._on_log(f"⚠ {msg}")
        QMessageBox.warning(self, "无法开启自动导入", msg)

    def _maybe_auto_import(self, reason: str):
        chk = getattr(self, "auto_import_check", None)
        if chk is None or not chk.isChecked():
            return
        if self._yzd_importing:
            return
        group = self._auto_import_group or {}
        if not group.get("id"):
            self._on_log("⚠ 已勾选自动导入但没拿到分组，本次不自动上传；可手动点「导入云展客户库」")
            return
        if not self.crawler.snapshot_results():
            self._on_log("ℹ 本次没采到数据，自动导入已跳过")
            return
        tid = self._current_task_id()
        if not tid:
            self._on_log("⚠ 本次采集没登记到任务索引（会话目录未创建），自动导入已跳过")
            return
        tasks = [t for t in _tasks.load_tasks() if t.get("id") == tid] or [
            {"id": tid, "name": f"本次采集 {tid}", "folder": str(self.session_dir)}]
        self._on_log(
            f"🤖 {reason}，按勾选自动导入云展客户库：分组「{group.get('groupName') or ''}」")
        self._start_yzd_import(group, tasks)

    def _build_jobs(self, tasks: list) -> list:
        """当前会话直接吃内存快照（避开"刚采完还没落盘"的空档），历史任务才回读文件。"""
        jobs = []
        tid = self._current_task_id()
        rows = self.crawler.snapshot_results()
        for t in tasks or []:
            name = t.get("name") or t.get("id") or "未命名任务"
            if t.get("id") and t.get("id") == tid and rows:
                jobs.append({"id": t.get("id"), "name": name, "records": rows, "raw_file": ""})
                continue
            raw_file = t.get("raw_file") or ""
            if not raw_file and t.get("folder"):
                raw_file = os.path.join(str(t.get("folder")), _tasks.RAW_FILE_NAME)
            jobs.append({"id": t.get("id") or "", "name": name,
                         "records": None, "raw_file": raw_file})
        return jobs

    def _start_yzd_import(self, group: dict, tasks):
        if self._yzd_importing:
            return
        if tasks is None:
            latest = _tasks.latest_task()
            if not latest:
                QMessageBox.warning(self, "提示", "暂无采集任务：请先在主界面完成一次采集")
                return
            jobs = self._build_jobs([latest])
        else:
            jobs = self._build_jobs(tasks)
        if not jobs:
            QMessageBox.warning(self, "提示", "没有可导入的采集任务")
            return
        empty = [j for j in jobs if not j["raw_file"] and not (j["records"] or [])]
        if len(empty) == len(jobs):
            QMessageBox.warning(self, "提示", "选中的采集任务里没有数据")
            return
        import tempfile
        try:
            tmp_dir = tempfile.mkdtemp(prefix="yzd_import_")
        except Exception as e:
            QMessageBox.critical(self, "导入失败", f"无法创建临时上传目录：{e}")
            return
        self._yzd_jobs = jobs
        self._yzd_importing = True
        group_name = group.get("groupName") or ""
        self._yzd_group_name = group_name
        flag = self._flag_from_group(group)
        self._on_log(f"⬆ 开始导入云展客户库：分组「{group_name}」（{'国外' if flag else '国内'}）"
                     f"· 共 {len(jobs)} 个任务")
        # 锁界面：父窗禁用后子控件有效状态恒为禁用，解锁后回到各自最新状态
        self.setEnabled(False)
        prog = YzdProgressDialog(None, f"正在导入云展客户库：{group_name}")
        self._yzd_progress = prog
        prog.show()
        worker = YzdImportWorker(self.crawler, self._yzd_phone, jobs,
                                 group.get("id"), flag, tmp_dir)
        worker.logged.connect(self._on_log)
        worker.progress.connect(self._on_yzd_import_progress)
        worker.finished_all.connect(self._on_yzd_import_finished)
        self._yzd_import_worker = worker
        worker.start()

    @staticmethod
    def _flag_from_group(group: dict) -> int:
        value = group.get("foreignFlag")
        try:
            return int(value) if value is not None else 0
        except (TypeError, ValueError):
            return 0

    def _on_yzd_import_progress(self, done: int, total: int, text: str):
        if self._yzd_progress is not None:
            self._yzd_progress.set_text(text, f"已完成 {done}/{total} 个任务。导入期间请勿关闭程序。")

    def _unlock_yzd_import(self):
        self._yzd_importing = False
        self.setEnabled(True)
        if self._yzd_progress is not None:
            try:
                self._yzd_progress.allow_close()
                self._yzd_progress.close()
            except Exception:
                pass
            self._yzd_progress = None
        self.activateWindow()

    def _on_yzd_import_finished(self, results: list, fatal: str):
        if fatal:
            self._on_log(f"❌ 导入异常：{fatal}")
            self._unlock_yzd_import()
            QMessageBox.critical(self, "导入失败", fatal)
            return
        lines = []
        saved_total = 0
        for idx, item in enumerate(results or []):
            name, imp, sav, fail, err = item
            saved_total += sav
            if err:
                lines.append(f"· {name}：导入失败（{err}）")
                self._on_log(f"❌ {name}：{err}")
            else:
                lines.append(f"· {name}：上传 {imp} 条 / 入库成功 {sav} 条 / 失败 {fail} 条")
                self._on_log(f"✅ {name}：上传 {imp} / 入库 {sav} / 失败 {fail}")
            if idx < len(self._yzd_jobs):
                job = self._yzd_jobs[idx]
                if job.get("id"):
                    try:
                        _tasks.update_task(
                            job["id"],
                            yzd_import_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                            yzd_group=self._yzd_group_name, yzd_saved=sav)
                    except Exception as e:
                        self._on_log(f"⚠ 任务导入结果回写失败：{e}")
        self._unlock_yzd_import()   # 先解锁再弹结果
        self._on_log("ℹ 云展接口不做重复校验，同一批数据重复导入会在客户库里产生多条记录")
        summary = "\n".join(lines) if lines else "没有任何任务被处理"
        if saved_total == 0:
            QMessageBox.warning(self, "导入结果",
                                f"未入库任何客户：\n{summary}\n\n请查看运行日志里的后端返回。")
        else:
            QMessageBox.information(self, "导入结果", f"共入库 {saved_total} 条客户：\n\n{summary}")
        import shutil
        worker, self._yzd_import_worker = self._yzd_import_worker, None
        if worker is not None:
            try:
                shutil.rmtree(worker.tmp_dir, ignore_errors=True)
            except Exception:
                pass

    def closeEvent(self, ev):
        if self._yzd_importing:
            # 导入中途退程序会在云展那边留下"传了一半"的数据，且线程被强杀只拿到半个结果
            QMessageBox.information(self, "请等一下",
                                    "正在导入云展客户库，请等导入完成后再关闭程序。")
            ev.ignore()
            return
        try:
            self.crawler.stop_crawl()
            if self.autosave and self.autosave.isRunning():
                self.autosave.wait(3000)
            self.crawler.close_browser()
        except Exception:
            pass
        super().closeEvent(ev)


def main():
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
