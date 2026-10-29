# -*- coding: utf-8 -*-
"""待采集公司名单：企知道补采通道的中转站（b2bcollect 六通道体系的第六路）。

两个入口：
  ① 任务开始时用户主动上传的公司名单文件（seed，识别出的公司名全部入列）；
  ② 关键词采集被清洗口径剔掉的「废数据」（手机号/固话/邮箱/微信号四项全空）——
     这些公司名白白丢掉太可惜，交给企知道按名补联系方式，变废为宝。
一个出口：企知道通道逐家 pop 公司名补采（名单通道，不参与关键词并行）。

三个台账（防死循环、防重复采，这是本模块存在的全部理由）：
  pending   等待补采的公司（同步落盘到任务文件夹「待采集公司名单.xlsx」，用户可见可改）；
  tried     已经送过企知道的公司——**无论查没查到都不再入列**。否则「查到但仍无
            联系方式 → 又被清洗成废数据 → 再入列」会无限循环烧请求；
  collected 已采到联系方式的公司（pop 时命中即跳过，关键词通道后面补到了也不重复补）。

并发：采集线程（写废数据）/ 企知道线程（pop）/ 保存线程（flush）三方共用，全锁保护；
落盘一律 tmp+os.replace 原子替换（进程被杀/磁盘满不会把上一份好名单写成半个坏文件）。
"""
import json
import os
import threading
import time
from collections import deque
from pathlib import Path

LIST_FILENAME = "待采集公司名单.xlsx"
LEDGER_FILENAME = "pending_ledger.json"


def extract_company_names(path, log=print):
    """从用户选择的名单文件里识别公司名（xlsx/csv/txt）。
    列定位三级判据（对齐企知道工具的选列思路，但不弹窗——后台通道要的是稳不是交互）：
    ① 表头含「公司/企业/单位名称」语义直接命中；② 否则按列内容抽样打分取最高；
    ③ 单列文件（txt/单列 csv）整列当公司名。识别后去重保序、剔非公司串。"""
    import customer_library_cleaner as _cleaner
    p = Path(path)
    ext = p.suffix.lower()
    grid = []
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True)   # 不用 read_only：实测会静默截行
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            grid.append(["" if c is None else str(c) for c in row])
        wb.close()
    elif ext == ".csv":
        import csv
        raw = p.read_bytes()
        text = None
        for enc in ("utf-8-sig", "gbk", "utf-8"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            text = raw.decode("utf-8", errors="ignore")
        for row in csv.reader(text.splitlines()):
            grid.append(list(row))
    else:  # txt 等：一行一个公司名
        for line in p.read_text(encoding="utf-8", errors="ignore").splitlines():
            grid.append([line])

    if not grid:
        return []
    width = max(len(r) for r in grid)
    grid = [r + [""] * (width - len(r)) for r in grid]

    def col_score(ci, header=""):
        vals = [r[ci].strip() for r in grid[1:] if r[ci].strip()] if header else \
               [r[ci].strip() for r in grid if r[ci].strip()]
        if not vals:
            return 0.0
        sample = vals[:60]
        hit = sum(1 for v in sample if _cleaner._is_company(v))
        return hit / len(sample)

    header = grid[0]
    best_ci, best, best_sem = -1, 0.0, False
    for ci in range(width):
        h = header[ci].strip() if ci < len(header) else ""
        sem = 0.45 if any(k in h for k in ("公司", "企业", "单位名称", "商户")) else 0.0
        score = sem + 0.55 * col_score(ci, header=h)
        if score > best:
            best, best_ci, best_sem = score, ci, bool(sem)
    if best_ci < 0 or best < 0.35:
        return []
    # 是否跳过首行：只有命中列的首格确实是表头才跳。
    # ① 语义命中但首格本身就是公司名（如「XX有限公司」含“公司”二字）→ 是数据不是表头；
    # ② 无表头的 xlsx/txt 首行也是数据，跳了就丢第一家。
    first_val = (header[best_ci] or "").strip()
    header_literal = first_val in ("公司名", "公司名称", "企业名称", "单位名称", "名称",
                                  "公司", "企业", "商户名称", "客户名称")
    skip_first = (best_sem and not _cleaner._EXT_COMPANY_TAIL_RE.search(first_val)) \
        or header_literal
    body = grid[1:] if skip_first else grid
    names, seen = [], set()
    for r in body:
        v = (r[best_ci] or "").strip()
        if not v or v == "None" or v in seen:
            continue
        if not _cleaner._is_company(v):
            continue
        seen.add(v)
        names.append(v)
    return names


class PendingCompanyList:
    """任务级待采集名单。bind() 之前所有写操作只记内存（允许先 seed 后绑定目录）。"""

    def __init__(self):
        self._lock = threading.RLock()
        self._pending = deque()
        self._in_pending = set()
        self._tried = set()
        self._collected = set()
        self._dir = None
        self._collected_probe = None     # 可选回调：关键词通道是否已采到该公司
        self._dirty = False
        self._unbound_seed = []          # bind 之前进来的公司名，bind 时落盘

    # ── 绑定任务文件夹 ────────────────────────────────────────
    def bind(self, session_dir):
        with self._lock:
            self._dir = Path(session_dir)
            self._dir.mkdir(parents=True, exist_ok=True)
            self._load_ledger()
            self._load_xlsx()
            if self._unbound_seed:
                added = self._extend(self._unbound_seed, "seed")
                self._unbound_seed = []
                self.flush()
                return added
            return 0

    @property
    def xlsx_path(self):
        return self._dir / LIST_FILENAME if self._dir else None

    @property
    def is_bound(self):
        """是否已绑定任务文件夹（绑定后废数据才能落盘到待采名单）。"""
        with self._lock:
            return self._dir is not None

    @property
    def ledger_path(self):
        return self._dir / LEDGER_FILENAME if self._dir else None

    def set_collected_probe(self, fn):
        with self._lock:
            self._collected_probe = fn

    # ── 入列 ─────────────────────────────────────────────────
    def seed(self, names, source="名单文件"):
        """用户名单整份入列（任务开始时调用）。"""
        with self._lock:
            if self._dir is None:
                self._unbound_seed.extend(names)
                return len(names)
            n = self._extend(names, source)
            self.flush()
            return n

    def add_waste(self, names):
        """清洗剔掉的废数据公司名入列（自动保存线程调用）。tried/collected 里的不收。"""
        with self._lock:
            n = self._extend(names, "waste")
            if n:
                self.flush()
            return n

    def _extend(self, names, source):
        added = 0
        for raw in names or []:
            name = str(raw or "").strip()
            if not name or name in self._in_pending or name in self._tried \
                    or name in self._collected:
                continue
            self._in_pending.add(name)
            self._pending.append(name)
            added += 1
        if added:
            self._dirty = True
        return added

    # ── 出列（企知道通道）────────────────────────────────────
    def pop(self):
        """取下一家待补采公司。命中 collected/探针 的直接跳过（不浪费企知道请求）。
        无待采返回 None。pop 即记 tried：查不到就丢弃，绝不回列（防死循环）。"""
        with self._lock:
            while self._pending:
                name = self._pending.popleft()
                self._in_pending.discard(name)
                self._dirty = True
                if name in self._collected:
                    self._tried.add(name)
                    continue
                if self._collected_probe:
                    try:
                        if self._collected_probe(name):
                            self._collected.add(name)
                            self._tried.add(name)
                            continue
                    except Exception:
                        pass
                self._tried.add(name)
                return name
            return None

    def mark_result(self, name, found):
        """企知道补采回执：查到→记 collected（其结果进主结果集参与清洗）；
        查不到→只留在 tried（直接丢弃，用户原话「企知道都查不出的数据直接丢弃」）。"""
        with self._lock:
            self._tried.add(name)
            if found:
                self._collected.add(name)
            self._flush_ledger()

    def requeue(self, name):
        """把「这一家根本没真正查成」的公司放回队首重排（区别于 mark_result(False) 的丢弃）。
        适用：通道就绪失败 / 被拦截 break / 未登录空壳——都不是「企知道查不到」，
        不该记 tried 白白丢掉。与 pop 的 tried 记账解耦：从 tried 移除后重新入列。
        已 collected 的公司不 requeue（已采到，无需重来）。"""
        with self._lock:
            name = str(name or "").strip()
            if not name or name in self._in_pending or name in self._collected:
                return False
            self._tried.discard(name)
            self._pending.appendleft(name)
            self._in_pending.add(name)
            self._dirty = True
            self.flush()
            return True

    def has_work(self):
        with self._lock:
            return bool(self._pending)

    def counts(self):
        with self._lock:
            return {"pending": len(self._pending), "tried": len(self._tried),
                    "collected": len(self._collected)}

    def snapshot(self):
        with self._lock:
            return list(self._pending)

    # ── 落盘 / 读盘 ──────────────────────────────────────────
    def flush(self):
        """名单 xlsx + 台账 json 原子落盘。保存线程收尾/停止时调用。"""
        with self._lock:
            if self._dir is None:
                return False
            self._flush_xlsx()
            self._flush_ledger()
            self._dirty = False
            return True

    def _flush_xlsx(self):
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "待采集公司名单"
            ws.append(["公司名"])
            for name in self._pending:
                ws.append([name])
            tmp = str(self.xlsx_path) + ".tmp"
            wb.save(tmp)
            os.replace(tmp, str(self.xlsx_path))
        except PermissionError:
            pass    # 名单被 Excel 打开着：本轮跳过，下轮补上（不阻断采集）
        except Exception:
            pass

    def _flush_ledger(self):
        try:
            payload = {"tried": sorted(self._tried), "collected": sorted(self._collected),
                       "at": time.strftime("%Y-%m-%d %H:%M:%S")}
            tmp = str(self.ledger_path) + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False)
            os.replace(tmp, str(self.ledger_path))
        except Exception:
            pass

    def _load_ledger(self):
        try:
            data = json.loads(self.ledger_path.read_text(encoding="utf-8"))
            self._tried = set(data.get("tried", []))
            self._collected = set(data.get("collected", []))
        except Exception:
            self._tried, self._collected = set(), set()

    def _load_xlsx(self):
        """续跑同一任务文件夹时把名单读回（用户也可能手工往表里加公司）。"""
        if not self.xlsx_path or not self.xlsx_path.exists():
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(self.xlsx_path), data_only=True)
            ws = wb.active
            names = []
            for row in ws.iter_rows(values_only=True):
                v = "" if not row or row[0] is None else str(row[0]).strip()
                if v and v != "公司名":
                    names.append(v)
            wb.close()
            self._extend(names, "reload")
        except Exception:
            pass
