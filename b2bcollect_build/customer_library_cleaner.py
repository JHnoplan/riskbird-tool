"""清洗后客户库入库数据生成模块。

对齐《客户库数据导入模板.xlsx》的列结构，把三个平台采集的原始结果做
schema mapping + 数据质量校验 + 跨平台合并去重，输出可直接入库的 Excel。

字段映射（模板列 → 采集字段，按优先级）：
    客户名称 ← 法定代表人（真实联系人）?? 公司名
    企业名称 ← 公司名
    岗位     ← 客户名称取自法定代表人时填「法定代表人」，否则留空
    手机号   ← 电话（正则校验 ^1[3-9]\\\\d{9}$）
    固话     ← 电话（区号-号码 / 分机号归一化）
    邮箱     ← 邮箱（正则校验 + 去重）
    微信号   ← 源数据无此字段，留空
    数据来源 ← 数据来源（多平台命中时逗号拼接）
    外呼备注/备注 ← 状态异常说明 + 官网 + 地址摘要

清洗规则：
    1. 剔除「未找到」「失败」「额度用尽待重试」记录；
    2. 同一公司跨平台合并为一条（字段取首个非空值，电话/邮箱取并集），
       避免入库侧重复；跨站岗位回捞：本站没岗位、别站有证据一致的岗位→补上；
       证据不一致→写进备注不丢；无明写岗位时按邮箱前缀/部门/公司名尾巴
       智能匹配出「岗位族」也挂进备注（字段名不叫岗位也认得：职位/职务/jobtitle…）；
    3. 电话、邮箱拆分为手机号 / 固话 / 邮箱三个独立字段并逐个校验格式，
       电话只分类不丢弃：识别不出标准格式的号码按原样落到固话；
    4. 文本归一化（去零宽字符与首尾空白、全角标点转半角）；
    5. 手机号/固话/邮箱全量保留，多个用 , 分隔，清洗环节不做条数截断
       （历史上每家只留 3 个号码，有用的手机号就是这么被洗没的）；
       上传云展时若超接口长度上限，由 yzd_client.clamp_fields 按整条号码收拢。
    6. 模板表头按列名动态定位，不写死列序；模板缺失时回退内置表头。
"""

import os
import re

DEFAULT_OUTPUT_NAME = "清洗后客户库入库数据.xlsx"

# 模板表头兜底（A-I 外呼块 / J-R 备注块），运行时优先读模板真实表头
FALLBACK_HEADER = [
    "客户名称", "企业名称", "岗位", "手机号", "固话", "邮箱", "微信号", "数据来源", "外呼备注",
    "客户名称", "企业名称", "岗位", "手机号", "固话", "邮箱", "微信号", "数据来源", "备注",
]

# 采集字段 → 模板列的取值优先级
# 「联系人/岗位/微信号/备注」只在外部数据智能清洗时才会出现（三个平台的采集结果里
# 没有这几列），采集链路的输出保持原样不受影响。
SOURCE_FIELDS = {
    "公司名": ("公司名", "企业名称", "公司名称"),
    "法定代表人": ("法定代表人", "法人代表", "法人", "负责人"),
    "联系人": ("联系人", "客户名称", "联系人姓名"),
    "岗位": ("岗位", "职位", "职务"),
    "电话": ("电话", "联系电话", "更多电话"),
    "邮箱": ("邮箱", "邮箱地址", "电子邮箱"),
    "微信号": ("微信号", "微信"),
    "官网": ("官网", "网址", "企业网址"),
    "地址": ("地址", "注册地址", "通信地址"),
    "状态": ("状态",),
    "数据来源": ("数据来源",),
    "备注": ("备注",),
}

# 岗位智能匹配助手（与采集层同一口径）：部门词/邮箱前缀/公司名尾巴等旁证归集成岗位族。
# 注意：cdp_base 顶部 import 了本模块（循环依赖），**不能在模块顶层 import cdp_base**，
# 否则本文件先加载时 cdp_base 未就绪、名字被降级。改为函数内惰性 import（调用时两边都已就绪）。
_FAMILY_HELPERS = None


def _load_family_helpers():
    global _FAMILY_HELPERS
    if _FAMILY_HELPERS is not None:
        return _FAMILY_HELPERS
    try:
        import cdp_base
        _FAMILY_HELPERS = (cdp_base._family_of, cdp_base._BARE_TITLE_RE,
                           cdp_base.dept_from_email, cdp_base.dept_from_company,
                           cdp_base.dept_only)
    except Exception:      # pragma: no cover - 独立分发环境降级为不标注
        _FAMILY_HELPERS = False
    return _FAMILY_HELPERS


_TEMPLATE_SEARCH_NAMES = ("客户库数据导入模板", "客户库导入模板", "客户库数据模 板")

_MOBILE_RE = re.compile(r"^1[3-9]\d{9}$")
# 固话：区号(3-4位) + 号码(7-8位) + 可选分机；也接受裸 7-8 位号码
_LANDLINE_RE = re.compile(r"^(?:0(\d{2,3}))?[-\s]?(\d{7,8})(?:[-\s转.]*(\d{1,6}))?$")
# 400/800 客服热线：不能按区号拆分，保留原样
_HOTLINE_RE = re.compile(r"^(0?)(400|800)[-\s]?(\d{3})[-\s]?(\d{4})$")
# 分组书写的手机号（139 0806 0502 / 139-0806-0502）：不先合回 11 位会被分隔符切碎后整段丢掉
_MOBILE_REGROUP_RE = re.compile(r"(?<!\d)(1[3-9]\d)[-\s]?(\d{4})[-\s]?(\d{4})(?!\d)")
# 国际码前缀：只在去掉后确实是个合法手机号时才动，避免误伤 86 开头的固话
_COUNTRY_CODE_RE = re.compile(r"^\+?0{0,2}86-?")
# 「前导 0 / 区号 + 手机号 + 可选分机」的连写：013900139000、075513900139000、13900139000转801
# 必须整串锚定匹配，绝不能在数字串里 search 11 位子串：固话带分机的连写里正好藏着
# 1[3-9] 开头的 11 位，一 search 就把正常固话判成手机并把号打坏。
_MOBILE_INLINE_RE = re.compile(
    r"^(?:\+?0{0,2}86|0\d{2,3}|0)?[-\s]?(1[3-9]\d{9})(?:[-\s转.、]*\d{1,6})?$")
# 「区号␣号码」只隔一个空格：不先连起来会被分隔符拆两段，区号整段丢掉（0755 88881111）
_AREA_SPACE_RE = re.compile(r"(?<!\d)(0\d{2,3})[\s]+(\d{7,8})(?!\d)")
# 公司名尾部杂质（采集页标题带分页版次：实测「…材料有限公司3.」）：只剥中文后的结尾纯数字/句点
_NAME_TAIL_JUNK_RE = re.compile(r"(?<=[\u4e00-\u9fa5\)）])\s*[0-9０-９]+\s*[.。·]*\s*$")
# 平台给公司名加的「（XX厂家）」式 SEO 尾巴（与 cdp_base.strip_name_marketing_tail 同口径，
# 清洗层再挡一道：已入库的历史脏名不必重采也能在清洗时洗回规范名）。
# 只剥位于结尾且含厂家/制造商等词的括号段，「东方鑫鸿（北京）科技有限公司」不受影响。
_NAME_MARKETING_TAIL_RE = re.compile(
    r"[（(][^（）()]{2,24}(?:厂家|制造商|生产厂家|供应商|批发商|价格|报价)[）)]\s*$")
_EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")
_EMAIL_EXTRACT_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
_SPLIT_RE = re.compile(r"[;；,，、\s/|]+")
_PLACEHOLDER_RE = re.compile(r"^(-{1,}|暂无|未知|无|待补充|NULL|null|N/?A|见详情)$")
_SEP_CHARS = "/／|｜,，、;；-－–—\\s"

# 零宽字符与全角标点
_ZERO_WIDTH_RE = re.compile(r"[\u200b-\u200f\ufeff\u00ad]")
# Excel/openpyxl 不接受的控制字符与非法代理对：采集到的字段可能携带，
# 不在归一化阶段剥掉，写清洗表时会抛 IllegalCharacterError 炸掉整份导出。
_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f\ud800-\udfff\ufffe\uffff]")
_FULLWIDTH_MAP = {ord("："): ":", ord("；"): ";", ord("，"): ",", ord("、"): "、",
                  ord("（"): "(", ord("）"): ")", ord("　"): " "}


def normalize_text(value) -> str:
    """文本归一化：非字符串转字符串、去零宽字符、压缩连续空白。"""
    if value is None:
        return ""
    text = str(value)
    text = _ILLEGAL_RE.sub("", text)
    text = _ZERO_WIDTH_RE.sub("", text)
    text = text.translate(_FULLWIDTH_MAP)
    text = re.sub(r"[ \t\u3000]+", " ", text)
    return text.strip()


def _get_field(record: dict, key: str) -> str:
    """按 SOURCE_FIELDS 优先级从记录里取第一个非空字段值。"""
    for name in SOURCE_FIELDS.get(key, (key,)):
        val = normalize_text(record.get(name, ""))
        if val and not _PLACEHOLDER_RE.match(val):
            return val
    return ""


def _clean_legal_rep(value: str) -> str:
    """法定代表人可能带「该方已注销 / 多自然人分隔」等脏值，取首个有效人名。"""
    if not value:
        return ""
    parts = [p.strip() for p in re.split(r"[/／|｜,，、;；]", value) if p.strip()]
    for p in parts:
        p = re.sub(r"\s+", "", p)
        if 2 <= len(p) <= 20 and not _PLACEHOLDER_RE.match(p):
            return p
    return re.sub(r"\s+", "", parts[0]) if parts else ""


# 重复称呼（实测八方资源网 contact.aspx「姓 名：徐先生 先生」：商家把「徐先生」当
# 名字填进姓名框、平台又在后面自动补一个「先生」）：不合并就入库外呼叫不出口的名字。
# 只缩同词重复（反引用），「徐总经理」不会被拆成「徐总」。
_DUP_HONORIFIC_RE = re.compile(r"(先生|女士|小姐|总监|经理|厂长|老板|主任|工程师|总)(?:\s*\1)+")
# 人名位不该出现的两类错位值：纯号码与公司名（采集层已拦一道，这里是入库前最后一道闸，
# 并兼顾「清洗数据」按钮直读外部表格的链路——那里没有人帮忙过滤）。
_CONTACT_NUMBER_RE = re.compile(r"^[\d\s\-+()转号]+$")
_CONTACT_NOT_NAME_RE = re.compile(r"(公司|企业|集团|有限|商行|商贸|工作室|营业部|店铺)")
# 地址尾部粘连的页面按钮词（与 cdp_base._ADDR_NOISE_RE 同口径的子集）
_ADDR_BUTTON_RE = re.compile(r"(公司主页|公司网址|访问商铺|查看更多|相关产品|产品索引|收藏夹|举报|扫一扫|扫码)")


def _clean_contact(value: str) -> str:
    """联系人/法定代表人入库前校验：合并重复称呼、剪粘在人名后的号码尾巴，
    纯号码（实测爱采购 cardInfo 直接填「18962536677」）与公司名（实测 MIC 把
    「沈阳微特通用技术开发有限公司」写进联系人列）一律作废，让上层退回公司名：
    客户名称列挂个号码或公司全称，入库后外呼名单直接不可用。"""
    s = re.sub(r"\s+", "", normalize_text(value))
    if not s:
        return ""
    s = _DUP_HONORIFIC_RE.sub(r"\1", s)
    if _CONTACT_NUMBER_RE.match(s):
        return ""
    mnum = re.search(r"\d{7,}", s)
    if mnum and mnum.start() >= 2:
        s = s[:mnum.start()]
    if _CONTACT_NOT_NAME_RE.search(s):
        return ""
    return s


def split_phones(raw: str):
    """把混合电话串拆成 (手机号列表, 固话列表)。

    只做分类、格式化与去重；唯一的丢弃是全 0 假号（见下方占位判定）。
    标准格式都认不出的号码也原样归到固话，
    否则一个有用的手机号会被静默洗没（实测：旧版剥掉分机字「转」后，
    00755-27908031转1 被区号贪婪误配重排成 0075-52790803转1，号码位置被打坏）。
    """
    mobiles, landlines = [], []
    if not raw:
        return mobiles, landlines
    # 先把空格/横线分隔的 400/800 热线归一为连字符形式，避免被分隔符拆开
    text = re.sub(r"\b([48]00)[-\s]?(\d{3})[-\s]?(\d{4})\b", r"\1-\2-\3", normalize_text(raw))
    # 分组写的手机号（139 0806 0502）先合回 11 位，否则会被分隔符切成三段全丢
    text = _MOBILE_REGROUP_RE.sub(r"\1\2\3", text)
    # 「区号 号码」空格连写（0755 88881111）补成横线：分隔符会把空格当断点，区号就此丢掉
    text = _AREA_SPACE_RE.sub(r"\1-\2", text)
    for token in _SPLIT_RE.split(text):
        # 分机字「转」必须留着：删了以后 `0519-8389830转6049` 就变成一串 15 位连写，
        # 固话正则的 (\d{7,8}) 贪婪吃掉前 8 位，洗出 `0519-83898306转049` 这种打坏的号码。
        token = re.sub(r"[^\d+\-转]", "", normalize_text(token))
        token = re.sub(r"转+(?!\d)", "", token)      # 只有「转」却没有分机号的空尾直接去掉
        if not token or not token.strip("-转"):
            continue
        digits = token.replace("-", "")
        # 商家随手灌的一串 0（实测顺企网「0515 80000000 / 0515-82807298」：区号真、
        # 号码主体灌 8 个 0）：与采集层 is_placeholder_number 同口径阈值 7 连 0，
        # 既不把空号当联系方式入库，也不误杀 18710000805 这类真实号段。
        # 历史脏数据不必重采，清洗时就能挡掉。
        if re.search(r"0{7,}", digits):
            continue
        if _MOBILE_RE.match(digits):
            if digits not in mobiles:
                mobiles.append(digits)
            continue
        # 带国家码的手机号（+8613800138000 / 0086-138-0013-8000）脱码后按手机归位
        stripped = _COUNTRY_CODE_RE.sub("", digits)
        if _MOBILE_RE.match(stripped):
            if stripped not in mobiles:
                mobiles.append(stripped)
            continue
        # 手机号优先：带前导 0/区号/分机号的连写得赶在固话贪婪匹配前整串锚定认出来，
        # 否则会被固话正则切成「13900139转000801」这种既不是手机也不是固话的碎片。
        mi = _MOBILE_INLINE_RE.match(digits)
        if mi:
            if mi.group(1) not in mobiles:
                mobiles.append(mi.group(1))
            continue
        hm = _HOTLINE_RE.match(token) or _HOTLINE_RE.match(digits)
        if hm:
            hotline = f"{hm.group(2)}-{hm.group(3)}-{hm.group(4)}"
            if hotline not in landlines:
                landlines.append(hotline)
            continue
        m = _LANDLINE_RE.match(token) or _LANDLINE_RE.match(digits)
        if m:
            area, number, ext = m.group(1), m.group(2), m.group(3)
            if len(number) < 7:
                continue
            formatted = f"0{area}-{number}" if area else number
            if ext:
                formatted += f"转{ext}"
            if formatted not in landlines:
                landlines.append(formatted)
            continue
        # 兜底：格式不规范但确实是号码的（异常区号如 00755、国外号等）原样进固话，
        # 绝不重排数字、绝不静默丢弃。
        plain = re.sub(r"\D", "", digits)
        if 7 <= len(plain) <= 15 and token not in landlines:
            landlines.append(token)
    return mobiles, landlines


def split_emails(raw: str):
    """邮箱拆分去重 + 格式校验，统一小写域名部分。"""
    result = []
    if not raw:
        return result
    for token in _SPLIT_RE.split(raw):
        email = normalize_text(token).strip("<>\"'()").lower()
        if not email:
            continue
        # 兼容 "a@b.com(公司)" / "a@b.com 官网" 之类粘连：提取所有合法邮箱
        for found in _EMAIL_EXTRACT_RE.findall(email):
            if found not in result:
                result.append(found)
    return result


def is_valid_record(record: dict) -> bool:
    """有效性判定：公司名非空，且状态不含 未找到/失败/异常/无结果/额度受限。"""
    name = _get_field(record, "公司名")
    if not name:
        return False
    status = _get_field(record, "状态")
    for kw in ("未找到", "失败", "异常", "无结果", "不存在", "未匹配", "额度", "待重试"):
        if kw in status:
            return False
    return True


def _first_non_blank(values):
    for v in values:
        if v:
            return v
    return ""


def merge_records(results: list):
    """跨平台按公司名合并去重：返回 (合并后的记录字典列表, 重复条数)。"""
    merged = {}
    order = []
    dup = 0
    for record in results:
        if not is_valid_record(record):
            continue
        name = _normalize_key(_get_field(record, "公司名"))
        if not name:
            continue
        source = _get_field(record, "数据来源") or "未知平台"
        if name in merged:
            dup += 1
            item = merged[name]
            if source not in item["sources"]:
                item["sources"].append(source)
            had_contact = bool(item["fields"].get("联系人"))
            for key in list(SOURCE_FIELDS):
                # 岗位不参与「首个非空」泛填：它必须锁联系人同源（见下），
                # 否则八方的「经理」会串到顺企网那个名字上。
                if key == "岗位":
                    continue
                if not item["fields"].get(key):
                    item["fields"][key] = _get_field(record, key)
            # 岗位必须跟联系人同源：各站可能给不同联系人（八方的「魏经理」+ 顺企网
            # 的「陈生」），若独立取首非空会把「岗位:经理」挂在另一个名字上。
            # 但同源约束只约束「岗位列」本身：别站的岗位不再静默丢弃——
            # 证据一致（同一联系人且主岗位还空）直接回捞进列，否则挂备注供外呼参考。
            other_pos = normalize_text(_get_field(record, "岗位"))
            kept_contact = normalize_text(item["fields"].get("联系人") or "")
            cur_contact = normalize_text(_get_field(record, "联系人") or "")
            main_pos = normalize_text(item["fields"].get("岗位") or "")
            if not had_contact and kept_contact:
                # 首个带联系人的源：岗位以它为准（覆盖更早无联系人源的串值）
                item["fields"]["岗位"] = other_pos
            elif other_pos and other_pos != main_pos:
                if kept_contact and kept_contact == cur_contact and not main_pos:
                    item["fields"]["岗位"] = other_pos
                else:
                    tag = (f"{other_pos}({cur_contact}@{source})"
                           if cur_contact and cur_contact != kept_contact
                           else f"{other_pos}({source})")
                    prev = item["fields"].get("岗位(跨站)", "")
                    item["fields"]["岗位(跨站)"] = ((prev + "、" + tag) if prev else tag)[:60]
            # 电话/邮箱取并集而非首个非空
            for key in ("电话", "邮箱"):
                more = _get_field(record, key)
                if more and key in item["fields"]:
                    base = item["fields"][key]
                    item["fields"][key] = (base + ";" + more) if base else more
            # 跨站联系人不再静默丢弃：主联系人进「客户名称」，其余归集到备注
            # （实测同一公司八方「陈」/11467「黄双华」/gys「黄先生」只留一个，
            # 外呼能用的名字就这么洗掉了）。
            other = _get_field(record, "联系人")
            if other:
                base = item["fields"].get("联系人") or ""
                others = item["fields"].setdefault("其他联系人", "")
                pool = [base] + [x for x in re.split(r"[,、]", others) if x]
                if other not in pool:
                    pool.append(other)
                primary, tail = pool[0], pool[1:]
                item["fields"]["联系人"] = primary
                if tail:
                    item["fields"]["其他联系人"] = "、".join(tail)[:60]
        else:
            merged[name] = {
                "sources": [source],
                "fields": {key: _get_field(record, key) for key in SOURCE_FIELDS},
            }
            order.append(name)
    return [merged[n] for n in order], dup


def _normalize_key(name: str) -> str:
    """公司名匹配键：去空白与全角括号转半角，避免同名不同格式重复入库。"""
    key = re.sub(r"\s+", "", name)
    return key.replace("（", "(").replace("）", ")").upper()


def _join_limited(values, limit=None, sep=","):
    """号码全量保留不截断（历史上每家只留 3 个，有用的手机号就是这么被洗没的）；
    上传云展超接口长度上限时由 yzd_client.clamp_fields 按整条号码收拢。"""
    vals = list(values)
    if limit:
        vals = vals[:limit]
    return sep.join(vals)


def _remark(fields: dict, contact_ok: bool = False, fam: str = ""):
    """备注：状态说明 + 官网 + 地址摘要（+ 外部数据自带的备注列），供外呼人员参考。

    contact_ok：跨平台合并后已有电话/邮箱时，某一路源采集时写的「无电话(未公开)/
    无联系方式」已不成立，不再写进备注误导外呼（实测合并记录带电话仍残留此注）。"""
    pieces = []
    status = fields.get("状态", "")
    if contact_ok and status in ("无电话(未公开)", "无联系方式"):
        status = ""
    if status and status not in ("存续", "在业", "开业", "正常", "成功"):
        # 「无电话(未公开)」「详情页导航失败」是**采集状态**，写成「经营状态」会让外呼
        # 人员误以为公司已注销（实测清洗后表里供应商网 64 行就这么挂在备注上）。
        label = "采集状态" if status.startswith(("无电话", "无联系方式", "类目不符", "详情页", "解析")) \
            else "经营状态"
        pieces.append(f"{label}:{status}")
    site = fields.get("官网", "")
    if site:
        pieces.append(f"官网:{site}")
    addr = fields.get("地址", "")
    if addr:
        # MIC 采集页把地址后面的按钮词一并吃了进来（实测 265/286 行结尾挂着
        # 「公司主页」），备注里读起来不知地址到哪里为止。与采集层 clean_address 同口径。
        addr = _ADDR_BUTTON_RE.split(addr)[0].strip(" ：:，,、")
    if addr:
        pieces.append(f"地址:{addr[:50]}")
    own = fields.get("备注", "")
    if own:
        pieces.append(own)
    others = normalize_text(fields.get("其他联系人", ""))
    if others:
        pieces.append(f"其他联系人:{others}")
    # 跨站回捞不上的岗位不静默丢弃：它属于另一个联系人，写明出处挂备注（即使
    # 主岗位已有值也并列展示，因为那是不同人的职务），外呼人员自己判断可信度
    borrowed = normalize_text(fields.get("岗位(跨站)", ""))
    if borrowed:
        pieces.append(f"岗位(跨站来源):{borrowed}")
    # 岗位族：没有明写岗位列但旁证能归类时，给外呼一个可用的话术分支
    if fam and not normalize_text(fields.get("岗位", "")):
        pieces.append(f"岗位族:{fam}")
    remark = " | ".join(pieces)
    return remark[:180]


def _position_family(position, email, company):
    """岗位族：明写岗位能分类就分类；否则依次用邮箱前缀(sales@/export@)、联系人值
    里的部门词、公司名尾巴「销售部」做智能匹配（只写页面能证实的线索，不臆造）。"""
    helpers = _load_family_helpers()
    if not helpers:
        return ""
    _family_of, _BARE_TITLE_RE, _dept_from_email, _dept_from_company, _dept_only = helpers
    pos = (position or "").strip()
    if pos and not _BARE_TITLE_RE.fullmatch(pos):
        fam = _family_of(pos) or _family_of(_dept_from_company(pos) or "")
        if fam:
            return fam
    for ev in (_dept_from_email(email or ""),
               _dept_only(pos) if pos else "",
               _dept_from_company(company or "")):
        if ev and _family_of(ev):
            return _family_of(ev)
    return ""


def make_row(fields: dict, sources: list) -> dict:
    """一份归一化字段 + 来源列表 → 模板行（采集链路与外部数据清洗共用这一套规则）"""
    mobiles, landlines = split_phones(fields.get("电话", ""))
    emails = split_emails(fields.get("邮箱", ""))
    legal_rep = _clean_contact(_clean_legal_rep(fields.get("法定代表人", "")))
    contact = _clean_contact(_clean_legal_rep(fields.get("联系人", "")))
    # 公司名尾部杂质剥离（实测采集页标题带分页版次：「…材料有限公司3.」）
    company = _NAME_MARKETING_TAIL_RE.sub(
        "", _NAME_TAIL_JUNK_RE.sub("", fields.get("公司名", ""))).strip()
    position = normalize_text(fields.get("岗位", ""))
    # 兜底：联系人/法定代表人取到工商状态值（状态列被误认成联系人列就是这样），
    # 一律作废并退回公司名——客户名称写成「存续」是致命的，宁可没有联系人。
    if _is_status(contact):
        contact = ""
    if _is_status(legal_rep):
        legal_rep = ""
    # 岗位族旁证只用同行证据算：联系人列被退回公司名时不拿公司名当部门依据
    _cname_for_fam = contact or legal_rep if (contact or legal_rep) != company else ""
    # 客户名称优先用显式联系人；没有联系人但采到法定代表人时，岗位才补「法定代表人」
    return {
        "客户名称": contact or legal_rep or company,
        "企业名称": company,
        "岗位": position or ("" if contact else ("法定代表人" if legal_rep else "")),
        "手机号": _join_limited(mobiles),
        "固话": _join_limited(landlines),
        "邮箱": _join_limited(emails),
        "微信号": normalize_text(fields.get("微信号", "")),
        "数据来源": ",".join(s for s in sources if s),
        "外呼备注": _remark(fields, contact_ok=bool(mobiles or landlines or emails),
                          fam=_position_family(position, fields.get("邮箱", ""), _cname_for_fam)),
        "备注": _remark(fields, contact_ok=bool(mobiles or landlines or emails),
                      fam=_position_family(position, fields.get("邮箱", ""), _cname_for_fam)),
    }


def build_clean_rows(results: list, log=print, waste_sink=None):
    """把三平台原始结果转换为模板行数据：返回 [{模板列名: 值}] 列表。
    waste_sink：可选回调(list[公司名])。被入库口径剔掉的「四项联系方式全空」
    合并记录的公司名会交给它——b2bcollect 用它把废数据灌进待采集名单交企知道
    补采（变废为宝），不传则行为与旧版完全一致。"""
    merged, dup = merge_records(results)
    if waste_sink is not None:
        waste = []
        for item in merged:
            f = item["fields"]
            mobiles, landlines = split_phones(f.get("电话", ""))
            if not mobiles and not landlines and not split_emails(f.get("邮箱", "")) \
                    and not normalize_text(f.get("微信号", "")):
                name = normalize_text(f.get("公司名", ""))
                if name:
                    waste.append(name)
        if waste:
            try:
                waste_sink(waste)
            except Exception:
                pass
    rows = [make_row(item["fields"], item["sources"]) for item in merged]
    no_mobile = sum(1 for r in rows if not r["手机号"] and not r["固话"])
    log(f"清洗完成：原始 {len(results)} 条 → 有效入库 {len(rows)} 条"
        f"（剔除无效 {len(results) - sum(1 for r in results if is_valid_record(r))} 条、"
        f"跨平台合并 {dup} 条、无联系方式 {no_mobile} 条）")
    return rows


def locate_template(start_dirs=None) -> str:
    """在常用目录里查找客户的入库模板文件，找到返回路径，否则空串。"""
    home = os.path.expanduser("~")
    dirs = list(start_dirs or []) + [
        os.getcwd(),
        os.path.dirname(os.path.abspath(__file__)),
        os.path.join(home, "Downloads"),
        os.path.join(home, "Desktop"),
        os.path.join(home, "Documents"),
    ]
    for d in dirs:
        if not d or not os.path.isdir(d):
            continue
        try:
            names = sorted(os.listdir(d))
        except OSError:
            continue
        for fn in names:
            if not fn.lower().endswith((".xlsx", ".xls")):
                continue
            # 跳过 Excel/WPS 打开时产生的临时锁文件（~$xxx / .~xxx / .xxx）
            if fn.startswith(("~$", ".~", ".")):
                continue
            stem = os.path.splitext(fn)[0]
            if any(k in stem for k in _TEMPLATE_SEARCH_NAMES):
                return os.path.join(d, fn)
    return ""


def read_template_layout(template_path: str, log=print):
    """读取模板表头，返回 (外呼块列索引, 备注块列索引, 完整表头)。

    列名在模板里出现两次（A-I 与 J-R），按出现顺序区分两块。
    找不到模板时使用内置表头兜底。
    """
    header = list(FALLBACK_HEADER)
    if template_path:
        try:
            from openpyxl import load_workbook
            ws = load_workbook(template_path, read_only=False, data_only=True).active
            first = [normalize_text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
            if any(first):
                header = first
                log(f"已读取入库模板表头：{os.path.basename(template_path)}（{len(header)} 列）")
        except Exception as e:
            log(f"⚠ 读取模板失败，改用内置表头：{e}")

    pos = {}
    for idx, col in enumerate(header):
        pos.setdefault(col, []).append(idx)
    data_cols = ("客户名称", "企业名称", "岗位", "手机号", "固话",
                 "邮箱", "微信号", "数据来源")
    has_second_block = any(len(idxs) > 1 for idxs in pos.values())
    # 外呼块：以「外呼备注」所在列界定，同名列取其之前的第一次出现
    outbound_remark = "外呼备注" if "外呼备注" in pos else ("备注" if "备注" in pos else "")
    outbound_idx = pos[outbound_remark][-1] if outbound_remark else None
    outbound = {}
    for col in data_cols:
        idxs = pos.get(col) or []
        pick = [i for i in idxs if outbound_idx is None or i < outbound_idx]
        outbound[col] = (pick or idxs)[0]
    if outbound_remark:
        outbound[outbound_remark] = outbound_idx
    # 备注块：同名列取最后一组（模板 A-I / J-R 重复列头时落到第二块）
    remark_remark = "备注" if "备注" in pos else "外呼备注"
    remark_block = {}
    if has_second_block:
        for col in data_cols:
            if col in pos:
                remark_block[col] = pos[col][-1]
        if remark_remark in pos:
            remark_block[remark_remark] = pos[remark_remark][-1]
    return outbound, remark_block, header


def write_clean_excel(rows: list, filepath: str, template_path: str = "", log=print):
    """按模板列结构写出清洗后的入库数据 Excel（两块列头均填充，保证入库兼容）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    outbound, remark_block, header = read_template_layout(template_path, log)
    wb = Workbook()
    ws = wb.active
    ws.title = "客户库"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    thin = Border(*[Side(style="thin")] * 4)

    for col_idx, title in enumerate(header, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    # 两块列头都填充同样数据：外呼块用「外呼备注」，备注块用「备注」
    def values_for(row, block):
        out = {}
        for col, idx in block.items():
            if col in ("外呼备注", "备注"):
                out[idx] = row.get("外呼备注") or row.get("备注", "")
            else:
                out[idx] = row.get(col, "")
        return out

    for r, row in enumerate(rows, 2):
        cells = {}
        for block in (outbound, remark_block):
            for idx, val in values_for(row, block).items():
                cells[idx] = val
        for idx in range(len(header)):
            cell = ws.cell(row=r, column=idx + 1, value=cells.get(idx, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin

    widths = {"客户名称": 16, "企业名称": 34, "岗位": 12, "手机号": 20,
              "固话": 20, "邮箱": 32, "微信号": 14, "数据来源": 12,
              "外呼备注": 40, "备注": 40}
    for idx, title in enumerate(header, 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = \
            widths.get(title, 16)

    ws.freeze_panes = "A2"
    wb.save(filepath)
    log(f"✅ 已生成清洗后入库数据：{filepath}（{len(rows)} 条）")
    return filepath


# ═══════════════════════════════════════════════════════
#  外部数据智能清洗（「清洗数据」按钮）
# ═══════════════════════════════════════════════════════
# 别的渠道拿到的表，列名千奇百怪（「客户名称」有时是公司、有时是联系人），
# 写死列名对不上。这里的做法是：表头别名匹配 + 单元格内容抽样判定，两路打分
# 取最高，再把每个字段只认一列（贪心，避免两列都叫“名称”时抢写头）。
# 实在对不上的列不隐式丢弃，在日志里逐列列出，用户能看到“没采哪些列”。

# 对外清洗只关心这些规范字段（其余列一律归入「未识别」）
EXTERNAL_FIELDS = ("公司名", "联系人", "岗位", "电话", "邮箱", "微信号",
                   "官网", "地址", "状态", "数据来源", "备注")

# 字段 → (强别名, 弱别名)：强别名归一化后包含即算高置信，弱别名只给低分
EXTERNAL_ALIASES = {
    "公司名": (("企业名称", "公司名称", "单位名称", "企业全称", "公司全称", "注册名称",
                "工商登记名称", "customername", "companyname", "company", "enterprise", "orgname"),
               ("企业名", "公司名", "单位名", "机构名", "客户公司", "客户单位", "名称", "企业", "公司", "单位")),
    "联系人": (("联系人姓名", "联系人", "客户名称", "客户联系人", "对接人", "业务联系人",
                "联络人", "主要联系人", "contactperson", "contactname", "contact"),
               ("姓名", "客户姓名", "联系人姓", "负责人")),
    "岗位": (("职位", "职务", "岗位", "担任职位", "客户职位", "jobtitle", "position", "title"), ()),
    "电话": (("联系电话", "联系方式", "手机号码", "手机号", "移动电话", "公司电话",
                "企业电话", "办公电话", "电话号码", "电话", "mobile", "phone", "tel",
                "telephone", "contactnumber", "cellphone"),
             ("手机", "固话", "座机", "号码", "直线", "电话")),
    "邮箱": (("电子邮箱", "邮箱地址", "企业邮箱", "公司邮箱", "邮箱", "email", "e-mail", "mail"),
             ("邮件",)),
    "微信号": (("微信号", "微信", "wechat", "weixin", "wx号", "wxid"), ()),
    "官网": (("公司网站", "企业网站", "官方网站", "官网网址", "公司网址", "企业网址",
                "网址", "网站", "website", "url", "web"), ("官网", "站点")),
    "地址": (("公司地址", "企业地址", "单位地址", "通信地址", "注册地址", "办公地址",
                "详细地址", "联系地址", "地址", "address", "addr"),
             ("地区", "所在地", "所在区域", "区域")),
    "状态": (("登记状态", "企业状态", "经营状态", "公司状态", "存续状态", "status"), ("状态",)),
    "数据来源": (("数据来源", "采集来源", "来源平台", "数据渠道", "source", "channel"),
                 ("来源", "渠道")),
    "备注": (("外呼备注", "备注说明", "备注", "remark", "note", "memo", "comment"),
             ("说明", "注记", "补充")),
}

# 这些列对入库没用（序号/注册资本/统一社会信用代码等），即使包含“名称/号码”也不给认
EXTERNAL_IGNORE = ("序号", "编号", "行号", "水号", "税号", "统一社会信用代码", "信用代码",
                   "组织机构代码", "工商注册号", "注册资本", "实缴资本", "成立日期", "成立时间",
                   "核准日期", "营业期限", "参保人数", "人员规模", "所属行业", "行业",
                   "企业类型", "类型", "币种", "经营范围", "登记机关", "邮编", "邮政编码",
                   "postal", "zip", "传真", "fax", "qq", "二维码", "标签", "分组", "创建时间",
                   "更新时间", "修改时间", "操作人", "评分", "风险", "诉讼", "商标", "专利",
                   "登录密码", "账号密码", "附件", "图片")

# 内容特征（只用于“这列装的是什么”的推断，不参与写盘）
_EXT_COMPANY_RE = re.compile(
    r"(公司|集团|企业|工厂|厂家|研究所|研究院|设计院|中心|商行|商店|合作社|事务所|银行|医院|"
    r"学校|学院|工作室|营业部|分公司|分店| Ltd|LTD|LIMITED|INC|LLC|GMBH|CORP|CO\.)", re.I)
# 公司名一般以组织后缀结尾；靠这个才能把“企业名称/公司简称”这类列名本身和真实公司名分开
_EXT_COMPANY_TAIL_RE = re.compile(
    r"(公司|集团|企业|工厂|厂家|研究所|研究院|设计院|中心|商行|商店|合作社|事务所|银行|医院|"
    r"学校|学院|工作室|营业部|分公司|分店|Ltd|LTD|LIMITED|Inc|LLC|GmbH|Corp|CO\.)\s*[）)]*$", re.I)
_EXT_PERSON_RE = re.compile(r"^[\u4e00-\u9fa5·]{2,4}$|^[A-Za-z]{2,12}( [A-Za-z]{1,12}){0,3}$")
_EXT_ADDRESS_RE = re.compile(r"(省|自治区|市|区|县|镇|乡|街道|路|街|巷|号|大道|工业园|科技园|开发区|村)")
_EXT_URL_RE = re.compile(r"^(https?://|www\.)|\.(com|cn|net|org|com\.cn|cc)(/|$)", re.I)
# 工商登记状态值。这类两三个字的词（存续/在业/注销/吊销未注销）会被人名正则当成
# 「2 字联系人」，状态列就此抢下联系人字段、客户名称整列写成「存续」。
_EXT_STATUS_WORDS = {
    "存续", "在业", "在营", "开业", "在册", "正常", "有效", "歇业", "经营", "运营",
    "注销", "已注销", "未注销", "吊销", "已吊销", "吊销未注销", "未注册", "撤销", "撤消",
    "关停", "关闭", "停产", "停业", "清算", "迁移", "迁出", "迁入", "移出", "转入",
    "待查", "其他", "上市", "未上市", "已上市",
    "active", "existing", "operational", "cancelled", "canceled", "revoked", "dormant",
}
# 只由状态字组成的变体（在营/已吊销/未注销…）也认；不含「正」「常」等常见姓氏用字，
# 免得「常正」这种真人名被判成状态后联系人被洗没。
_EXT_STATUS_CHARS = set("存续在业开注销吊未已停歇清算迁移闭查效市营搁撤关")
_EXT_STATUS_SPLIT_RE = re.compile(r"[、，,/|()（）\s;；]+")


def _cell_text(value) -> str:
    """单元格→文本：Excel 里手机号经常是数字类型，直接 str() 会拿到 1.38001e+10 这种断号。

    int/float 一律展开成完整数字串，日期取年月日，None/空值给空串。
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.10f}".rstrip("0").rstrip(".")
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%Y-%m-%d")
        except Exception:
            return normalize_text(value)
    return normalize_text(value)


def _norm_col(name) -> str:
    """列名归一化：先去掉括号里的注释，再去空白标点、转小写，让“联系电话(手机)”也能对上"""
    s = re.sub(r"[（(][^）)]*[）)]", "", str(name or ""))
    s = re.sub(r"[\s：:、,，.。_\-/#*]+", "", s)
    return s.lower()


def _is_company(v: str) -> bool:
    """这串内容像不像一个真实公司名（不能只靠“含企业/公司”，否则“企业名称”这种列名自己就会被当成数据）"""
    if not v or _EXT_URL_RE.search(v):
        return False
    if _EXT_COMPANY_TAIL_RE.search(v):
        return True
    # 中间带主体词的长串（“XX公司深圳分公司”、括号后缀等）也算
    return len(v) >= 8 and bool(_EXT_COMPANY_RE.search(v))


def _is_status(v: str) -> bool:
    """是否是工商状态值：存续 / 存续(在营、开业、在册) / 吊销未注销 / active。"""
    parts = [p for p in _EXT_STATUS_SPLIT_RE.split(normalize_text(v)) if p]
    if not parts or len(parts) > 4:
        return False
    for p in parts:
        low = p.lower()
        if low in _EXT_STATUS_WORDS:
            continue
        if 2 <= len(low) <= 6 and all(ch in _EXT_STATUS_CHARS for ch in low):
            continue
        return False
    return True


def _is_person(v: str) -> bool:
    return bool(v) and bool(_EXT_PERSON_RE.match(v)) and not _is_company(v) and not _is_status(v)


def _is_email(v: str) -> bool:
    if not v:
        return False
    low = v.lower()
    # “邮箱: a@b.com” 这种带前缀的脏表头也得认出来
    return bool(_EMAIL_RE.match(low)) or bool(_EMAIL_EXTRACT_RE.search(low))


def _is_phone(v: str) -> bool:
    if not v or not re.search(r"\d", v):
        return False
    mob, land = split_phones(v)
    return bool(mob or land)


def _is_url(v: str) -> bool:
    return bool(v) and bool(_EXT_URL_RE.search(v.strip().lower()))


def _is_address(v: str) -> bool:
    return bool(v) and len(v) >= 6 and bool(_EXT_ADDRESS_RE.search(v)) and not _is_company(v)


def _ratio(values, pred) -> float:
    vals = [v for v in values if v]
    if not vals:
        return 0.0
    return sum(1 for v in vals if pred(v)) / len(vals)


def content_score(field: str, values: list) -> float:
    """按抽样内容给这一列属于 `field` 的置信度（0~100）"""
    if not values:
        return 0.0
    if field == "公司名":
        return 100 * _ratio(values, _is_company)
    if field == "联系人":
        return 100 * _ratio(values, _is_person)
    if field == "电话":
        return 100 * _ratio(values, _is_phone)
    if field == "邮箱":
        return 100 * _ratio(values, _is_email)
    if field == "官网":
        return 100 * _ratio(values, _is_url)
    if field == "地址":
        return 100 * _ratio(values, _is_address)
    if field == "状态":
        # 状态列的表头「状态」只是弱别名（39 分够不到 40 的录用线），全靠内容把它认领，
        # 否则这一列会被「存续」像人名这点抢去当联系人。
        return 100 * _ratio(values, _is_status)
    # 岗位/微信号/备注/数据来源靠内容猜不出来，只看表头
    return 0.0


def header_score(field: str, name: str) -> float:
    """表头别名打分：强别名 100 / 弱别名 65 / 命中忽略词 -60"""
    norm = _norm_col(name)
    if not norm:
        return 0.0
    strong, weak = EXTERNAL_ALIASES.get(field, ((), ()))
    score = 0.0
    if any(k in norm for k in (x.lower() for x in strong)):
        score = 100.0
    elif any(k in norm for k in (x.lower() for x in weak)):
        score = 65.0
    if score < 100 and any(k in norm for k in EXTERNAL_IGNORE):
        return -60.0        # 低置信又命中忽略词 → 这列本来就不是入库字段
    return score


def match_columns(headers: list, data_rows: list, log=None):
    """智能匹配列名：返回 (映射 {列下标: 规范字段}, 已识别描述, 未识别列名)

    表头占 6 成、内容占 4 成：“客户名称”列里装的是公司还是联系人，单看表头判不了。
    同一字段只认一列、同一列只认一个字段，按总分贪心分配。
    """
    log = log or (lambda _m: None)
    ncols = len(headers)
    samples = {}
    for c in range(ncols):
        vals = []
        for row in data_rows[:20]:
            if c < len(row):
                text = _cell_text(row[c])
                if text:
                    vals.append(text)
        samples[c] = vals

    claims = []
    for c in range(ncols):
        for field in EXTERNAL_FIELDS:
            hs = header_score(field, headers[c])
            cs = content_score(field, samples[c])
            total = hs * 0.6 + cs * 0.4
            if hs < 0:
                continue
            if total >= 40:
                claims.append((total, c, field))
    claims.sort(key=lambda x: -x[0])

    mapping, used_cols, used_fields = {}, set(), set()
    for total, c, field in claims:
        if c in used_cols or field in used_fields:
            continue
        mapping[c] = field
        used_cols.add(c)
        used_fields.add(field)

    # 公司名是入库的硬性必要条件：没认出来就退而求其次取“最像公司名”的一列
    for field, pick in (("公司名", lambda c: content_score("公司名", samples[c])),
                        ("电话", lambda c: content_score("电话", samples[c]))):
        if field in used_fields:
            continue
        best, best_v = -1, 0.45 if field == "电话" else 0.30
        for c in range(ncols):
            if c in used_cols:
                continue
            v = pick(c)
            if v > best_v:
                best, best_v = c, v
        if best >= 0:
            mapping[best] = field
            used_cols.add(best)
            used_fields.add(field)
            log(f"  ↳ 没找到「{field}」列名，按内容把“{headers[best] or '第%d列' % (best + 1)}”当{field}使用")

    identified = [f"“{headers[c] or '第%d列' % (c + 1)}”→{field}"
                  for c, field in sorted(mapping.items())]
    skipped = [str(headers[c] or "第%d列" % (c + 1)).strip() for c in range(ncols) if c not in mapping]
    return mapping, identified, skipped


def _is_label(name: str) -> bool:
    """单元格是否就是个字段标题（命中任一入库别名）——用于把表头行“拉”出来"""
    norm = _norm_col(name)
    if not norm:
        return False
    for strong, weak in EXTERNAL_ALIASES.values():
        for kw in strong + weak:
            if kw in norm:
                return True
    return False


def detect_header_row(rows: list):
    """定位表头行：返回 (表头列名, 数据行)。无表头的文件也能走（全列名置空、从第 1 行读数据）

    真实表头的特征：非空单元格多、内容短、命中字段别名；反过来，一行里出
    现多个公司名/邮箱/网址，就说明它本身就是数据行。
    """
    if not rows:
        return [], []
    scan = min(len(rows), 12)
    best_i, best_s = 0, None
    for i in range(scan):
        cells = [_cell_text(v) for v in rows[i]]
        cells = [c for c in cells if c]
        if not cells:
            continue
        labelish = sum(1 for c in cells if len(c) <= 20 and not _is_phone(c) and not _is_email(c))
        labelhit = sum(1 for c in cells if _is_label(c))
        dataish = sum(1 for c in cells if _is_company(c) or _is_email(c) or _is_url(c))
        score = len(cells) * 2 + labelish * 1.5 + labelhit * 3 - dataish * 6 - i * 0.5
        if best_s is None or score > best_s:
            best_s, best_i = score, i
    header = [_cell_text(v) for v in rows[best_i]]
    cells = [c for c in header if c]
    dataish = sum(1 for c in cells if _is_company(c) or _is_email(c) or _is_url(c))
    # 候选行本身全是公司名/邮箱 → 它其实是数据行（无表头文件）
    if not cells or dataish > len(cells) * 0.4:
        ncols = max(len(r) for r in rows)
        return [""] * ncols, list(rows)
    return header, list(rows[best_i + 1:])


def _read_csv_tables(path: str):
    import csv

    last_err = None
    for enc in ("utf-8-sig", "gbk", "utf-8", "gb18030", "big5"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                raw = fh.read()
            sample = raw[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|\u3000，、")
                sep = dialect.delimiter
            except Exception:
                sep = ","
            rows = [list(r) for r in csv.reader(raw.splitlines(), delimiter=sep)]
            return [(os.path.basename(path), rows)]
        except UnicodeDecodeError as e:
            last_err = e
            continue
        except Exception as e:
            last_err = e
            break
    raise ValueError(f"CSV 读取失败（{last_err}），请确认编码为 UTF-8 或 GBK")


def read_tables(path: str):
    """读外部数据文件 → [(工作表名, [[单元格, …], …])]，支持 .xlsx/.xlsm/.csv（.xls 不支持）"""
    if not path or not os.path.exists(path):
        raise FileNotFoundError(f"文件不存在：{path}")
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        raise ValueError(f"旧版 .xls 格式读不了：{os.path.basename(path)}\n"
                        f"请先用 Excel/WPS 另存为 .xlsx 再清洗")
    if ext == ".csv":
        return _read_csv_tables(path)
    if ext not in (".xlsx", ".xlsm"):
        raise ValueError(f"不支持的文件类型：{ext or '无后缀'}（只支持 .xlsx / .xlsm / .csv）")
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    tables = []
    try:
        for ws in wb.worksheets:
            rows = [[_cell_text(v) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c for c in r)]
            if len(rows) >= 2:
                tables.append((ws.title, rows))
    finally:
        wb.close()
    return tables


def records_from_table(headers: list, data_rows: list, mapping: dict, default_source: str):
    """按列映射把一张表转成规范字段记录（列名为空的无表头文件也能转）"""
    records = []
    for row in data_rows:
        if not row or all(not _cell_text(v) for v in row):
            continue
        rec = {}
        for col, field in mapping.items():
            if col >= len(row):
                continue
            text = _cell_text(row[col])
            if text:
                rec[field] = text
        if not any(rec.values()):
            continue
        if not rec.get("数据来源"):
            rec["数据来源"] = default_source
        records.append(rec)
    return records


def merge_external_records(records: list):
    """外部数据合并：同名公司且同联系人才合并（一家公司多个联系人不能压成一条）

    采集链路按公司名合并是安全的（一家公司一条），外部名单经常是“一家公司
    好几个人”，只按公司名合会把电话归给同一个人、其他联系人直接消失。
    """
    merged, order, dup, invalid = {}, [], 0, 0
    for rec in records:
        if not is_valid_record(rec):
            invalid += 1
            continue
        name = _normalize_key(_get_field(rec, "公司名"))
        if not name:
            invalid += 1
            continue
        contact = _normalize_key(_get_field(rec, "联系人") or _get_field(rec, "法定代表人"))
        source = _get_field(rec, "数据来源") or "外部导入"
        key = (name, contact)
        if key in merged:
            dup += 1
            item = merged[key]
            if source not in item["sources"]:
                item["sources"].append(source)
            for f in list(SOURCE_FIELDS):
                if not item["fields"].get(f):
                    item["fields"][f] = _get_field(rec, f)
            for f in ("电话", "邮箱", "微信号"):
                more = _get_field(rec, f)
                if more and f in item["fields"] and more not in item["fields"][f]:
                    base = item["fields"][f]
                    item["fields"][f] = (base + ";" + more) if base else more
        else:
            merged[key] = {"sources": [source],
                           "fields": {f: _get_field(rec, f) for f in SOURCE_FIELDS}}
            order.append(key)
    return [merged[k] for k in order], dup, invalid


def build_external_rows(records: list, log=print):
    """外部记录 → 模板行（内容与采集数据同一套清洗规则，只换合并口径）"""
    merged, dup, invalid = merge_external_records(records)
    rows = [make_row(item["fields"], item["sources"]) for item in merged]
    no_contact = sum(1 for r in rows if not r["手机号"] and not r["固话"] and not r["邮箱"])
    log(f"清洗完成：原始 {len(records)} 行 → 有效入库 {len(rows)} 条"
        f"（剔除无效 {invalid} 行、重复合并 {dup} 行、无任何联系方式 {no_contact} 条）")
    return rows


def clean_external_records(paths: list, log=print):
    """读多个外部数据文件，智能匹配列名并归一成规范记录

    返回 (records, report)：report 逐文件逐表记录“哪些列认成了什么、哪些列被丢了、
    各多少行”，清洗不是黑盒，用户得能对上账。识别错误的列名全部列在日志里，
    比静默丢列后“莫名少了数据”好排查。未知格式只报错不中断，一个文件坏了不影响其他文件。
    """
    records, report = [], []
    for path in paths or []:
        name = os.path.basename(path)
        stem = os.path.splitext(name)[0]
        try:
            tables = read_tables(path)
        except Exception as e:
            report.append({"file": name, "sheet": "-", "rows": 0, "error": str(e)})
            log(f"⚠ 跳过「{name}」：{e}")
            continue
        if not tables:
            report.append({"file": name, "sheet": "-", "rows": 0,
                           "error": "表里没有数据"})
            log(f"⚠ 跳过「{name}」：没读到任何数据行")
            continue
        for sheet, rows in tables:
            headers, data_rows = detect_header_row(rows)
            if not any(headers):
                headers = [""] * max(len(r) for r in rows)
            while len(headers) < max((len(r) for r in rows), default=0):
                headers.append("")
            mapping, identified, skipped = match_columns(headers, data_rows, log=log)
            if not any(c in mapping for c in range(len(headers))):
                report.append({"file": name, "sheet": sheet, "rows": len(data_rows),
                               "error": "没认到任何可用列"})
                log(f"⚠ 「{name}」·{sheet}：列名和内容都没对上入库字段，本表跳过")
                continue
            part = records_from_table(headers, data_rows, mapping,
                                      f"外部导入:{stem}"[:60])
            records.extend(part)
            info = {"file": name, "sheet": sheet, "rows": len(part),
                    "identified": identified, "skipped": skipped}
            report.append(info)
            log(f"📄 「{name}」·{sheet}：{len(data_rows)} 行→{len(part)} 条· "
                f"识别列：{'，'.join(identified) or '无'}"
                + (f"；未使用列：{'、'.join(skipped[:8])}" if skipped else ""))
    return records, report
