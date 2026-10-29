"""
b2bcollect_gui.py —— B2B 多平台企业数据自动采集工具（PyQt5 GUI）

一个共享 Chrome、需浏览器的站各占一个标签并行采集：
  · 关键词并行（5 站）：中国制造网 made-in-china.com / 八方资源网 b2b168.com /
    顺企网 11467.com / 爱采购 b2b.aipage.com（后三站纯 HTTP 不占标签）/ 1688（需登录，默认不勾）；
  · 名单驱动（企知道）：从「待采集公司名单」逐家补采，看门狗语义——名单空则休眠，
    关键词采集清洗掉的无联系方式废数据会自动回流到名单、唤醒企知道补采（变废为宝）。
导入或粘贴一批关键词（或上传公司名单）→ 智能匹配行业列表页 → 进详情页取联系方式 →
边采边自动保存 → 导出「原始 + 清洗后入库」两份 Excel；支持断点续采、被验证码/风控拦截时
弹窗等人工过验证后自动继续、导入云展动力客户库。

运行：python b2bcollect_gui.py
"""
import os
import re
import sys
import time
import zipfile
import threading
from pathlib import Path
from datetime import datetime

from PyQt5.QtCore import Qt, QThread, pyqtSignal, QTimer
from PyQt5.QtGui import QFont, QPalette
from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QTextEdit, QFileDialog, QSpinBox, QDoubleSpinBox,
    QProgressBar, QMessageBox, QGroupBox, QPlainTextEdit, QDialog,
    QLineEdit, QCheckBox, QListWidget, QListWidgetItem, QAbstractItemView,
    QStackedWidget, QButtonGroup, QInputDialog, QDialogButtonBox,
)

from parallel_session import ParallelSession, SITE_ORDER, SITE_NAMES
import customer_library_cleaner as _cleaner
import keyword_expander as _kwexp

# 状态行站点状态标记（图标+颜色+文字三通道，对齐 GitHub Actions/Jenkins 任务实况惯例）：
# done=✅绿粗「已完成」/ stopped=■橙粗「已停止」/ running=●蓝「进行中」
# 字形选高覆盖度的 ✅/●/■（避开 ⏳ 这类小字号下会糊成方块的 emoji）
STATS_STYLE = {
    "done": ('#2e7d32', "✅"),
    "stopped": ('#e65100', "■"),
    "running": ('#1a73e8', "●"),
}
# 深色主题下的提亮版：色相完全一致，只把亮度抬起来。
# #2e7d32 这种深绿在近黑底上对比度不足（看着像「已完成」压根没点亮），
# 状态行是最需要一眼扫到的信息，绝不能靠用户去猜颜色。
STATS_STYLE_DARK = {
    "done": ('#66bb6a', "✅"),
    "stopped": ('#ffa726', "■"),
    "running": ('#8ab4f8', "●"),
}


def _muted_color():
    """次要文字颜色按系统深浅色自适应。
    写死 #666/#999 在深色模式下会「字在但看不见」——实测深色系统上五行站点看板
    在近黑底上完全隐形，只剩一块空白把输出日志挤出屏幕。取前景与底色的中间值，
    深浅两套主题下都保证可读，同时仍是视觉上的次要层级。"""
    try:
        pal = QApplication.palette()
        fg = pal.color(QPalette.WindowText).getRgb()[:3]
        bg = pal.color(QPalette.Window).getRgb()[:3]
        return "#%02x%02x%02x" % tuple((f + b) // 2 for f, b in zip(fg, bg))
    except Exception:
        return "#888888"


def _muted_style(extra=""):
    return f"color:{_muted_color()};{extra}"


def _is_dark_theme():
    """当前是否深色主题（写死颜色在深/浅底上的可读性差很多，必须先看一眼背景亮度）。"""
    try:
        return QApplication.palette().color(QPalette.Window).lightness() < 128
    except Exception:
        return False


def _danger_style(extra=""):
    """错误/提醒红：深色底用亮红，浅色底用深红（同一个 #c62828 在深底上几乎看不清）。"""
    return f"color:{'#ff6b6b' if _is_dark_theme() else '#c62828'};{extra}"


def _tab_toggle_style():
    """分组页「国内/国外」切换按钮的默认态配色按主题给两套。
    写死浅蓝底会在深色面板上亮出一块补丁，和主窗口风格不一致。"""
    checked = "QPushButton:checked { background: #4472C4; color: white; font-weight: bold; }"
    if _is_dark_theme():
        return ("QPushButton { background: #3a3d42; color: #d5d8dc; border: 1px solid #4a4e55; }"
                + checked)
    return ("QPushButton { background: #eef1f6; color: #1a3c6e; border: 1px solid #d0d8e8; }"
            + checked)


def _status_badge_style(tone):
    """云展绑定状态徽标（error / ok / neutral 三态）。
    写死浅底深字在深色主题下像贴了张便签，三态亮度差还会被放大成一个很跳两个很暗；
    这里按主题给成套的低饱和底色加同层级前景，三态只靠色相区分、不靠亮度抢视线。"""
    dark = _is_dark_theme()
    if tone == "error":
        bg, fg = ("#3a2224", "#ff6b6b") if dark else ("#fdecea", "#c62828")
    elif tone == "ok":
        bg, fg = ("#1f3324", "#7fc97f") if dark else ("#e8f5e9", "#2e7d32")
    else:
        bg, fg = ("#2f3134", _muted_color()) if dark else ("#f0f0f0", _muted_color())
    return ("QLabel { padding: 4px 10px; border-radius: 4px; background: %s; color: %s; }" % (bg, fg))


# 各站链路、防采策略说明：从面板标题里抽出来挂到 tooltip，界面只留短句，
# 详细机制悬停可查（避免一段长文字把整个分区挤乱）。
SITE_TIPS = {
    "mic": "中国制造网 cn.made-in-china.com：中文公司搜索→showroom 联系页（meta 明文中文联系人/电话）。"
           "搜索需浏览器标签过验证码；联系方式采集会复用会话 Cookie 走后台 HTTP，"
           "命中 302 自动回落标签，标签越少可见时间越短。验证码同样接看门狗，过完自动续采。",
    "b2b168": "八方资源网 b2b168.com：纯 HTTP 搜索(Index.aspx POST)→公司自有子域商铺首页+「关于我们」页，"
              "完全不占用浏览器标签（所以看不到它的标签是正常的），速度最快。\n"
              "防采：搜索接口限流时返 4xx/5xx 插页，识别到即停手冷却等恢复（不拉浏览器），绝不当空数据写库。",
    "11467": "顺企网 11467.com：纯 HTTP 关键词搜索(b2b.11467.com)→公司商铺子域联系方式页(contact.asp)+商铺首页，"
              "联系人/手机/固话/邮箱/QQ 微信都在标签行里直接解析；平时不占用浏览器标签。\n"
              "防采：请求过快会返 HTTP 429 + v.js 人机挑战（纯 HTTP 不执行 JS，拿到的就是空页）——"
              "这不是「没有公司」，已接看门狗：弹窗提醒 + 临时借一个标签把验证页摆到眼前，"
              "人工过验后把 clearance Cookie 抄回 HTTP 会话自动续采；抄不回来就降级用标签继续采。",
    "aipage": "爱采购 b2b.aipage.com（百度爱采购镜像站）：纯 HTTP 关键词搜索→搜索接口翻页(/s/a)→"
              "进店联系方式页(/shop?xzhid=&tpath=contact)，联系人/手机/固话/邮箱/微信/地址服务端直接注水解析；"
              "完全不占用浏览器标签。公司按店铺 id 去重，同店多产品只采一次。\n"
              "防采：限流插页识别到即停手冷却等恢复（不拉浏览器），绝不当空数据写库。",
    "1688": "1688 www.1688.com：关键词→商品/店铺搜索→公司详情页。需先登录（有登录限制，故默认不勾选）；"
            "勾选后会弹窗提示并可一键打开 1688 登录页，请先在浏览器里确认已登录再开始采集。\n"
            "与其它需浏览器的站共享同一个 Chrome（各占自己标签），绝不会另开浏览器。",
    "qzd": "企知道 www.qizhidao.com：**名单驱动**通道（不参与关键词并行）。从「待采集公司名单」逐家模拟人工查询，"
           "拿工商详情/联系方式。需登录（点右侧「打开企知道」按钮先登录）。\n"
           "看门狗：名单空则休眠零占用，关键词采集清洗掉的无联系方式废数据自动回流入列后自动苏醒补采；"
           "采过/查不到的公司名会从名单剔除，避免重复采集。",
}


def build_stats_html(text, site_status=None):
    """把状态行文本「已采 N 家　站名 x｜站名 y」渲染为带图标+着色的富文本。
    纯函数、无副作用，便于离线断言；无站点状态时原样返回纯文本（降级兼容）。"""
    status = site_status or {}
    if not status or "<" in (text or ""):
        return text or ""
    head, _, tail = text.partition("　")
    palette = STATS_STYLE_DARK if _is_dark_theme() else STATS_STYLE
    segs = []
    for s in SITE_ORDER:
        name = SITE_NAMES.get(s)
        if not name or name not in tail:
            continue
        i = tail.index(name)
        j = tail.find("｜", i)
        raw = tail[i:j].strip() if j >= 0 else tail[i:].strip()
        st = status.get(s, "running")
        color, icon = palette.get(st, palette["running"])
        mark = {"done": "已完成", "stopped": "已停止"}.get(st, "进行中")
        # 终态整段加粗（扫一眼就能扫到），进行中保持常规字重避开视觉噪声
        bold = "" if st == "running" else "font-weight:bold;"
        segs.append(f'<span style="color:{color};{bold}">{icon} <b>{raw}</b>·{mark}</span>')
    return f'{head}　{"｜".join(segs)}' if segs else text

try:
    import yzd_client as _yzd
    import task_registry as _tasks
    HAS_YZD = True
except ImportError:
    HAS_YZD = False


# ── 关键词文件读取（无第三方依赖，docx 走 zip+xml）────────────────
def read_keywords_from_file(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".txt", ".csv", ".md", ".json", ""):
        return p.read_text(encoding="utf-8", errors="ignore")
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=False, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None and str(cell).strip():
                        lines.append(str(cell).strip())
        wb.close()
        return "\n".join(lines)
    if ext == ".docx":
        with zipfile.ZipFile(p) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        lines = []
        for para in re.split(r"</w:p>", xml):
            texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", para, flags=re.S)
            line = "".join(texts)
            line = (line.replace("&amp;", "&").replace("&lt;", "<")
                        .replace("&gt;", ">").replace("&quot;", '"'))
            if line.strip():
                lines.append(line.strip())
        return "\n".join(lines)
    return p.read_text(encoding="utf-8", errors="ignore")


# ── 采集线程：各平台并行由 ParallelSession 内部各起一个线程 ─────────
class CrawlWorker(QThread):
    def __init__(self, session, keywords):
        super().__init__()
        self.session = session
        self.keywords = keywords

    def run(self):
        try:
            self.session.crawl(self.keywords, start_index=0)
        except Exception as e:
            self.session._log(f"❌ 采集线程异常：{e}")


# ── 自动保存线程：各平台合并写「原始 + 清洗后入库」两份 ──────────────
class AutoSaveWorker(QThread):
    saved = pyqtSignal(str, int)   # 路径, 条数
    failed = pyqtSignal(str)

    def __init__(self, session, rows, session_dir):
        super().__init__()
        self.session = session
        self.rows = rows
        self.session_dir = session_dir

    def run(self):
        try:
            ts_dir = Path(self.session_dir)
            ts_dir.mkdir(parents=True, exist_ok=True)
            raw = ts_dir / (_tasks.RAW_FILE_NAME if HAS_YZD else "企业数据_原始.xlsx")
            self._write_raw(raw, self.rows)
            cl = ts_dir / "企业数据_清洗后入库数据.xlsx"
            self._write_cleaned(cl, self.rows)
            self.saved.emit(str(raw), len(self.rows))
        except PermissionError:
            self.failed.emit("文件正被 Excel/WPS 打开占用，本次自动保存跳过，关闭后下轮补上")
        except Exception as e:
            self.failed.emit(f"自动保存失败：{e}")

    @staticmethod
    def _write_raw(path, rows):
        from openpyxl import Workbook
        from parallel_session import RAW_HEADERS
        from cdp_base import clean_cell
        tmp = str(path) + ".tmp"
        wb = Workbook(); ws = wb.active; ws.title = "B2B多平台企业数据"
        ws.append(RAW_HEADERS)
        for i, r in enumerate(rows, 1):
            # 过 clean_cell：采集到的地址/主营可能带控制字符，
            # 不过滤会让 openpyxl 抛 IllegalCharacterError 炸掉整轮自动保存。
            ws.append([i] + [clean_cell(r.get(h, "")) for h in RAW_HEADERS[1:]])
        wb.save(tmp)
        os.replace(tmp, str(path))

    def _write_cleaned(self, path, rows):
        clean_rows = self.session.clean_records(rows, log=lambda m: None)
        template = _cleaner.locate_template([self.session_dir, str(Path.home() / "Downloads")])
        tmp = str(path) + ".tmp"
        _cleaner.write_clean_excel(clean_rows, tmp, template_path=template, log=lambda m: None)
        os.replace(tmp, str(path))


# ══════════════════════════════════════════════
#  云展动力：账号绑定 / 客户库分组 / 导入
# ═══════════════════════════════════════════════
YZD_IMPORT_CHUNK = 5000


def _yzd_font(size=10, bold=False):
    return QFont("Microsoft YaHei" if sys.platform == "win32" else "Arial",
                 size, QFont.Bold if bold else QFont.Normal)


class YzdBindWorker(QThread):
    finished_bind = pyqtSignal(bool, str)

    def __init__(self, phone):
        super().__init__()
        self.phone = phone

    def run(self):
        try:
            _yzd.YzdClient().check_phone(self.phone)
            self.finished_bind.emit(True, self.phone)
        except _yzd.YzdError as e:
            self.finished_bind.emit(False, str(e))
        except Exception as e:
            self.finished_bind.emit(False, f"校验手机号异常：{e}")


class YzdGroupsWorker(QThread):
    finished_groups = pyqtSignal(int, list, str)

    def __init__(self, phone, foreign_flag, keyword=""):
        super().__init__()
        self.phone = phone
        self.foreign_flag = foreign_flag
        self.keyword = keyword

    def run(self):
        try:
            groups = _yzd.YzdClient(self.phone).list_groups(
                foreign_flag=self.foreign_flag, group_name=self.keyword)
            self.finished_groups.emit(self.foreign_flag, groups, "")
        except Exception as e:
            self.finished_groups.emit(self.foreign_flag, [], str(e))


class YzdCreateGroupWorker(QThread):
    finished_create = pyqtSignal(dict, str)

    def __init__(self, phone, group_name, foreign_flag):
        super().__init__()
        self.phone = phone
        self.group_name = group_name
        self.foreign_flag = foreign_flag

    def run(self):
        try:
            group = _yzd.YzdClient(self.phone).create_group(
                self.group_name, foreign_flag=self.foreign_flag)
            self.finished_create.emit(group, "")
        except Exception as e:
            self.finished_create.emit({}, str(e))


class YzdImportWorker(QThread):
    logged = pyqtSignal(str)
    progress = pyqtSignal(int, int, str)
    task_done = pyqtSignal(str, int, int, int, str)
    finished_all = pyqtSignal(list, str)

    def __init__(self, session, phone, jobs, group_id, foreign_flag, tmp_dir):
        super().__init__()
        self.session = session
        self.phone = phone
        self.jobs = jobs
        self.group_id = group_id
        self.foreign_flag = foreign_flag
        self.tmp_dir = tmp_dir

    def _records(self, job):
        if job.get("records") is not None:
            return job["records"]
        return _tasks.read_records_from_xlsx(job.get("raw_file", ""))

    def run(self):
        try:
            self._run_all()
        except Exception as e:
            self.finished_all.emit([], f"导入线程异常：{e}")

    def _run_all(self):
        total = len(self.jobs)
        results = []
        client = _yzd.YzdClient(self.phone, log=self.logged.emit)
        for idx, job in enumerate(self.jobs, 1):
            name = job.get("name") or f"任务{idx}"
            self.progress.emit(idx - 1, total, f"正在导入 {idx}/{total}：{name}")
            try:
                records = self._records(job)
            except Exception as e:
                self.logged.emit(f"⚠ 读取「{name}」数据失败：{e}")
                results.append((name, 0, 0, 0, f"读取数据失败：{e}"))
                self.task_done.emit(name, 0, 0, 0, f"读取数据失败：{e}")
                continue
            if not records:
                msg = "该任务没有原数据"
                self.logged.emit(f"⚠ 「{name}」跳过：{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue
            try:
                rows = self.session.clean_records(records, log=self.logged.emit)
                rows = _yzd.dedupe_rows(rows, log=self.logged.emit)
            except Exception as e:
                msg = f"清洗失败：{e}"
                self.logged.emit(f"⚠ 「{name}」{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue
            if not rows:
                msg = "清洗后没有可导入的记录"
                self.logged.emit(f"⚠ 「{name}」跳过：{msg}")
                results.append((name, 0, 0, 0, msg))
                self.task_done.emit(name, 0, 0, 0, msg)
                continue
            chunks = [rows[i:i + YZD_IMPORT_CHUNK] for i in range(0, len(rows), YZD_IMPORT_CHUNK)] or [[]]
            imp = sav = fail = 0
            err = ""
            for c_idx, chunk in enumerate(chunks, 1):
                tag = f"{name}（{c_idx}/{len(chunks)}）" if len(chunks) > 1 else name
                self.progress.emit(idx - 1, total, f"正在导入 {idx}/{total}：{tag} · {len(chunk)} 条")
                file_path = os.path.join(self.tmp_dir, f"yzd_import_{idx}_{c_idx}.xlsx")
                try:
                    written = _yzd.write_import_workbook(chunk, file_path, log=self.logged.emit)
                    if written == 0:
                        raise RuntimeError("清洗结果全为空行，没有可上传的数据")
                    self.logged.emit(f"⬆ {tag}：已生成上传文件 {written} 行，正在提交云展…")
                    res = client.import_excel(file_path, group_id=self.group_id,
                                              foreign_flag=self.foreign_flag)
                    imp += res["importTotal"]; sav += res["saveTotal"]; fail += res["failTotal"]
                    self.logged.emit(f"  ↳ 云展返回：导入 {res['importTotal']} / 成功 {res['saveTotal']} / 失败 {res['failTotal']}")
                except Exception as e:
                    err = str(e)
                    self.logged.emit(f"❌ {tag} 导入失败：{err}")
                    break
                finally:
                    try:
                        if os.path.exists(file_path):
                            os.remove(file_path)
                    except Exception:
                        pass
            results.append((name, imp, sav, fail, err))
            self.task_done.emit(name, imp, sav, fail, err)
        self.progress.emit(total, total, "导入结束")
        self.finished_all.emit(results, "")


class YzdProgressDialog(QDialog):
    def __init__(self, parent, title="正在导入云展客户库"):
        super().__init__(parent)
        self._allow_close = False
        self.setWindowTitle(title)
        self.setModal(True)
        self.setWindowFlags(Qt.Dialog | Qt.CustomizeWindowHint | Qt.WindowTitleHint)
        self.setMinimumWidth(460)
        layout = QVBoxLayout(self)
        layout.setSpacing(14)
        layout.setContentsMargins(24, 22, 24, 20)
        self.info_lbl = QLabel("准备中…")
        self.info_lbl.setFont(_yzd_font(11))
        self.info_lbl.setWordWrap(True)
        layout.addWidget(self.info_lbl)
        self.bar = QProgressBar()
        self.bar.setRange(0, 0)
        self.bar.setTextVisible(False)
        self.bar.setMinimumHeight(20)
        layout.addWidget(self.bar)
        self.detail_lbl = QLabel("导入期间请勿关闭程序，也不要执行其他操作。")
        self.detail_lbl.setFont(_yzd_font(9))
        self.detail_lbl.setStyleSheet(_muted_style())
        self.detail_lbl.setWordWrap(True)
        layout.addWidget(self.detail_lbl)

    def set_text(self, main, detail=""):
        self.info_lbl.setText(main)
        if detail:
            self.detail_lbl.setText(detail)

    def allow_close(self):
        self._allow_close = True

    def closeEvent(self, event):
        event.accept() if self._allow_close else event.ignore()


class YzdGroupDialog(QDialog):
    """「导入云展客户库」弹窗：分组选择 + 历史任务多选。"""
    MODE_LATEST = "latest"
    MODE_OTHERS = "others"
    HINT_IMPORT = "选完分组后点右下「导入最新」或「导入其他」"
    HINT_PICK = "选一个分组作为采集完自动导入的目标，然后点右下「选用该分组」"

    def __init__(self, phone, parent=None, pick_group_only=False):
        super().__init__(parent)
        self.phone = phone
        self._pick_only = bool(pick_group_only)
        self._hint_text = self.HINT_PICK if self._pick_only else self.HINT_IMPORT
        self.mode = ""
        self.chosen_group = None
        self.chosen_tasks = []
        self._groups = {0: None, 1: None}
        self._flag = 0
        self._worker = None
        self._create_worker = None
        self._pending_name = ""
        self._created_id = ""
        self._task_map = {}
        self.setWindowTitle("选择自动导入的客户库分组" if self._pick_only else "导入云展客户库")
        self.setMinimumSize(560, 460)
        self.resize(600, 500)
        self._init_ui()
        self._load_groups()

    def _init_ui(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(18, 16, 18, 16)
        root.setSpacing(10)
        self.stack = QStackedWidget()
        root.addWidget(self.stack, 1)
        self.stack.addWidget(self._build_group_page())
        self.stack.addWidget(self._build_task_page())

    def _build_group_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setSpacing(10)
        tab_row = QHBoxLayout()
        self.tab_btns = {}
        self.tab_group = QButtonGroup(self)
        self.tab_group.setExclusive(True)
        for flag, text in ((0, "国内客户库"), (1, "国外客户库")):
            btn = QPushButton(text)
            btn.setCheckable(True)
            btn.setChecked(flag == 0)
            btn.setMinimumHeight(32)
            btn.setStyleSheet(_tab_toggle_style())
            btn.clicked.connect(lambda _c, f=flag: self._switch_flag(f))
            self.tab_group.addButton(btn, flag)
            self.tab_btns[flag] = btn
            tab_row.addWidget(btn)
        tab_row.addStretch()
        self.refresh_btn = QPushButton("刷新")
        self.refresh_btn.setMinimumHeight(32)
        self.refresh_btn.setMinimumWidth(64)
        self.refresh_btn.clicked.connect(self._refresh_groups)
        tab_row.addWidget(self.refresh_btn)
        self.new_group_btn = QPushButton("+ 新建分组")
        self.new_group_btn.setMinimumHeight(32)
        self.new_group_btn.clicked.connect(self._create_group)
        tab_row.addWidget(self.new_group_btn)
        lay.addLayout(tab_row)

        search_row = QHBoxLayout()
        self.search_edit = QLineEdit()
        self.search_edit.setPlaceholderText("搜索分组名称（模糊匹配）")
        self.search_edit.setMinimumHeight(30)
        self.search_edit.returnPressed.connect(self._search)
        search_row.addWidget(self.search_edit, 1)
        self.search_btn = QPushButton("搜索")
        self.search_btn.setMinimumHeight(30)
        self.search_btn.setMinimumWidth(70)
        self.search_btn.clicked.connect(self._search)
        search_row.addWidget(self.search_btn)
        lay.addLayout(search_row)

        self.group_status_lbl = QLabel("正在加载分组…")
        self.group_status_lbl.setFont(_yzd_font(9))
        self.group_status_lbl.setStyleSheet(_muted_style())
        self.group_status_lbl.setWordWrap(True)
        lay.addWidget(self.group_status_lbl)

        self.group_list = QListWidget()
        self.group_list.setMinimumHeight(180)
        self.group_list.setSelectionMode(QAbstractItemView.SingleSelection)
        self.group_list.itemSelectionChanged.connect(self._update_hint)
        lay.addWidget(self.group_list, 1)

        self.hint_lbl = QLabel(self._hint_text)
        self.hint_lbl.setFont(_yzd_font(9))
        self.hint_lbl.setStyleSheet(_muted_style())
        self.hint_lbl.setWordWrap(True)
        lay.addWidget(self.hint_lbl)

        btn_row = QHBoxLayout()
        self.cancel_btn = QPushButton("取消")
        self.cancel_btn.setMinimumHeight(36)
        self.cancel_btn.setMinimumWidth(90)
        self.cancel_btn.clicked.connect(self.reject)
        btn_row.addWidget(self.cancel_btn)
        btn_row.addStretch()
        self.others_btn = QPushButton("导入其他")
        self.others_btn.setMinimumHeight(36)
        self.others_btn.setMinimumWidth(110)
        self.others_btn.clicked.connect(self._show_tasks)
        btn_row.addWidget(self.others_btn)
        self.latest_btn = QPushButton("导入最新")
        self.latest_btn.setMinimumHeight(36)
        self.latest_btn.setMinimumWidth(110)
        self.latest_btn.setDefault(True)
        self.latest_btn.clicked.connect(self._accept_latest)
        btn_row.addWidget(self.latest_btn)
        if self._pick_only:
            self.others_btn.setVisible(False)
            self.latest_btn.setText("选用该分组")
        lay.addLayout(btn_row)
        return page

    def _build_task_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setSpacing(10)
        top = QHBoxLayout()
        self.back_btn = QPushButton("← 返回")
        self.back_btn.setMinimumHeight(32)
        self.back_btn.clicked.connect(lambda: self.stack.setCurrentIndex(0))
        top.addWidget(self.back_btn)
        title = QLabel("选择要导入的采集任务（可多选，必选）")
        title.setFont(_yzd_font(11, True))
        top.addWidget(title, 1)
        self.all_chk = QCheckBox("全选")
        self.all_chk.toggled.connect(self._toggle_all)
        top.addWidget(self.all_chk)
        lay.addLayout(top)
        self.task_count_lbl = QLabel("共 0 个任务")
        self.task_count_lbl.setFont(_yzd_font(9))
        self.task_count_lbl.setStyleSheet(_muted_style())
        lay.addWidget(self.task_count_lbl)
        self.task_list = QListWidget()
        self.task_list.setMinimumHeight(200)
        self.task_list.setSelectionMode(QAbstractItemView.NoSelection)
        self.task_list.itemChanged.connect(lambda _i: self._update_task_hint())
        lay.addWidget(self.task_list, 1)
        self.task_hint_lbl = QLabel("至少选择一个任务")
        self.task_hint_lbl.setFont(_yzd_font(9))
        self.task_hint_lbl.setStyleSheet(_muted_style())
        self.task_hint_lbl.setWordWrap(True)
        lay.addWidget(self.task_hint_lbl)
        btn_row = QHBoxLayout()
        btn_row.addStretch()
        self.task_cancel_btn = QPushButton("取消")
        self.task_cancel_btn.setMinimumHeight(36)
        self.task_cancel_btn.clicked.connect(self.reject)
        btn_row.addWidget(self.task_cancel_btn)
        self.task_ok_btn = QPushButton("确认导入")
        self.task_ok_btn.setMinimumHeight(36)
        self.task_ok_btn.setMinimumWidth(110)
        self.task_ok_btn.clicked.connect(self._accept_others)
        btn_row.addWidget(self.task_ok_btn)
        lay.addLayout(btn_row)
        return page

    def _set_busy(self, busy):
        for w in (self.new_group_btn, self.search_btn, self.search_edit,
                  self.refresh_btn, self.latest_btn, self.others_btn):
            w.setEnabled(not busy)
        if busy:
            QApplication.setOverrideCursor(Qt.WaitCursor)
        else:
            QApplication.restoreOverrideCursor()

    def _switch_flag(self, flag):
        if flag == self._flag and self._groups[flag] is not None:
            return
        self._flag = flag
        cached = self._groups.get(flag)
        if cached is None:
            self._load_groups()
        else:
            self._render_groups(cached, "")

    def _search(self):
        self._groups[self._flag] = None
        self._load_groups()

    def _refresh_groups(self):
        self._groups[self._flag] = None
        self._load_groups()

    def _load_groups(self):
        if self._worker is not None and self._worker.isRunning():
            return
        keyword = self.search_edit.text().strip()
        self.group_status_lbl.setText("正在加载分组…")
        self._set_busy(True)
        self._worker = YzdGroupsWorker(self.phone, self._flag, keyword)
        self._worker.finished_groups.connect(self._on_groups)
        self._worker.start()

    def _on_groups(self, flag, groups, error):
        self._groups[flag] = groups if not error else None
        if flag != self._flag:
            self._set_busy(False)
            return
        self._render_groups(groups, error)
        self._set_busy(False)

    def _render_groups(self, groups, error):
        self.group_list.clear()
        if error:
            self.group_status_lbl.setText(f"分组加载失败：{error}")
            self.group_status_lbl.setStyleSheet(_danger_style())
            self._update_hint()
            return
        self.group_status_lbl.setStyleSheet(_muted_style())
        if not groups:
            keyword = self.search_edit.text().strip()
            self.group_status_lbl.setText(
                "没搜到分组，换个关键词试试" if keyword else
                "该账号下还没有分组，请先点「+ 新建分组」创建一个")
            self._update_hint()
            return
        self.group_status_lbl.setText(
            f"共 {len(groups)} 个分组，选中一个后点右下角导入" if not self._pick_only else
            f"共 {len(groups)} 个分组，选一个当自动导入目标")
        for g in groups:
            num = g.get("clientNum")
            text = f"{g.get('groupName', '')}"
            if num is not None:
                text += f"　（客户数 {num}）"
            item = QListWidgetItem(text)
            item.setData(Qt.UserRole, g)
            self.group_list.addItem(item)
        if self._created_id and self._select_group_by_id(self._created_id):
            self._created_id = ""
        self._update_hint()

    def _current_group(self):
        items = self.group_list.selectedItems()
        return items[0].data(Qt.UserRole) if items else None

    def _update_hint(self, error=False):
        if self._current_group():
            self.hint_lbl.setText("")
            self.hint_lbl.setStyleSheet(_muted_style())
        elif error:
            self.hint_lbl.setText("请先选择一个分组（必选）")
            self.hint_lbl.setStyleSheet(_danger_style())
        else:
            self.hint_lbl.setText(self._hint_text)
            self.hint_lbl.setStyleSheet(_muted_style())

    def _need_group(self):
        group = self._current_group()
        if group is None:
            self._update_hint(error=True)
            QMessageBox.warning(self, "提示", "请选择一个分组")
        return group

    def _create_group(self):
        if not self.phone:
            QMessageBox.warning(self, "提示", "请先绑定云展动力账号")
            return
        name, ok = QInputDialog.getText(self, "新建分组", "群组名称（最多 10 个字符）：")
        if not ok:
            return
        name = (name or "").strip()
        if not name:
            QMessageBox.warning(self, "提示", "群组名称不能为空！")
            return
        if len(name) > _yzd.GROUP_NAME_MAX:
            QMessageBox.warning(self, "提示", f"群组名称最多输入{_yzd.GROUP_NAME_MAX}个字符！")
            return
        if self._create_worker is not None and self._create_worker.isRunning():
            return
        self._set_busy(True)
        self.group_status_lbl.setText(f"正在创建分组「{name}」…")
        self._pending_name = name
        self._create_worker = YzdCreateGroupWorker(self.phone, name, self._flag)
        self._create_worker.finished_create.connect(self._on_group_created)
        self._create_worker.start()

    def _on_group_created(self, group, error):
        if error:
            self._set_busy(False)
            self.group_status_lbl.setText(f"新建分组失败：{error}")
            QMessageBox.warning(self, "新建分组失败", error)
            return
        self._groups = {0: None, 1: None}
        self._created_id = str(group.get("id") or "")
        self.search_edit.clear()
        self.group_status_lbl.setText(f"已创建分组「{group.get('groupName') or self._pending_name}」")
        self._load_groups()

    def _select_group_by_id(self, gid):
        for i in range(self.group_list.count()):
            item = self.group_list.item(i)
            data = item.data(Qt.UserRole) or {}
            if str(data.get("id")) == gid:
                self.group_list.setCurrentItem(item)
                return True
        return False

    def _accept_latest(self):
        group = self._need_group()
        if group is None:
            return
        self.mode = self.MODE_LATEST
        self.chosen_group = group
        self.accept()

    def _show_tasks(self):
        group = self._need_group()
        if group is None:
            return
        self.chosen_group = group
        self._fill_tasks()
        self.stack.setCurrentIndex(1)

    def _fill_tasks(self):
        self.task_list.clear()
        self.all_chk.blockSignals(True)
        self.all_chk.setChecked(False)
        self.all_chk.blockSignals(False)
        tasks = _tasks.sorted_for_import(_tasks.load_tasks())
        self._task_map = {}
        for t in tasks:
            name = t.get("name") or t.get("id") or "未命名任务"
            text = (f"{name}\n采集时间 {t.get('started_at') or '—'}"
                    f"　{t.get('count') or 0} 条　[{t.get('status') or '—'}]")
            item = QListWidgetItem(text)
            item.setFlags(item.flags() | Qt.ItemIsUserCheckable)
            item.setCheckState(Qt.Unchecked)
            item.setData(Qt.UserRole, t)
            self.task_list.addItem(item)
            self._task_map[id(item)] = t
        if not tasks:
            self.task_count_lbl.setText("暂无采集任务记录：先在主界面完成一次「开始采集 → 采集完成」")
        else:
            self.task_count_lbl.setText(f"共 {len(tasks)} 个任务（按采集时间倒序）")
        self._update_task_hint()

    def _toggle_all(self, checked):
        for i in range(self.task_list.count()):
            self.task_list.item(i).setCheckState(Qt.Checked if checked else Qt.Unchecked)

    def _checked_tasks(self):
        out = []
        for i in range(self.task_list.count()):
            item = self.task_list.item(i)
            if item.checkState() == Qt.Checked:
                task = item.data(Qt.UserRole)
                if task:
                    out.append(task)
        return out

    def _update_task_hint(self, error=False):
        n = len(self._checked_tasks())
        if n:
            self.task_hint_lbl.setText("")
            self.task_hint_lbl.setStyleSheet(_muted_style())
        elif error:
            self.task_hint_lbl.setText("请至少选择一个任务（必选）")
            self.task_hint_lbl.setStyleSheet(_danger_style())
        else:
            self.task_hint_lbl.setText("勾选要导入的采集任务（可多选）")
            self.task_hint_lbl.setStyleSheet(_muted_style())
        self.task_ok_btn.setText(f"确认导入（{n}）" if n else "确认导入")

    def _accept_others(self):
        tasks = self._checked_tasks()
        if not tasks:
            self._update_task_hint(error=True)
            QMessageBox.warning(self, "提示", "请至少选择一个任务")
            return
        self.mode = self.MODE_OTHERS
        self.chosen_tasks = tasks
        self.accept()

    def closeEvent(self, event):
        for w in (self._worker, self._create_worker):
            if w is not None and w.isRunning():
                w.wait(3000)
        event.accept()


# ── 主窗口 ─────────────────────────────────────────────────────
class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        # 标题只写工具名：平台清单在面板「③ 采集平台」里已逐站列出，
        # 塞进标题栏会被系统截断，反而看不清。
        self.setWindowTitle("B2B 多平台企业数据自动采集工具")
        # 小屏自适应：初始高度绝不超出可用屏高。写死尺寸在 13 寸笔电上会把
        # 底部日志裁到屏幕外（就是「输出日志看不到、得把窗口拉大」的另一半原因）。
        self.setMinimumSize(880, 620)
        try:
            avail = QApplication.primaryScreen().availableGeometry()
            self.resize(min(1000, max(880, avail.width() - 120)),
                        min(720, max(620, avail.height() - 120)))
        except Exception:
            self.resize(1000, 720)
        self.session = ParallelSession()
        self.worker = None
        self.autosave = None
        self._pending_save = False
        self._opening_browser = False
        self._paused = False
        self.keywords = []
        self._kw_source_file = ""
        self._co_source_file = ""   # 公司名单文件（企知道补采）
        self._opening_login = False  # 防重复点「打开企知道/1688」同时拉多个后台线程
        self.session_dir = None
        self._first_start = None
        if HAS_YZD:
            _tasks.set_desktop_dir(self._get_desktop_dir())
        self._yzd_phone = _tasks.load_account() if HAS_YZD else ""
        self._yzd_worker = None
        self._yzd_import_worker = None
        self._yzd_dialog = None
        self._yzd_progress = None
        self._yzd_importing = False
        self._yzd_jobs = []
        self._yzd_group_name = ""
        self._auto_import_group = None
        self._task_id = ""
        # 防采/验证码人工提醒弹窗（非模态，全局唯一）：只在被拦时存在，过验后关闭
        self._block_alert = None
        self._build_ui()
        self._connect_signals()
        self._refresh_yzd_status()
        # 自动保存双触发：①每满 AUTOSAVE_EVERY 条（见 _on_progress，采集驱动即时落盘）；
        # ②定时兜底（长时间单关键词深翻页也能按时间保存），另在开始/结束时各存一次。
        self.AUTOSAVE_EVERY = 20
        self._last_autosave_bucket = 0
        self._timer = QTimer(self)
        self._timer.timeout.connect(self.request_autosave)
        self._timer.start(3 * 60 * 1000)
        # 采集计时（对齐企知道/1688 工具）：秒级刷新；验证等待/手动暂停时接管计时，
        # 避免人工等待把平均耗时算虚高。
        self._crawl_timer = QTimer(self)
        self._crawl_timer.timeout.connect(self._update_timer)
        self._crawl_start_ts = 0.0
        self._crawl_elapsed = 0.0
        self._hold_from = 0.0
        self._held_seconds = 0.0

    def _get_desktop_dir(self):
        desktop = Path.home() / "Desktop"
        if desktop.exists():
            return desktop
        for env in ("OneDrive", "OneDriveConsumer", "OneDriveCommercial"):
            od = os.environ.get(env)
            if od:
                d = Path(od) / "Desktop"
                if d.exists():
                    return d
        try:
            desktop.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        return desktop

    def _build_ui(self):
        root = QWidget()
        lay = QVBoxLayout(root)
        # 统一节奏：分区之间留固定空隙、内边距收紧，避免控件贴边挤成一堆。
        lay.setContentsMargins(14, 12, 14, 12)
        lay.setSpacing(10)
        self.setCentralWidget(root)

        # ── 顶部双输入区：左＝关键词（5 平台并行搜索），右＝公司名单（企知道补采）──
        top = QHBoxLayout()
        top.setSpacing(10)

        gb_kw = QGroupBox("① 关键词（5 平台并行搜索）")
        gb_kw.setToolTip("上传文档或直接粘贴，自动识别多个关键词（换行/逗号/顿号分隔）。")
        v = QVBoxLayout(gb_kw)
        hr = QHBoxLayout()
        self.btn_file = QPushButton("导入关键词文件")
        self.btn_file.setToolTip("支持 txt / csv / xlsx / docx / md，自动识别里面的关键词。")
        self.btn_paste = QPushButton("从剪贴板粘贴")
        self.btn_parse = QPushButton("解析关键词")
        self.btn_smart = QPushButton("智能解析")
        self.btn_smart.setToolTip(
            "把「行业类别名」自动展开成一批「具体产品词」，帮你采到更多数据。\n"
            "例：「制冷与热泵核心部件」→ 制冷压缩机、冷凝器、蒸发器、膨胀阀、热泵机组……\n"
            "纯服务/流程词（研发、咨询、招投标、物流、运维等）会自动忽略，因为站内搜索没有对应商品。\n"
            "注意：解析结果由内置行业知识库生成，仅供参考，不一定完全符合你的需求，请检查后再采集。")
        hr.addWidget(self.btn_file); hr.addWidget(self.btn_paste)
        hr.addWidget(self.btn_parse); hr.addWidget(self.btn_smart)
        hr.addStretch(1)
        v.addLayout(hr)
        self.kw_text = QPlainTextEdit()
        self.kw_text.setPlaceholderText("每行一个关键词（也可用逗号/顿号分隔），例如：太阳能路灯、光伏储能电池、LED显示屏")
        self.kw_text.setMaximumHeight(96)
        v.addWidget(self.kw_text)
        self.lbl_kw = QLabel("识别到 0 个关键词")
        self.lbl_kw.setStyleSheet(_muted_style())
        # 解析完会拼出「识别到 12 个关键词：A、B、C…」长串：不折行就会把窗口
        # 最小宽度顶宽，小屏上直接得横向拖动才能看到开始采集按钮。
        self.lbl_kw.setWordWrap(True)
        v.addWidget(self.lbl_kw)
        top.addWidget(gb_kw, 3)

        gb_co = QGroupBox("② 公司名单（企知道逐家补采·可选）")
        gb_co.setToolTip(
            "可选：上传一份公司名文件（xlsx/csv/txt，有一列是公司名即可），企知道会逐家模拟人工查询补采。\n"
            "开始采集后任务文件夹里会生成「待采集公司名单.xlsx」（与原数据/清洗入库数据同层级）；\n"
            "关键词采集清洗掉的无联系方式废数据会自动回流进这份名单、被企知道补采（变废为宝）。\n"
            "采过/查不到的公司名会自动剔除，不重复采集。")
        vc = QVBoxLayout(gb_co)
        hc = QHBoxLayout()
        self.btn_co_file = QPushButton("选择公司名单文件")
        self.btn_co_file.setToolTip("选择需补采的公司名文件（xlsx/csv/txt，自动识别公司名列）。")
        self.btn_open_qzd = QPushButton("打开企知道")
        self.btn_open_qzd.setToolTip("企知道需登录后才能补采。点此在共享浏览器打开企知道首页，请完成登录。")
        hc.addWidget(self.btn_co_file); hc.addWidget(self.btn_open_qzd)
        hc.addStretch(1)
        vc.addLayout(hc)
        self.lbl_co = QLabel("未选择公司名单（可选：不传也能靠废数据回流自动补采）")
        self.lbl_co.setStyleSheet(_muted_style())
        self.lbl_co.setWordWrap(True)
        vc.addWidget(self.lbl_co)
        self.chk_waste = QCheckBox("启用企知道补采（废数据自动回流·挂机过夜推荐）")
        self.chk_waste.setChecked(True)
        self.chk_waste.setToolTip(
            "勾选后：① 上传的公司名单会逐家补采；② 关键词采集清洗掉的无联系方式废数据自动回流名单并被补采。\n"
            "企知道通道为看门狗语义——名单为空时休眠、零资源占用；有数据才附加浏览器查询。\n"
            "需先登录企知道（点「打开企知道」）。不勾则废数据仍会累积进名单表，供下次开启时使用。")
        vc.addWidget(self.chk_waste)
        top.addWidget(gb_co, 2)

        lay.addLayout(top)

        gb_site = QGroupBox("③ 采集平台（关键词并行）")
        gb_site.setToolTip(
            "任一站点命中防采验证都会弹窗提醒并由看门狗自动续采。\n"
            "八方资源网 / 顺企网 / 爱采购为纯 HTTP，不占浏览器标签（看不到它们的标签是正常的）；\n"
            "中国制造网 / 1688 各占一个浏览器标签。1688 需登录，默认不勾选。\n\n"
            + "\n\n".join(SITE_TIPS.get(s, "") for s in SITE_ORDER))
        vs = QVBoxLayout(gb_site)
        hs = QHBoxLayout()
        self.site_checks = {}
        for slug in SITE_ORDER:
            chk = QCheckBox(SITE_NAMES[slug])
            # 1688 有登录限制，默认不勾（勾选时弹窗提示登录）
            chk.setChecked(slug != "1688")
            chk.setToolTip(SITE_TIPS.get(slug, ""))
            if slug == "1688":
                chk.toggled.connect(lambda on, s=slug: self._on_site_toggled(s, on))
            hs.addWidget(chk)
            self.site_checks[slug] = chk
        hs.addStretch(1)
        vs.addLayout(hs)
        lay.addWidget(gb_site)

        gb_p = QGroupBox("④ 采集参数")
        # 分区标题只留短句，三个开关的机制差异集成交互说明挂在分区本身：
        # 悬停就能看清「勾了会少发多少请求」，不用把长句子平铺到界面上。
        gb_p.setToolTip(
            "采集间隔：相邻请求等待秒数。各站共用同一个出口 IP，挂着过夜建议调到 6~10 秒更安全。\n"
            "翻页一律不限制：一直翻到没有新数据，再自动切下一个关键词。\n"
            "过滤行业不相关商家：拿关键词和公司名+主营比对，明显无关的连详情页都不进（省请求=提速+防封）；"
            "信号取不到时一律放行不误杀。\n"
            "增量模式：跨运行记住已采到/已判不相关的商家，再跑同一批关键词命中即跳过；"
            "因拦截没采成的不记，下轮仍重试。\n"
            "后台运行：开始采集后自动最小化共享 Chrome；命中验证码会自动还原窗口递到眼前，过完再最小化。")
        vp = QVBoxLayout(gb_p)
        hp = QHBoxLayout()
        hp.addWidget(QLabel("采集间隔(秒)："))
        self.spin_delay = QDoubleSpinBox()
        self.spin_delay.setRange(2.0, 60.0); self.spin_delay.setValue(5.0); self.spin_delay.setSingleStep(1.0)
        # 页数不再对外暴露：各平台一律「翻到没新数据为止」（max_pages 固定 0），
        # 旧版那个「不限页数 / 限定时最大页数」开关实际从不关，留着只增加认知负担。
        self.spin_delay.setToolTip("相邻请求的等待秒数。各站共享同一个 IP，挂着过夜建议调到 6~10 秒更安全。\n"
                                   "翻页一律不限制：一直翻到没有新数据再自动切下一个关键词。")
        hp.addWidget(self.spin_delay)
        hp.addSpacing(10)
        self.chk_guard = QCheckBox("过滤行业不相关商家")
        self.chk_guard.setChecked(True)
        self.chk_guard.setToolTip("用关键词与「公司名+主营」做相关性比对，明显无关的连详情页都不进（省请求=提速+防封）。"
                                  "信号取不到时一律放行不误杀。")
        self.chk_incr = QCheckBox("增量模式·跳过历史已采商家")
        self.chk_incr.setChecked(True)
        self.chk_incr.setToolTip("跨运行记住已采到/已判不相关的商家，再跑同一批关键词命中即跳过、不重复请求"
                                 "（省请求=防封），只补新出现的商家；因拦截没采成的不记，下轮仍重试。")
        self.chk_bg = QCheckBox("后台运行·最小化浏览器窗口")
        # 默认勾选：多平台并行时用户根本不需要盯浏览器，挂着自动最小化才能真「无人值守」；
        # 命中验证码仍会自动还原窗口把验证递到眼前。
        self.chk_bg.setChecked(True)
        self.chk_bg.setToolTip("开始采集后自动最小化共享 Chrome（八方资源网/顺企网/爱采购本就不开标签）。\n"
                               "命中验证码时会自动还原窗口提醒人工过验证，过后重新最小化；\n"
                               "启动浏览器已带防后台节流开关，最小化不影响页面渲染。\n"
                               "注：对启动前已在跑的旧 Chrome 进程无效（开关只能启动时传入），会改用前置窗口方式。")
        hp.addWidget(self.chk_guard); hp.addWidget(self.chk_incr); hp.addWidget(self.chk_bg)
        hp.addStretch(1)
        self.btn_clear_seen = QPushButton("清空历史种子库")
        self.btn_clear_seen.setToolTip("删除各站跨运行去重记录，让下次对同一批关键词从头重采。")
        hp.addWidget(self.btn_clear_seen)
        vp.addLayout(hp)
        lay.addWidget(gb_p)

        hb = QHBoxLayout()
        self.btn_open = QPushButton("打开浏览器")
        self.btn_start = QPushButton("开始采集")
        self.btn_pause = QPushButton("暂停")
        self.btn_stop = QPushButton("停止")
        self.btn_clean = QPushButton("清洗数据")
        self.btn_clean.setToolTip(
            "立即把当前已采数据跑一遍清洗：生成/刷新任务文件夹里的「原始」与「清洗后入库数据」两份 Excel，\n"
            "无联系方式的废数据会自动回流到「待采集公司名单」（启用企知道补采时会被逐家补采）。\n"
            "采集过程中每满 20 条也会自动清洗保存一次，此按钮用于随时手动触发。")
        self.btn_export = QPushButton("导出 Excel")
        self.btn_reset = QPushButton("开始新任务")
        self.btn_reset.setToolTip(
            "换一批关键词从零重采时用：清空本次已采数据、各站进度账目、进度条/计时/看板，\n"
            "并释放当前数据文件夹——下次点「开始采集」会新建一个以时间戳命名的文件夹保存新数据。\n"
            "复位前先把剩余数据同步落盘（不满一个自动保存批次的尾部也不丢）。\n"
            "跨运行历史种子库默认保留（省请求=防封），需要连去重记录一起清就在对话框里勾选。")
        # 主操作唯一化：只有「开始采集」吃强调色，其余保持默认样式（满屏彩色=没有重点）；
        # 按「准备 / 跑 / 产出」三组拉开间距，一眼能看出动作顺序。
        self.btn_start.setObjectName("primary")
        self.btn_start.setMinimumWidth(120)
        for b in (self.btn_open, self.btn_start, self.btn_pause, self.btn_stop,
                  self.btn_clean, self.btn_export, self.btn_reset):
            b.setMinimumHeight(36)
        hb.addWidget(self.btn_open)
        hb.addSpacing(14)
        for b in (self.btn_start, self.btn_pause, self.btn_stop):
            hb.addWidget(b)
        hb.addStretch(1)
        for b in (self.btn_clean, self.btn_export, self.btn_reset):
            hb.addWidget(b)
        self.btn_pause.setEnabled(False); self.btn_stop.setEnabled(False)
        lay.addLayout(hb)

        self.auto_import_check = QCheckBox("采集完自动导入云展客户库")
        self.auto_import_check.setToolTip(
            "勾选时先让你选一个客户库分组；采集完成后自动把本次数据上传到该分组并在日志提示结果。\n"
            "需先绑定云展动力账号。手动点「停止」不触发自动导入。")
        self.auto_import_check.toggled.connect(self._on_auto_import_toggled)

        # 绑定状态 + 自动导入勾选 + 两个动作合并成一行，不再各占一块（省下一行高度给日志）
        yzd_group = QGroupBox("云展动力客户库")
        yzd_group.setToolTip(
            "把本次采到的企业按联系人口径清洗后导入云展动力客户库指定分组。\n"
            "需先绑定云展动力账号；手动点「导入客户库」可选分组和历史任务，"
            "勾选「采集完自动导入」则先选一个分组、跑完自动上传并在日志提示结果。\n"
            "手动点「停止」不触发自动导入（中途停手的数据不会静默上传）。")
        yl = QHBoxLayout(yzd_group)
        self.yzd_status_lbl = QLabel("未绑定云展动力账号")
        self.yzd_status_lbl.setFont(_yzd_font(10))
        self.yzd_status_lbl.setMinimumWidth(180)
        yl.addWidget(self.yzd_status_lbl)
        yl.addSpacing(12)
        yl.addWidget(self.auto_import_check)
        yl.addStretch()
        self.bind_yzd_btn = QPushButton("绑定账号")
        self.bind_yzd_btn.setMinimumHeight(30)
        self.bind_yzd_btn.setToolTip("绑定云展动力手机号账号（导入客户库前需先绑定）")
        self.bind_yzd_btn.clicked.connect(self._bind_yzd_account)
        yl.addWidget(self.bind_yzd_btn)
        self.import_yzd_btn = QPushButton("导入客户库")
        self.import_yzd_btn.setMinimumHeight(30)
        self.import_yzd_btn.clicked.connect(self._open_yzd_import)
        yl.addWidget(self.import_yzd_btn)
        lay.addWidget(yzd_group)

        self.banner = QLabel("")
        self.banner.setVisible(False)
        self.banner.setWordWrap(True)
        lay.addWidget(self.banner)

        prog_row = QHBoxLayout()
        self.progress = QProgressBar()
        self.progress.setMinimumHeight(26)
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setFormat("就绪")
        prog_row.addWidget(self.progress, 1)
        self.timer_lbl = QLabel("⏱ 00:00:00")
        f = self.timer_lbl.font(); f.setBold(True); f.setPointSize(11)
        self.timer_lbl.setFont(f)
        prog_row.addWidget(self.timer_lbl)
        lay.addLayout(prog_row)
        self.lbl_stats = QLabel("就绪")
        # 状态行支持富文本：每站按状态加图标+颜色+文字（✅绿已完成 / ■橙已停止 / ●蓝进行中），
        # 三通道编码不只依赖颜色，色觉障碍同样可辨（任务实况仪表惯例）
        self.lbl_stats.setTextFormat(Qt.AutoText)
        lay.addWidget(self.lbl_stats)
        self.lbl_board = QLabel("")
        self.lbl_board.setWordWrap(True)
        self.lbl_save_info = QLabel("")
        # 次要信息（看板明细/保存回显）统一小一号 + muted 色：与主状态行拉开层级，
        # 一进门只需先读「已采多少家」，其余悬停/扫读即可。
        for _lbl in (self.lbl_board, self.lbl_save_info, self.lbl_kw):
            _f = _lbl.font(); _f.setPointSize(10); _lbl.setFont(_f)
        self.lbl_board.setStyleSheet(_muted_style())
        # 空标签在 Qt 里照样占一整行高：未开始采集/未落盘时不该占位。
        # 有内容才显形，把省下的行高全让给输出日志（小屏上日志多一行是一行）。
        self.lbl_board.setVisible(False)
        lay.addWidget(self.lbl_board)
        self.lbl_save_info.setStyleSheet(_muted_style())
        self.lbl_save_info.setVisible(False)
        lay.addWidget(self.lbl_save_info)

        self.log = QTextEdit(); self.log.setReadOnly(True)
        self.log.setFont(QFont("Menlo", 10) if sys.platform != "win32" else QFont("Consolas", 9))
        # 日志是小屏电脑的底线：上面所有控件都压缩过，这里给硬性最小高度，
        # 绝不再出现「日志被顶出屏幕、得把窗口拉大才能看到输出」。
        self.log.setMinimumHeight(180)
        # 性能（同事反馈挂机后电脑卡）：过夜采集日志能刷上万行，QTextEdit 无上限
        # 累积会让文档越来越大、每次 append 重排越来越慢、内存持续上涨。
        # 设最大块数：只保留最近 N 行，超出自动从头丢弃，渲染/内存恒定。
        self.log.document().setMaximumBlockCount(3000)
        lay.addWidget(self.log, 1)

    def _connect_signals(self):
        s = self.session.signals
        s.log_signal.connect(self._on_log)
        s.progress_signal.connect(self._on_progress)
        s.kw_progress_signal.connect(self._on_kw_progress)
        s.stats_signal.connect(self._render_stats)
        s.crawl_finished_signal.connect(self._on_finished)
        s.site_finished_signal.connect(lambda name: self._on_log(f"🏁 {name} 采集完成"))
        s.blocked_signal.connect(self._on_blocked)
        s.block_cleared_signal.connect(self._on_block_cleared)
        if hasattr(s, "login_required_signal"):
            s.login_required_signal.connect(lambda m: self._on_log("🔐 " + m))
        self.btn_file.clicked.connect(self._on_import_file)
        self.btn_paste.clicked.connect(self._on_paste)
        self.btn_parse.clicked.connect(self._on_parse)
        self.btn_smart.clicked.connect(self._on_smart_parse)
        self.btn_co_file.clicked.connect(self._on_import_company_file)
        self.btn_open_qzd.clicked.connect(lambda: self._open_site_login("qzd"))
        self.btn_open.clicked.connect(self._on_open_browser)
        self.btn_start.clicked.connect(self._on_start)
        self.btn_pause.clicked.connect(self._on_pause)
        self.btn_stop.clicked.connect(self._on_stop)
        self.btn_clean.clicked.connect(self._on_clean_now)
        self.btn_export.clicked.connect(self._on_export)
        self.btn_clear_seen.clicked.connect(self._on_clear_seen_store)
        self.btn_reset.clicked.connect(self._on_reset_task)

    def _on_log(self, msg):
        self.log.append(msg)
        self.log.verticalScrollBar().setValue(self.log.verticalScrollBar().maximum())

    # ── 进度条 / 计时 / 看板 ──────────────────────────────
    @staticmethod
    def _format_time(seconds):
        seconds = int(max(0, seconds))
        return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"

    def _update_bar(self, prefix="采集中"):
        pct = getattr(self, "_bar_pct", 0)
        cur = getattr(self, "_bar_companies", 0)
        self.progress.setValue(pct)
        # 时间只给右侧独立计时器显示：旧版进度条文本里也嵌一份⏱，同一屏出现两个
        # 不同数值的时间（一个是累计耗时、一个是刷新时刻），看着像数据对不上。
        self.progress.setFormat(f"{prefix} {pct}%｜已采 {cur} 家")

    def _on_kw_progress(self, _slug, done_units, total_units):
        self._bar_pct = int(done_units / total_units * 100) if total_units else 0
        self._update_bar()

    def _refresh_board(self):
        """量化看板压成**一行**总计（成功/无联系方式/失败可重采/早过滤/详情不相关），
        逐站明细移到本行 tooltip 悬停可查。旧版每站一行共 5 行、且写死深色文字，
        在深色系统上整块隐形却照样占位，把输出日志挤到屏幕外。"""
        try:
            stats = self.session.site_stats()
        except Exception:
            return
        if not stats:
            return
        keys = ("got", "ok", "pages", "listed", "no_contact", "failed", "early", "filtered")
        tot = {k: sum(s[k] for s in stats.values()) for k in keys}
        self.lbl_board.setText(
            f"明细：成功 {tot['ok']}｜无联系方式 {tot['no_contact']}｜失败可重采 {tot['failed']}｜"
            f"早过滤 {tot['early']}｜详情不相关 {tot['filtered']}｜翻 {tot['pages']} 页见 {tot['listed']} 条")
        self.lbl_board.setToolTip("\n".join(
            f"{name}：采到 {s['got']} 家（成功 {s['ok']}）｜翻 {s['pages']} 页见 {s['listed']} 条｜"
            f"无联系方式 {s['no_contact']}｜失败/可重采 {s['failed']}｜早过滤 {s['early']}｜"
            f"详情不相关 {s['filtered']}" for name, s in stats.items()))
        self.lbl_board.setVisible(True)

    def _render_stats(self, text, site_status=None):
        """状态行渲染：交给纯函数 build_stats_html（图标+颜色双通道），降级纯文本不报错。"""
        self.lbl_stats.setText(build_stats_html(text, site_status))

    def _on_progress(self, cur, _total):
        self._bar_companies = cur
        # 采集驱动自动保存：每满 AUTOSAVE_EVERY 条触发一次（与 3 分钟定时兜底双触发）
        bucket = cur // self.AUTOSAVE_EVERY
        if bucket > self._last_autosave_bucket:
            self._last_autosave_bucket = bucket
            self.request_autosave()
        self._update_bar()
        self._refresh_board()

    def _show_banner(self, text, bg, fg):
        self.banner.setText(text)
        self.banner.setStyleSheet(f"background:{bg};color:{fg};padding:8px;border-radius:4px;font-weight:bold;")
        self.banner.setVisible(True)

    def _hide_banner(self):
        self.banner.setVisible(False)

    def _on_blocked(self, reason):
        """被拦：醒目横幅 + 独立弹窗提醒 + 接管计时（验证等待不计入采集耗时）+
        后台化时自动还原窗口，把验证码递到用户眼前（人在环不变）。
        弹窗必须显式弹：仅靠顶部横幅在窗口后台/最小化时用户根本看不到，
        就会出现「验证码一直拦着但没人过、关键词进度空转」的白跑现象。"""
        self._show_banner(
            f"⚠ 触发安全验证/风控（{reason}）：请在弹出的 Chrome 窗口里手动完成验证，通过后会自动继续采集。",
            "#f8d7da", "#721c24")
        self._show_block_alert(reason)
        self._hold_timer(True)
        if self.chk_bg.isChecked():
            threading.Thread(target=lambda: self.session.set_browser_background(False),
                             daemon=True).start()

    def _show_block_alert(self, reason):
        """防采/验证码人工提醒弹窗：非模态（模态会阻塞主窗口事件循环，用户反而不
        能操作），全局只留一个实例，多站同时被拦时只更新文案不叠窗。"""
        box = getattr(self, "_block_alert", None)
        if box is None:
            box = QMessageBox(self)
            box.setWindowTitle("需要人工过验证")
            box.setIcon(QMessageBox.Warning)
            box.setStandardButtons(QMessageBox.Ok)
            box.setModal(False)
            self._block_alert = box
        box.setText(
            "⚠ 站点触发防采拦截，已在 Chrome 里打开对应页面。\n\n"
            f"{reason}\n\n"
            + ("本次是 IP 级限流/冷却：页面里没有滑块就不需要操作（别对着页面干找），\n"
               "确实有滑块/验证码就顺手过完。\n"
               if "HTTP" in (reason or "") else
               "请切到那个 Chrome 窗口，把滑块/验证码手动过完。\n")
            + "看门狗每 15 秒实测一次「能不能真拿到数据」，确认后自动续采，\n"
              "不需要回来点任何按钮（关掉本窗也不会停等待）。")
        box.show()
        box.raise_()
        box.activateWindow()

    def _close_block_alert(self):
        box = getattr(self, "_block_alert", None)
        if box is not None:
            box.hide()
            box.deleteLater()
            self._block_alert = None

    def _on_block_cleared(self, platform):
        self._hide_banner()
        self._close_block_alert()
        self._hold_timer(False)
        if self.chk_bg.isChecked():
            threading.Thread(target=lambda: self.session.set_browser_background(True),
                             daemon=True).start()

    # ── 采集计时 ──────────────────────────────────────
    def _update_timer(self):
        if self._crawl_start_ts:
            self._crawl_elapsed = time.time() - self._crawl_start_ts - self._held_seconds
        self.timer_lbl.setText(f"⏱ {self._format_time(self._crawl_elapsed)}")
        self._update_bar()

    def _start_timer(self):
        self._crawl_start_ts = time.time()
        self._crawl_elapsed = 0.0
        self._hold_from = 0.0
        self._held_seconds = 0.0
        self._crawl_timer.start(1000)

    def _stop_timer(self):
        self._crawl_timer.stop()
        self._update_timer()

    def _hold_timer(self, on):
        """验证等待/手动暂停期间接管计时：暂停段不计入采集耗时（否则人工等待
        会把平均耗时算虚高，与企知道工具同口径）。"""
        if not hasattr(self, "_held_seconds"):
            self._held_seconds = 0.0
        if on:
            if not self._hold_from:
                self._hold_from = time.time()
        elif self._hold_from:
            self._held_seconds += time.time() - self._hold_from
            self._hold_from = 0.0

    def _enabled_sites(self):
        return [slug for slug in SITE_ORDER if self.site_checks[slug].isChecked()]

    # ── 关键词 ─────────────────────────────────────────────────
    def _on_import_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择关键词文件", "", "支持的文件 (*.txt *.csv *.xlsx *.xlsm *.docx *.md);;所有文件 (*)")
        if not path:
            return
        try:
            text = read_keywords_from_file(path)
            self.kw_text.setPlainText(text)
            self._kw_source_file = path
            self._on_parse()
            self._on_log(f"已导入文件：{path}")
        except Exception as e:
            QMessageBox.warning(self, "导入失败", f"读取文件出错：{e}")

    def _on_paste(self):
        t = QApplication.clipboard().text()
        if t:
            self.kw_text.appendPlainText(t)
            self._on_parse()

    def _on_parse(self):
        self.keywords = ParallelSession.parse_keywords(self.kw_text.toPlainText())
        self.lbl_kw.setText(f"识别到 {len(self.keywords)} 个关键词：" +
                            ("、".join(self.keywords[:12]) + ("…" if len(self.keywords) > 12 else "")
                             if self.keywords else "无"))
        self._warn_abstract_keywords(self.keywords)

    # 环节/服务/流程类词（非具体产品词）：B2B 站内搜索几乎无商品与之对应，
    # 实测这类词每词只能翻出 0~2 家，是「采了很多关键词却只有几百条数据」的主因。
    _ABSTRACT_KW_MARKS = ("研发", "标准支持", "咨询", "方案设计", "招投标", "招标",
                          "交付验收", "客户开发", "点位开发", "工程配套", "系统集成",
                          "安装调试", "现场安装", "跨境供应链", "解决方案", "关键词")

    def _warn_abstract_keywords(self, keywords):
        """关键词质量提醒（只提醒、不删改）：把看起来是「行业报告目录/价值链环节名」
        而非「具体产品词」的关键词列出来，建议换成产品词，避免用户白采。
        低误报：只命中明确的环节/服务/流程词，不碰含材料/部件/设备等的品类词。"""
        abstract = [k for k in keywords
                    if any(m in k for m in self._ABSTRACT_KW_MARKS)]
        if not abstract:
            return
        show = "、".join(abstract[:8]) + ("…" if len(abstract) > 8 else "")
        self._on_log(
            f"⚠ 关键词质量提醒：{len(abstract)}/{len(keywords)} 个词像是「环节/服务名」而非具体产品词（{show}）。\n"
            "   B2B 站内搜索按产品/商品匹配，这类抽象词每词往往只翻出 0~2 家，是数据量偏少的主因。\n"
            "   建议换成具体产品词（例：把「制冷与热泵核心部件」拆成「制冷压缩机」「冷凝器」「蒸发器」「膨胀阀」「热泵机组」），数据量可提升一个量级。\n"
            "   也可直接点「智能解析」按钮，自动把这些类别名展开成产品词。")

    # 「智能解析」结果提示语（润色后·通俗易懂）：内置知识库展开，结果仅供参考，
    # 不一定符合用户预期，务必让用户先检查再采集——放在界面标签 + 日志各一份。
    _SMART_HINT = ("智能解析已把行业类别名自动换成常见产品词，结果仅供参考，"
                   "不一定完全符合你的需求，请对照下方列表检查、增删后再开始采集。")

    def _on_smart_parse(self):
        """智能解析：把「行业报告目录/价值链环节名」按内置知识库展开成具体产品词，
        结果写回输入框供用户检查编辑，并给出「仅供参考、请检查」的提示。
        与「解析关键词」的区别：后者只做切分清洗、不改词；本功能会做语义展开。"""
        raw = self.kw_text.toPlainText()
        if not raw.strip():
            QMessageBox.information(self, "请先输入关键词",
                                    "智能解析需要先有内容：请粘贴或导入一批行业类别名/关键词，再点「智能解析」。")
            return
        try:
            rep = _kwexp.expand_keywords(raw)
        except Exception as e:
            self._on_log(f"⚠ 智能解析失败：{str(e).splitlines()[0][:120]}")
            return
        kws = rep["keywords"]
        if not kws:
            self._on_log("⚠ 智能解析：没识别到可展开的内容，请检查输入。")
            return
        # 写回输入框（一行一个），方便用户直接检查、增删
        self.kw_text.setPlainText("\n".join(kws))
        self.keywords = list(kws)
        self.lbl_kw.setText(
            f"智能解析完成：得到 {len(kws)} 个产品词（展开 {len(rep['expanded'])} 个类别"
            f"{('、忽略 ' + str(len(rep['dropped'])) + ' 个服务/流程词') if rep['dropped'] else ''}）。"
            f"⚠ {self._SMART_HINT}")
        # 日志给出可核对的明细：展开了哪些类别、忽略了哪些词、保留了哪些原词
        self._on_log(f"🪄 智能解析：{len(kws)} 个产品词（由 {len(rep['expanded'])} 个类别展开）。")
        for seg, name, added in rep["expanded"]:
            self._on_log(f"   ・「{seg}」→ 归入【{name}】，新增 {added} 个产品词")
        if rep["kept"]:
            self._on_log(f"   ・原样保留（本就是产品词）：{'、'.join(rep['kept'][:20])}"
                         + ("…" if len(rep["kept"]) > 20 else ""))
        if rep["dropped"]:
            self._on_log(f"   ・已忽略（纯服务/流程词，站内无对应商品）：{'、'.join(rep['dropped'])}")
        self._on_log(f"⚠ {self._SMART_HINT}")

    # ── 公司名单（企知道补采）/ 登录页 ───────────────────
    def _on_import_company_file(self):
        """选择需补采的公司名文件：先预览识别到多少家（真正入列在开始采集时
        由 session.bind_pending 完成，那里才知道任务文件夹）。"""
        path, _ = QFileDialog.getOpenFileName(
            self, "选择公司名单文件", "", "支持的文件 (*.xlsx *.xlsm *.csv *.txt);;所有文件 (*)")
        if not path:
            return
        self._co_source_file = path
        fname = Path(path).name
        try:
            from pending_companies import extract_company_names
            n = len(extract_company_names(path))
        except Exception as e:
            n = -1
            self._on_log(f"⚠ 公司名单预览失败：{str(e).splitlines()[0][:120]}")
        if n > 0:
            self.lbl_co.setText(f"已选择：{fname}（识别到 {n} 家公司，开始采集后写入待采名单）")
            self._on_log(f"📋 已选择公司名单：{fname}，识别到 {n} 家公司")
            # 传了名单就是想用企知道补采：自动勾上补采开关（免用户漏勾）
            if not self.chk_waste.isChecked():
                self.chk_waste.setChecked(True)
        elif n == 0:
            self.lbl_co.setText(f"已选择：{fname}（未识别到公司名，请确认有一列是公司名）")
            self._on_log(f"⚠ {fname} 里没识别到公司名（需有一列是公司名）")
        else:
            self.lbl_co.setText(f"已选择：{fname}")

    def _on_site_toggled(self, slug, checked):
        """勾选 1688 时弹窗提示需登录，确认后打开 1688 登录页（通俗易懂的话术）。"""
        if slug != "1688" or not checked:
            return
        ret = QMessageBox.information(
            self, "1688 需要先登录",
            "1688 有登录限制，登录后才能采集（所以默认没帮你勾选）。\n\n"
            "点「确定」会立即在共享浏览器里打开 1688，请你先看一眼是否已登录：\n"
            "・已登录 → 直接关掉提示，开始采集即可；\n"
            "・未登录 → 在打开的 1688 页面扫码/输账号登录，登录好再开始采集。\n\n"
            "点「取消」则暂不打开，你可以稍后自己在浏览器登录。",
            QMessageBox.Ok | QMessageBox.Cancel, QMessageBox.Ok)
        if ret == QMessageBox.Ok:
            self._open_site_login("1688")

    def _open_site_login(self, slug):
        """打开 1688/企知道 登录页：附加浏览器有网络 IO，一律走后台线程，
        绝不在主线程同步启动（否则窗口会「未响应」）。回控只走信号/主线程轮询。"""
        if getattr(self, "_opening_login", False):
            return
        label = SITE_NAMES.get(slug, slug)
        self._opening_login = True
        btn = self.btn_open_qzd if slug == "qzd" else None
        if btn:
            btn.setEnabled(False)
        self._on_log(f"正在打开{label}登录页（首次会拉起共享浏览器，请稍候）…")

        def task():
            try:
                ok, msg = self.session.open_login_page(slug)
                self.session._log(("✅ " if ok else "⚠ ") + msg)
            except Exception as e:
                self.session._log(f"⚠ 打开{label}失败：{str(e).splitlines()[0][:120]}")
            finally:
                self._opening_login = False

        threading.Thread(target=task, daemon=True).start()

        def poll():
            if getattr(self, "_opening_login", False):
                QTimer.singleShot(400, poll)
            elif btn:
                btn.setEnabled(True)
        QTimer.singleShot(400, poll)

    # ── 浏览器 ─────────────────────────────────────────────────
    def _on_open_browser(self):
        if getattr(self, "_opening_browser", False):
            return
        enabled = self._enabled_sites()
        if not enabled:
            QMessageBox.information(self, "提示", "请至少勾选一个采集站点。")
            return
        self.session.set_enabled(enabled)
        self._opening_browser = True
        self.btn_open.setEnabled(False)
        self._on_log("正在启动共享浏览器（需浏览器的站点各占一个标签）…")

        def task():
            try:
                ok = self.session.start_browser(headed=True)
                if ok:
                    self.session._log("✅ 共享浏览器已就绪，可点「开始采集」")
                else:
                    self.session._log("⚠ 浏览器启动失败，请确认已安装 Google Chrome 后重试")
            finally:
                self._opening_browser = False

        threading.Thread(target=task, daemon=True).start()

        def poll():
            if getattr(self, "_opening_browser", False):
                QTimer.singleShot(500, poll)
            else:
                self.btn_open.setEnabled(True)
        QTimer.singleShot(500, poll)

    # ── 采集控制 ───────────────────────────────────────────────
    def _on_clear_seen_store(self):
        if self.worker and self.worker.isRunning():
            QMessageBox.information(self, "提示", "采集进行中，请先停止后再清空种子库。")
            return
        if QMessageBox.question(self, "确认", "确定清空全部平台历史种子库？清空后将对同一批关键词从头重采。") \
                == QMessageBox.Yes:
            self.session.clear_seen_store()
            QMessageBox.information(self, "完成", "历史种子库已清空。")

    def _final_save_sync(self):
        """同步收尾保存：复用自动保存的写盘方法，把最后不满一个批次的尾部数据也落盘。
        不然「停止→立刻开新任务」会让这批内存数据直接消失（自动保存是每 20 条才触发）。"""
        if not self.session_dir:
            return False
        if self.autosave and self.autosave.isRunning():
            self.autosave.wait(8000)   # 等上一个写盘线程退出，避免同文件互踩
        rows = self.session.snapshot_results()
        if not rows:
            return False
        try:
            ts_dir = Path(self.session_dir)
            ts_dir.mkdir(parents=True, exist_ok=True)
            raw = ts_dir / (_tasks.RAW_FILE_NAME if HAS_YZD else "企业数据_原始.xlsx")
            cl = ts_dir / "企业数据_清洗后入库数据.xlsx"
            worker = AutoSaveWorker(self.session, rows, self.session_dir)
            worker._write_raw(raw, rows)
            worker._write_cleaned(cl, rows)
            self._on_log(f"💾 收尾保存 {len(rows)} 条 → {raw}")
            return True
        except PermissionError:
            self._on_log("⚠ 收尾保存失败：Excel/WPS 正打开两份结果表占用，请关闭后手动点「导出 Excel」兼容尾部数据")
            return False
        except Exception as e:
            self._on_log(f"⚠ 收尾保存失败：{e}")
            return False

    def _on_reset_task(self):
        """「开始新任务」（需求：换一批关键词从零重采）：
        先同步落盘收尾→清空本次已采数据/各站账目/进度条/计时/看板 →
        释放当前时间戳文件夹（_first_start/session_dir 置空，下次「开始采集」自动新建）。
        跨运行历史种子库默认不动（清了就丢去重成果，下轮全量重请反而更容易被封 IP），
        要不要一并清由用户在对话框里勾选决定。"""
        if self.worker and self.worker.isRunning():
            QMessageBox.information(self, "提示",
                                    "采集进行中。请先点「停止」并等日志出现「采集完成/已停止」后再开新任务。")
            return
        total = len(self.session.snapshot_results())
        dlg = QDialog(self)
        dlg.setWindowTitle("开始新任务 · 确认复位范围")
        dlg.setModal(True)
        dv = QVBoxLayout(dlg)
        tip = (f"当前已采 {total} 条数据、已保存到：\n{self.session_dir or '（尚未建立数据文件夹）'}\n\n"
               "复位会执行：先把尾部数据落盘 → 清空内存结果与各站进度账目 →\n释放当前文件夹（下次「开始采集」新建时间戳文件夹）→\n"
               "进度条/计时/看板/状态行归零。已保存的 Excel 与文件夹本身不会被删除。\n\n请勾选额外项：" if total else
               "当前内存里没有已采数据（本次无需落盘）。\n复位会清空各站进度账目并释放当前数据文件夹，\n"
               "下次点「开始采集」会新建一个时间戳文件夹。\n\n请勾选额外项：")
        dv.addWidget(QLabel(tip))
        chk_seen = QCheckBox("同时清空历史种子库（跨运行去重记录一起清，同一批关键词也彻底从头重采）")
        chk_seen.setToolTip("不勾：新关键词仍从零开文件夹、进度归零，但老关键词已采过的商家仍会命中跳过（省请求=防封）。\n"
                            "勾选：完全当全新任务，所有商家重新请求一遍（短时间重请会抬高防采风险）。")
        chk_kw = QCheckBox("同时清空关键词输入框（换一批关键词时用）")
        chk_log = QCheckBox("同时清空日志区")
        for c in (chk_seen, chk_kw, chk_log):
            dv.addWidget(c)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("确认开始新任务")
        bb.button(QDialogButtonBox.Cancel).setText("取消")
        bb.accepted.connect(dlg.accept)
        bb.rejected.connect(dlg.reject)
        dv.addWidget(bb)
        if dlg.exec_() != QDialog.Accepted:
            return

        if total:
            self._final_save_sync()
        if chk_seen.isChecked():
            self.session.clear_seen_store()
            self._on_log("🧹 历史种子库已清空（各平台跨运行去重记录全部删除）")
        self.session.reset_results()

        # 释放当前数据文件夹：_on_start 里 not self._first_start 分支会重建新时间戳目录
        self._first_start = None
        self.session_dir = None
        self._task_id = ""
        if chk_kw.isChecked():
            self.kw_text.clear()
            self.keywords = []
            self._kw_source_file = ""
            self.lbl_kw.setText("识别到 0 个关键词")
        if chk_log.isChecked():
            self.log.clear()
        self._last_autosave_bucket = 0
        self._bar_pct = 0
        self._bar_companies = 0
        self._manual_stop = False
        self._paused = False
        self._held_seconds = 0.0
        self._hold_from = 0.0
        self._crawl_elapsed = 0.0
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setFormat("就绪")
        self.timer_lbl.setText("⏱ 00:00:00")
        self.lbl_stats.setText("就绪")
        self.lbl_board.setText("")
        self.lbl_save_info.setText("")
        self.lbl_board.setVisible(False)
        self.lbl_save_info.setVisible(False)
        self._hide_banner()
        self._close_block_alert()
        self.btn_start.setEnabled(True)
        self.btn_pause.setEnabled(False)
        self.btn_stop.setEnabled(False)
        self.btn_pause.setText("暂停")
        self._on_log("🆕 已清空任务：本次数据已落盘并从内存清空，采集进度/计时/看板归零，"
                     "数据文件夹已释放——下次点「开始采集」会新建一个时间戳文件夹从头重采。")
        QMessageBox.information(self, "已复位",
                                "任务已清空：进度、内存数据、计时、看板均归零。\n"
                                "下次点「开始采集」会新建一个时间戳文件夹保存新数据。")

    def _on_start(self):
        if not self.keywords:
            self._on_parse()
        has_co = bool(self._co_source_file)
        want_qzd = self.chk_waste.isChecked()
        # 关键词、公司名单、废数据补采三者全空才无事可做（允许“只跑企知道名单”模式）
        if not self.keywords and not has_co and not want_qzd:
            QMessageBox.information(self, "提示", "请先输入/导入关键词，或选择公司名单文件。")
            return
        enabled = self._enabled_sites()
        if not enabled and not want_qzd:
            QMessageBox.information(self, "提示", "请至少勾选一个采集平台。")
            return
        if self.worker and self.worker.isRunning():
            return
        self.session.set_enabled(enabled)
        self.session.configure(delay_between=self.spin_delay.value(),
                               max_pages=0,      # 不限页数：翻到没有新数据自然切下一个关键词
                               relevance_guard=self.chk_guard.isChecked(),
                               incremental=self.chk_incr.isChecked())
        if not self._first_start:
            self._first_start = datetime.now().strftime("%Y%m%d_%H%M%S")
            base = _tasks.tool_dir(self._get_desktop_dir()) if HAS_YZD else \
                (self._get_desktop_dir() / "B2B多平台企业数据自动采集工具")
            self.session_dir = str(base / self._first_start)
        # 绑定待采集公司名单（任务文件夹里生成「待采集公司名单.xlsx」，与原数据/清洗入库同层级）：
        # 上传名单入列 + 清洗废数据持续回流；want_qzd=界面“启用企知道补采”开关。
        try:
            qzd_on, _added = self.session.bind_pending(self.session_dir, self._co_source_file, want_qzd)
        except Exception as e:
            qzd_on = False
            self._on_log(f"⚠ 待采名单绑定失败（不影响关键词采集）：{str(e).splitlines()[0][:120]}")
        self._register_task()
        self._paused = False
        self.btn_start.setEnabled(False)
        self.btn_pause.setEnabled(True); self.btn_pause.setText("暂停")
        self.btn_stop.setEnabled(True)
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self._bar_pct = 0
        self._bar_companies = 0
        self._last_autosave_bucket = 0
        self._start_timer()
        self._update_bar()
        _sites = "、".join(SITE_NAMES[s] for s in enabled) if enabled else "无关键词平台"
        _tail = "｜企知道名单补采已启用" if qzd_on else ""
        self._on_log(f"▶ 开始采集：{_sites}｜{len(self.keywords)} 个关键词{_tail}")
        self.worker = CrawlWorker(self.session, self.keywords)
        self.worker.start()
        self.request_autosave()
        if self.chk_bg.isChecked():
            # 标签是各站线程启动后才 pin 的：延迟 5s/30s 各试一次最小化，
            # 确保晚建立的标签所在窗口也能后台化。
            QTimer.singleShot(5000, self._apply_background)
            QTimer.singleShot(30000, self._apply_background)

    def _apply_background(self):
        if self.chk_bg.isChecked() and self.worker and self.worker.isRunning():
            threading.Thread(target=lambda: self.session.set_browser_background(True),
                             daemon=True).start()

    def _on_pause(self):
        if self._paused:
            self.session.resume_crawl()
            self._paused = False
            self.btn_pause.setText("暂停")
            self._hold_timer(False)
        else:
            self.session.pause_crawl()
            self._paused = True
            self.btn_pause.setText("继续")
            self._hold_timer(True)

    def _on_stop(self):
        self._manual_stop = True
        self.session.stop_crawl()

    def _on_finished(self):
        self.btn_start.setEnabled(True)
        self.btn_pause.setEnabled(False)
        self.btn_stop.setEnabled(False)
        self.btn_pause.setText("暂停")
        self._stop_timer()
        self.progress.setRange(0, 100)
        self.progress.setValue(100)
        total = len(self.session.snapshot_results())
        self.progress.setFormat(f"完成｜共 {total} 家｜耗时 {self._format_time(self._crawl_elapsed)}")
        self._refresh_board()
        self._hide_banner()
        self._close_block_alert()   # 停止/完成都要收掉验证弹窗，不然留着会误导用户还在等
        self.request_autosave()
        stopped = bool(getattr(self, "_manual_stop", False))
        self._manual_stop = False
        self._finish_task(_tasks.TASK_STATUS_STOPPED if stopped else _tasks.TASK_STATUS_DONE,
                          len(self.session.snapshot_results()))
        if not stopped:
            self._maybe_auto_import("采集完成")

    # ── 自动保存 / 导出 ────────────────────────────────────────
    def request_autosave(self):
        if not self.session_dir:
            return
        if self.autosave and self.autosave.isRunning():
            self._pending_save = True
            return
        rows = self.session.snapshot_results()
        if not rows:
            return
        self.autosave = AutoSaveWorker(self.session, rows, self.session_dir)
        self.autosave.saved.connect(self._on_saved)
        self.autosave.failed.connect(lambda m: self._on_log("⚠ " + m))
        self.autosave.finished.connect(self._after_autosave)
        self.autosave.start()

    def _after_autosave(self):
        if self._pending_save:
            self._pending_save = False
            self.request_autosave()

    def _on_saved(self, path, n):
        now = time.strftime("%H:%M:%S")
        self._on_log(f"💾 已自动保存 {n} 条 → {path}")
        self.lbl_save_info.setText(f"💾 上次自动保存：{now}（{n} 条）→ {path}")
        self.lbl_save_info.setVisible(True)
        if HAS_YZD and self._task_id:
            try:
                _tasks.update_task(self._task_id, count=n)
            except Exception:
                pass

    def _on_clean_now(self):
        """「清洗数据」：把当前已采数据立即跑一遍清洗，落盘原始+清洗入库两份，
        废数据经 clean_records 的 waste_sink 回流待采名单。数据量有限，同步执行秒级完成。"""
        rows = self.session.snapshot_results()
        if not rows:
            QMessageBox.information(self, "提示", "还没有采集到数据，无法清洗。")
            return
        if self.autosave and self.autosave.isRunning():
            self.autosave.wait(8000)   # 等上一个写盘线程退出，避免同文件互踩
        if not self.session_dir:
            self._first_start = datetime.now().strftime("%Y%m%d_%H%M%S")
            base = _tasks.tool_dir(self._get_desktop_dir()) if HAS_YZD else \
                (self._get_desktop_dir() / "B2B多平台企业数据自动采集工具")
            self.session_dir = str(base / self._first_start)
        try:
            ts_dir = Path(self.session_dir)
            ts_dir.mkdir(parents=True, exist_ok=True)
            raw = ts_dir / (_tasks.RAW_FILE_NAME if HAS_YZD else "企业数据_原始.xlsx")
            cl = ts_dir / "企业数据_清洗后入库数据.xlsx"
            AutoSaveWorker._write_raw(raw, rows)
            # 清洗只跑一次（避免重复触发 waste_sink）：直接写已算好的 clean_rows
            clean_rows = self.session.clean_records(rows, log=lambda m: None)
            template = _cleaner.locate_template([self.session_dir, str(Path.home() / "Downloads")])
            tmp = str(cl) + ".tmp"
            _cleaner.write_clean_excel(clean_rows, tmp, template_path=template, log=lambda m: None)
            os.replace(tmp, str(cl))
            self.lbl_save_info.setText(f"🧼 上次清洗：{time.strftime('%H:%M:%S')}（入库 {len(clean_rows)} 条）→ {cl}")
            self.lbl_save_info.setVisible(True)
            self._on_log(f"🧼 已清洗 {len(rows)} 条 → 入库 {len(clean_rows)} 条｜{cl}")
            tail = "，企知道通道会逐家补采。" if self.session.qzd_enabled else "（启用企知道补采后可被逐家补采）。"
            QMessageBox.information(
                self, "清洗完成",
                f"原始 {len(rows)} 条已清洗，入库 {len(clean_rows)} 条。\n\n"
                f"清洗后入库数据：\n{cl}\n\n"
                f"无联系方式的废数据已自动加入「待采集公司名单」{tail}")
        except PermissionError:
            self._on_log("⚠ 清洗失败：结果表正被 Excel/WPS 打开占用，请关闭后重试")
            QMessageBox.warning(self, "清洗失败", "结果表正被 Excel/WPS 打开占用，请关闭后重试。")
        except Exception as e:
            self._on_log(f"⚠ 清洗失败：{e}")
            QMessageBox.warning(self, "清洗失败", str(e))

    def _on_export(self):
        rows = self.session.snapshot_results()
        if not rows:
            QMessageBox.information(self, "提示", "还没有采集到数据。")
            return
        default = f"B2B多平台企业数据_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        path, _ = QFileDialog.getSaveFileName(self, "导出 Excel",
                                              str(self._get_desktop_dir() / default), "Excel 文件 (*.xlsx)")
        if not path:
            return
        ok = self.session.export_excel(path)
        if ok:
            cpath = str(Path(path).with_name(Path(path).stem + "_清洗后入库数据.xlsx"))
            try:
                res = self.session.export_cleaned_excel(cpath)
            except Exception as e:
                res = ""
                self._on_log(f"清洗导出失败：{e}")
            if res:
                QMessageBox.information(self, "导出完成",
                                        f"原始明细：\n{path}\n\n清洗后入库数据（对齐客户库模板）：\n{cpath}")
            else:
                QMessageBox.information(self, "导出完成",
                                        f"原始明细已导出（无满足入库条件的数据）：\n{path}")

    # ── 云展动力：任务登记 / 绑定 / 导入 ───────────────────────
    def _current_task_id(self):
        return Path(self.session_dir).name if self.session_dir else ""

    def _register_task(self):
        if not HAS_YZD or not self.session_dir:
            return
        try:
            tid = self._current_task_id()
            name_src = [self._kw_source_file] if self._kw_source_file else list(self.keywords)
            task = _tasks.new_task(tid, name_src, self.session_dir,
                                   count=len(self.session.snapshot_results()))
            existing = next((t for t in _tasks.load_tasks() if t.get("id") == tid), None)
            if existing:
                task["started_at"] = existing.get("started_at") or task["started_at"]
            _tasks.upsert_task(task)
            self._task_id = tid
            if existing is None:
                self._on_log(f"📁 已登记采集任务：{task['name']}")
        except Exception as e:
            self._on_log(f"⚠ 采集任务登记失败：{e}")

    def _finish_task(self, status, count):
        if not HAS_YZD or not self._task_id:
            return
        try:
            if _tasks.finish_task(self._task_id, status, count):
                self._on_log(f"📁 任务「{status}」已记录：合计 {count} 条")
        except Exception as e:
            self._on_log(f"⚠ 任务结束状态写入失败：{e}")

    def _refresh_yzd_status(self):
        if not hasattr(self, "yzd_status_lbl"):
            return
        if not HAS_YZD:
            self.yzd_status_lbl.setText("云展模块未加载（缺 yzd_client / task_registry）")
            self.yzd_status_lbl.setStyleSheet(_status_badge_style("error"))
            self.bind_yzd_btn.setEnabled(False)
            self.import_yzd_btn.setEnabled(False)
            return
        if self._yzd_phone:
            self.yzd_status_lbl.setText(f"已绑定云展账号：{self._yzd_phone}")
            self.yzd_status_lbl.setStyleSheet(_status_badge_style("ok"))
            self.bind_yzd_btn.setText("更换绑定账号")
        else:
            self.yzd_status_lbl.setText("未绑定云展动力账号")
            self.yzd_status_lbl.setStyleSheet(_status_badge_style("neutral"))
            self.bind_yzd_btn.setText("绑定云展动力账号")

    def _bind_yzd_account(self):
        if not HAS_YZD:
            QMessageBox.warning(self, "提示", "云展模块未加载，无法绑定账号")
            return
        if self._yzd_worker is not None and self._yzd_worker.isRunning():
            return
        tip = "请输入云展动力手机号"
        if self._yzd_phone:
            tip += f"（当前绑定：{self._yzd_phone}）"
        phone, ok = QInputDialog.getText(self, "绑定云展动力账号", tip, QLineEdit.Normal, self._yzd_phone)
        if not ok:
            return
        phone = (phone or "").strip().replace(" ", "").replace("-", "")
        if not phone or not _yzd.PHONE_RE.match(phone):
            QMessageBox.warning(self, "提示", "请输入正确的手机号")
            return
        self.bind_yzd_btn.setEnabled(False)
        self._on_log(f"🔗 正在校验云展动力账号 {phone} …")
        self._yzd_worker = YzdBindWorker(phone)
        self._yzd_worker.finished_bind.connect(self._on_bind_finished)
        self._yzd_worker.start()

    def _on_bind_finished(self, ok, msg):
        self.bind_yzd_btn.setEnabled(HAS_YZD)
        if ok:
            self._yzd_phone = msg
            saved = _tasks.save_account(msg)
            self._on_log(f"✅ 已绑定云展动力账号：{msg}" + ("" if saved else "（本地保存失败，下次启动需重新绑定）"))
            QMessageBox.information(self, "绑定成功", f"已绑定云展动力账号：{msg}")
        else:
            self._on_log(f"⚠ 绑定云展账号失败：{msg}")
            QMessageBox.warning(self, "绑定失败", msg)
        self._refresh_yzd_status()

    def _open_yzd_import(self):
        if not HAS_YZD:
            QMessageBox.warning(self, "提示", "云展模块未加载，无法导入客户库")
            return
        if self._yzd_importing:
            QMessageBox.information(self, "提示", "正在导入云展客户库，请等当前导入完成")
            return
        if not self._yzd_phone:
            QMessageBox.warning(self, "提示", "请先绑定云展动力账号")
            return
        try:
            added = _tasks.backfill_tasks()
            if added:
                self._on_log(f"📁 已补登记历史采集会话 {added} 个，可在「导入其他」里选中")
        except Exception as e:
            self._on_log(f"⚠ 历史会话补登记失败：{e}")
        dlg = YzdGroupDialog(self._yzd_phone, self)
        self._yzd_dialog = dlg
        if dlg.exec_() != QDialog.Accepted:
            return
        group = dlg.chosen_group or {}
        tasks = dlg.chosen_tasks if dlg.mode == YzdGroupDialog.MODE_OTHERS else None
        self._start_yzd_import(group, tasks)

    def _on_auto_import_toggled(self, checked):
        if not checked:
            self._auto_import_group = None
            self._on_log("已关闭「采集完自动导入云展客户库」，采集完成后需手动点导入")
            return
        if not HAS_YZD:
            self._revert_auto_import("云展模块未加载，无法自动导入客户库")
            return
        if self._yzd_importing:
            self._revert_auto_import("当前正在导入，请等它结束后再开启自动导入")
            return
        if not self._yzd_phone:
            self._revert_auto_import("请先点「绑定云展动力账号」，绑定后才能选自动导入的目标分组")
            return
        dlg = YzdGroupDialog(self._yzd_phone, self, pick_group_only=True)
        self._yzd_dialog = dlg
        if dlg.exec_() != QDialog.Accepted or not dlg.chosen_group:
            self._auto_import_group = None
            chk = getattr(self, "auto_import_check", None)
            if chk is not None:
                chk.blockSignals(True); chk.setChecked(False); chk.blockSignals(False)
            self._on_log("ℹ 没选定分组，「采集完自动导入」已取消")
            return
        self._auto_import_group = dlg.chosen_group
        flag = self._flag_from_group(self._auto_import_group)
        self._on_log(f"🤖 已设置采集完自动导入：分组「{self._auto_import_group.get('groupName') or ''}」"
                     f"（{'国外' if flag else '国内'}客户库），采集完成后自动上传并在日志提示结果")

    def _revert_auto_import(self, msg):
        self._auto_import_group = None
        chk = getattr(self, "auto_import_check", None)
        if chk is not None:
            chk.blockSignals(True); chk.setChecked(False); chk.blockSignals(False)
        self._on_log(f"⚠ {msg}")
        QMessageBox.warning(self, "无法开启自动导入", msg)

    def _maybe_auto_import(self, reason):
        chk = getattr(self, "auto_import_check", None)
        if chk is None or not chk.isChecked():
            return
        if self._yzd_importing:
            return
        group = self._auto_import_group or {}
        if not group.get("id"):
            self._on_log("⚠ 已勾选自动导入但没拿到分组，本次不自动上传；可手动点「导入云展客户库」")
            return
        if not self.session.snapshot_results():
            self._on_log("ℹ 本次没采到数据，自动导入已跳过")
            return
        tid = self._current_task_id()
        if not tid:
            self._on_log("⚠ 本次采集没登记到任务索引（会话目录未创建），自动导入已跳过")
            return
        tasks = [t for t in _tasks.load_tasks() if t.get("id") == tid] or [
            {"id": tid, "name": f"本次采集 {tid}", "folder": str(self.session_dir)}]
        self._on_log(f"🤖 {reason}，按勾选自动导入云展客户库：分组「{group.get('groupName') or ''}」")
        self._start_yzd_import(group, tasks)

    def _build_jobs(self, tasks):
        jobs = []
        tid = self._current_task_id()
        rows = self.session.snapshot_results()
        for t in tasks or []:
            name = t.get("name") or t.get("id") or "未命名任务"
            if t.get("id") and t.get("id") == tid and rows:
                jobs.append({"id": t.get("id"), "name": name, "records": rows, "raw_file": ""})
                continue
            raw_file = t.get("raw_file") or ""
            if not raw_file and t.get("folder"):
                raw_file = os.path.join(str(t.get("folder")), _tasks.RAW_FILE_NAME)
            jobs.append({"id": t.get("id") or "", "name": name, "records": None, "raw_file": raw_file})
        return jobs

    def _start_yzd_import(self, group, tasks):
        if self._yzd_importing:
            return
        if tasks is None:
            latest = _tasks.latest_task()
            if not latest:
                QMessageBox.warning(self, "提示", "暂无采集任务：请先在主界面完成一次采集")
                return
            jobs = self._build_jobs([latest])
        else:
            jobs = self._build_jobs(tasks)
        if not jobs:
            QMessageBox.warning(self, "提示", "没有可导入的采集任务")
            return
        empty = [j for j in jobs if not j["raw_file"] and not (j["records"] or [])]
        if len(empty) == len(jobs):
            QMessageBox.warning(self, "提示", "选中的采集任务里没有数据")
            return
        import tempfile
        try:
            tmp_dir = tempfile.mkdtemp(prefix="yzd_import_")
        except Exception as e:
            QMessageBox.critical(self, "导入失败", f"无法创建临时上传目录：{e}")
            return
        self._yzd_jobs = jobs
        self._yzd_importing = True
        group_name = group.get("groupName") or ""
        self._yzd_group_name = group_name
        flag = self._flag_from_group(group)
        self._on_log(f"⬆ 开始导入云展客户库：分组「{group_name}」（{'国外' if flag else '国内'}）· 共 {len(jobs)} 个任务")
        self.setEnabled(False)
        prog = YzdProgressDialog(None, f"正在导入云展客户库：{group_name}")
        self._yzd_progress = prog
        prog.show()
        worker = YzdImportWorker(self.session, self._yzd_phone, jobs, group.get("id"), flag, tmp_dir)
        worker.logged.connect(self._on_log)
        worker.progress.connect(self._on_yzd_import_progress)
        worker.finished_all.connect(self._on_yzd_import_finished)
        self._yzd_import_worker = worker
        worker.start()

    @staticmethod
    def _flag_from_group(group):
        value = group.get("foreignFlag")
        try:
            return int(value) if value is not None else 0
        except (TypeError, ValueError):
            return 0

    def _on_yzd_import_progress(self, done, total, text):
        if self._yzd_progress is not None:
            self._yzd_progress.set_text(text, f"已完成 {done}/{total} 个任务。导入期间请勿关闭程序。")

    def _unlock_yzd_import(self):
        self._yzd_importing = False
        self.setEnabled(True)
        if self._yzd_progress is not None:
            try:
                self._yzd_progress.allow_close()
                self._yzd_progress.close()
            except Exception:
                pass
            self._yzd_progress = None
        self.activateWindow()

    def _on_yzd_import_finished(self, results, fatal):
        if fatal:
            self._on_log(f"❌ 导入异常：{fatal}")
            self._unlock_yzd_import()
            QMessageBox.critical(self, "导入失败", fatal)
            return
        lines = []
        saved_total = 0
        for idx, item in enumerate(results or []):
            name, imp, sav, fail, err = item
            saved_total += sav
            if err:
                lines.append(f"· {name}：导入失败（{err}）")
                self._on_log(f"❌ {name}：{err}")
            else:
                lines.append(f"· {name}：上传 {imp} 条 / 入库成功 {sav} 条 / 失败 {fail} 条")
                self._on_log(f"✅ {name}：上传 {imp} / 入库 {sav} / 失败 {fail}")
            if idx < len(self._yzd_jobs):
                job = self._yzd_jobs[idx]
                if job.get("id"):
                    try:
                        _tasks.update_task(job["id"],
                                           yzd_import_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                                           yzd_group=self._yzd_group_name, yzd_saved=sav)
                    except Exception as e:
                        self._on_log(f"⚠ 任务导入结果回写失败：{e}")
        self._unlock_yzd_import()
        self._on_log("ℹ 云展接口不做重复校验，同一批数据重复导入会在客户库里产生多条记录")
        summary = "\n".join(lines) if lines else "没有任何任务被处理"
        if saved_total == 0:
            QMessageBox.warning(self, "导入结果", f"未入库任何客户：\n{summary}\n\n请查看运行日志里的后端返回。")
        else:
            QMessageBox.information(self, "导入结果", f"共入库 {saved_total} 条客户：\n\n{summary}")
        import shutil
        worker, self._yzd_import_worker = self._yzd_import_worker, None
        if worker is not None:
            try:
                shutil.rmtree(worker.tmp_dir, ignore_errors=True)
            except Exception:
                pass

    def closeEvent(self, ev):
        if self._yzd_importing:
            QMessageBox.information(self, "请等一下", "正在导入云展客户库，请等导入完成后再关闭程序。")
            ev.ignore()
            return
        try:
            self.session.stop_crawl()
            if self.autosave and self.autosave.isRunning():
                self.autosave.wait(3000)
            self.session.close_browser()
        except Exception:
            pass
        super().closeEvent(ev)


def _app_stylesheet():
    """全局样式：只动「层级」不动颜色主题，跟随系统深浅色。
    颜色一律不写死背景/文字（写死 #333 在深色模式下会隐形，就是旧版看板那块空白），
    只给主按钮上强调色、给分区标题降一级字重不抢戏。"""
    return """
    QGroupBox {
        font-weight: bold;
        padding-top: 10px;
        margin-top: 6px;
        border: 1px solid palette(mid);
        border-radius: 6px;
    }
    QGroupBox::title {
        subcontrol-origin: margin;
        left: 10px;
        padding: 0 4px;
    }
    QPushButton#primary {
        background-color: #1a73e8;
        color: #ffffff;
        border: none;
        padding: 0 18px;
        border-radius: 5px;
        font-weight: bold;
    }
    QPushButton#primary:hover   { background-color: #1765cc; }
    QPushButton#primary:pressed { background-color: #1454a8; }
    QPushButton#primary:disabled { background-color: palette(mid); color: palette(text); }
    QProgressBar {
        border: 1px solid palette(mid);
        border-radius: 5px;
        text-align: center;
        height: 26px;
    }
    QProgressBar::chunk { background-color: #1a73e8; border-radius: 4px; }
    QTextEdit, QPlainTextEdit {
        border: 1px solid palette(mid);
        border-radius: 6px;
    }
    """


def main():
    app = QApplication(sys.argv)
    app.setStyleSheet(_app_stylesheet())
    w = MainWindow()
    w.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
