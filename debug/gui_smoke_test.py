# GUI 冒烟自测（offscreen 跑，不在用户屏幕弹窗）：本轮四项需求的界面端到端验证
# ①入库口径：清洗后表剔掉「手机号/固话/邮箱/微信号」全空行，原始表仍全量
# ②③被拦→弹窗提醒 / 解除→收窗（横幅+弹窗+计时接管三处联动）
# ④「开始新任务」：复位前同步落盘、进度/看板/文件夹归零、历史种子库默认不动
import os
import sys
import json
import shutil
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")   # 无头跑 Qt，绝不打扰用户桌面
sys.path.insert(0, "/Users/jhnoplan/Desktop/PythonDemo/b2bcollect_build")

from PyQt5.QtWidgets import QApplication, QDialog, QCheckBox, QMessageBox
from openpyxl import load_workbook

import b2bcollect_gui as gui

ok = 0


def check(name, cond):
    global ok
    assert cond, f"FAIL: {name}"
    ok += 1
    print(f"  ✓ {name}")


app = QApplication.instance() or QApplication([])
w = gui.MainWindow()
w.show()        # offscreen 下窗口不会真上屏，但必须 show 才能让子控件 isVisible 生效

tmp = Path(tempfile.mkdtemp(prefix="b2b_gui_smoke_"))

# 站点与数据目录就位：结果表与种子库全部改指临时目录——绝不写用户桌面的历史
# 文件夹，也不跟用户正跑着的 GUI 争同一份 seen_store（并发下它一直在写）
seen_dir_tmp = tmp / "seen"
seen_dir_tmp.mkdir(parents=True, exist_ok=True)
SEEN_FIXTURE = json.dumps({"顺企网": {"测试旧公司": {"ct": 1, "st": "成功"}}}, ensure_ascii=False)
w.session.set_enabled(["mic", "b2b168", "11467", "aipage"])   # gys 供应商网已下线
for c in w.session.crawlers.values():
    c._commit_to_store = lambda rec: None      # 假记录绝不写库
    c._data_dir = seen_dir_tmp
    c._seen_store_path = seen_dir_tmp / f"seen_store_{c.PLATFORM_SLUG}.json"
    c._seen_store_path.write_text(SEEN_FIXTURE, encoding="utf-8")
    c._seen_store = json.loads(SEEN_FIXTURE)
seen_dir = w.session.crawlers["11467"]._seen_store_path
seen_bytes = seen_dir.read_bytes()
w.session_dir = str(tmp / "20261028_160000")
w._first_start = "20261028_160000"

# 三条记录：两条有联系方式，一条四项全空（必须从清洗后表里被剔掉，但留在原始表）
SAMPLES = [
    {"公司名称": "深圳市甲乙太阳能科技有限公司", "数据来源": "顺企网",
     "手机": "13912345678", "电话": "0755-88886666", "邮箱": "a@b.com", "状态": "成功"},
    {"公司名称": "佛山市丙丁照明器材厂", "数据来源": "八方资源网",
     "手机": "", "电话": "0757-82807298", "邮箱": "", "状态": "成功"},
    {"公司名称": "某从未公开联系方式的公司", "数据来源": "中国制造网",
     "手机": "", "电话": "", "邮箱": "", "状态": "无电话(未公开)"},
]
SLUG_OF = {"顺企网": "11467", "八方资源网": "b2b168", "中国制造网": "mic"}
for rec in SAMPLES:
    w.session.crawlers[SLUG_OF[rec["数据来源"]]]._add_result(rec)

# ══ ① 收尾保存：原始全量 + 清洗后按新口径剔除 ══
check("界面：三条记录进入内存结果", len(w.session.snapshot_results()) == 3)
saved = w._final_save_sync()
raw_p = tmp / "20261028_160000" / gui._tasks.RAW_FILE_NAME if gui.HAS_YZD else \
    tmp / "20261028_160000" / "企业数据_原始.xlsx"
clean_p = tmp / "20261028_160000" / "企业数据_清洗后入库数据.xlsx"
check("收尾保存成功落盘（停止后不满一个自动保存批次的数据不丢）", saved and raw_p.exists())
wb = load_workbook(raw_p, read_only=True)
ws = wb.worksheets[0]
raw_rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(x for x in r)]
wb.close()
check(f"原始表全量 3 行不过滤（补查线索不丢）（实际 {len(raw_rows)} 行）", len(raw_rows) == 3)
wb2 = load_workbook(clean_p, read_only=True)
ws2 = wb2.worksheets[0]
data_rows = [r for r in ws2.iter_rows(min_row=2, values_only=True) if any(x for x in r)]
wb2.close()
names = [str(r[1] or "") for r in data_rows]
check(f"清洗后入库表只剩有联系方式的 2 行（实际 {len(data_rows)} 行）", len(data_rows) == 2)
check("四项全空那条确实被剔除", "某从未公开联系方式的公司" not in "".join(names))
check("有手机号那条仍在（不误杀）", any("甲乙" in n for n in names))

# ══ ②③ 被拦提醒：横幅 + 独立弹窗 + 计时接管；解除后全部收起 ══
# 拦截分型（IP 限流 vs 可人工过验证）走两套文案，1688 与八方不能再共用一句“已开页”
from cdp_base import BLOCK_TAG_IP, BLOCK_TAG_CAP

check("初始没有提醒弹窗", w._block_alert is None)
w._on_blocked(BLOCK_TAG_IP + "八方资源网：站点返回 HTTP 502（IP 级限流/冷却，非验证码）")
app.processEvents()
check("被拦时顶部横幅可见（带上站名与分型，不带看不完的细节）",
      w.banner.isVisible() and "八方资源网" in w.banner.text() and "IP 级限流" in w.banner.text())
check("IP 类横幅直说「不需要你操作」（客户就是被「手动完成验证」误导到满屏找入口）",
      "不需要你操作" in w.banner.text() and "手动完成验证" not in w.banner.text())
check("被拦时弹出独立提醒窗（旧版只改横幅，最小化后用户根本看不见）",
      w._block_alert is not None and w._block_alert.isVisible())
check("弹窗是非模态（不阻塞主窗口，用户能去操作 Chrome）", w._block_alert.isModal() is False)
txt = w._block_alert.text()
check("弹窗带上被拦站点与原因", "八方资源网" in txt and "502" in txt)
check("IP 类弹窗不再声称「已在 Chrome 里打开对应页面」（实际根本没开）",
      "已在 Chrome 里打开对应页面" not in txt)
check("IP 类弹窗不叫用户去拉滑块，而是给出真正有效的恢复手段",
      all(k in txt for k in ("重拨宽带", "采集间隔", "断点续采")))
check("弹窗里提供用户主动点的「打开该站页面自查」按钮",
      any("自查" in b.text() for b in w._block_alert.buttons()))
first_box = w._block_alert

check("IP 类弹窗不会自动拉浏览器：没点按钮就不会去借标签",
      w._block_site == "八方资源网" and w.session.crawlers["b2b168"].pin_tab_id == "")
check_calls = []
w.session.open_site_check = lambda name="": (check_calls.append(name),
                                             (True, "已在 Chrome 里打开（测试替身）"))[1]
for b in w._block_alert.buttons():
    if "自查" in b.text():
        b.click()
app.processEvents()
import time as _t
_t.sleep(0.6)
app.processEvents()
check("点「自查」会把当前被拦站交给会话去开页（用户主动才能开）",
      check_calls == ["八方资源网"])

w._on_blocked(BLOCK_TAG_CAP + "中国制造网：页面弹出验证控件（#tcaptcha_transform）")
app.processEvents()
check("真验证码场景弹窗仍明确要求人工过完", "手动过完" in w._block_alert.text())
check("验证码类弹窗不再把「已打开」说死（借不到标签的站会被谎话误导）",
      "已把该站页面在 Chrome 里打开" not in w._block_alert.text()
      and "没看到就点下面的自查" in w._block_alert.text())
check("验证码类横幅仍明确要求去 Chrome 里手动过验证，并提「没看到页面就自查」",
      "Chrome 窗口里手动过" in w.banner.text() and "自查" in w.banner.text())
check("多站同时被拦不叠窗（复用同一弹窗实例）", w._block_alert is first_box)
check("复用弹窗文案更新成最新站点", "中国制造网" in w._block_alert.text())
check("验证等待期间接管采集计时", w._hold_from > 0)
w._on_block_cleared("八方资源网")
app.processEvents()
check("解除后提醒弹窗关闭", w._block_alert is None)
check("解除后计时交还给采集耗时", w._hold_from == 0.0)
w._show_banner("测试残留", "#fff", "#000")
w._on_blocked(BLOCK_TAG_IP + "顺企网：HTTP 429（IP 级限流/冷却，非验证码）")
app.processEvents()
w._on_finished()
check("结束/停止同样收回提醒弹窗与横幅（不留挂着误导用户）",
      w._block_alert is None and w.banner.isVisible() is False)

# ══ ④ 开始新任务：对话框可确认，复位彻底但不碰历史去重成果 ══
accepted_box = {}


def fake_exec(self):
    accepted_box["dlg"] = self
    for cb in self.findChildren(QCheckBox):
        if "种子库" in cb.text():
            cb.setChecked(False)          # 默认口径：保留历史种子库
    return QDialog.Accepted


QDialog.exec_ = fake_exec
QMessageBox.information = staticmethod(lambda *a, **k: QMessageBox.Ok)
w.btn_start.setEnabled(False)
w.btn_pause.setEnabled(True)
w.btn_stop.setEnabled(True)
w.progress.setValue(66)
w._bar_pct = 66
w._bar_companies = 3
w._last_autosave_bucket = 2
w.lbl_board.setText("上一批看板")
w.lbl_stats.setText("上一批状态行")
w.kw_text.setPlainText("太阳能路灯")
w._on_parse()
before_seen = w.session.crawlers["11467"]._seen_store

w._on_reset_task()
check("复位对话框弹出过（用户能勾额外项）", "dlg" in accepted_box)
check("复位后内存数据清空", w.session.snapshot_results() == [])
check("复位后各站账目清空（看板数字不再延续上一批）",
      all(c.stats()["ok"] == 0 for c in w.session.crawlers.values()))
check("复位释放数据文件夹（下次「开始采集」新建时间戳目录）",
      w.session_dir is None and w._first_start is None)
check("复位后进度条/百分比/家数归零",
      w.progress.value() == 0 and w._bar_pct == 0 and w._bar_companies == 0)
check("复位后自动保存批次计数归零（新任务第一满 20 条就落盘）", w._last_autosave_bucket == 0)
check("复位后计时显示与看板/状态行清空",
      w.timer_lbl.text().endswith("00:00:00") and w.lbl_board.text() == ""
      and w.lbl_stats.text() == "就绪")
check("复位后按钮回到可开始状态",
      w.btn_start.isEnabled() and not w.btn_pause.isEnabled() and not w.btn_stop.isEnabled())
check("未勾选时关键词保留（可换关键词前先复用同一批）", w.keywords == ["太阳能路灯"])
check("未勾选时日志保留（上一轮结果仍可回看）", w.log.toPlainText() != "")
check("历史种子库内存条目未被复位清掉", w.session.crawlers["11467"]._seen_store == before_seen)
check("历史种子库磁盘文件一字未变（跨运行去重成果保住=少请求=防封）", seen_dir.read_bytes() == seen_bytes)
check("上一批的 Excel 文件没被删（历史成果保持）", raw_p.exists() and clean_p.exists())
check("复位日志写明新任务会新建文件夹",
      "新建一个时间戳文件夹" in w.log.toPlainText())

# 关键词清空分支
w.kw_text.setPlainText("光伏储能电池")
w._on_parse()


def fake_exec_kw(self):
    for cb in self.findChildren(QCheckBox):
        # 只勾「关键词输入框/日志区」：注意种子库那条文案里也带“关键词”三字，
        # 按子串乱匹会误清种子库，必须按完整选项名区分
        cb.setChecked("关键词输入框" in cb.text() or "日志区" in cb.text())
    return QDialog.Accepted


QDialog.exec_ = fake_exec_kw
w._on_reset_task()
check("勾选后关键词输入框与解析结果一并清空",
      w.kw_text.toPlainText() == "" and w.keywords == [] and "识别到 0 个关键词" in w.lbl_kw.text())
log_now = w.log.toPlainText()
check("勾选后日志区清空（只留本次复位提示，上一轮痕迹不残）",
      "💾 收尾保存" not in log_now and "🆕 已清空任务" in log_now)
check("没勾种子库选项就绝不删去重文件（文案里也带“关键词”不能误伤）",
      seen_dir.read_bytes() == seen_bytes)

# 第三轮：明确勾选「同时清空历史种子库」→ 磁盘与内存去重记录都应清空
for c in w.session.crawlers.values():
    c._seen_store = json.loads(SEEN_FIXTURE)
    c._seen_store_path.write_text(SEEN_FIXTURE, encoding="utf-8")


def fake_exec_seen(self):
    for cb in self.findChildren(QCheckBox):
        cb.setChecked("历史种子库" in cb.text())
    return QDialog.Accepted


QDialog.exec_ = fake_exec_seen
w._on_reset_task()
check("勾选后五站去重磁盘文件确实删除（下轮对同一批关键词从头重采）",
      not seen_dir.exists() or list(seen_dir_tmp.glob("seen_store_*.json")) == [])
check("勾选后内存去重条目一并清空", w.session.crawlers["11467"]._seen_store == {})
check("清空种子库仍不动已采的结果表文件", raw_p.exists() and clean_p.exists())

shutil.rmtree(tmp, ignore_errors=True)
print(f"\n全部 {ok} 项断言通过 ✅")
